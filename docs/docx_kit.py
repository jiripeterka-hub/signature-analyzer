# -*- coding: utf-8 -*-
"""Typografická sada pro wordové dokumenty o AnEP — sazba, ne obsah.

Používají ji dva generátory a každý odpovídá na jinou otázku:

    gen-schopnosti-s-obrazky.py  →  AnEP-schopnosti-s-obrazky.docx  (co aplikace umí)
    gen-manual-pokrocily.py      →  AnEP-manual-pokrocily.docx      (jak to použít)

Bydlí zvlášť, protože jinak by tu byla dvě znění týchž pěti set řádků, která se při první
změně rozejdou — a hlavně proto, že tři pasti níž nejsou vidět, dokud se dokument nevysází.
Kdo je jednou obejde, nemá je obcházet podruhé.

PASTI PYTHON-DOCX, KTERÉ TAHLE SADA ŘEŠÍ

 1. Vestavěné nadpisové styly jmenují písmo přes motiv (``w:asciiTheme``) a ten má ve Wordu
    přednost před ``font.name``. Samotné nastavení jména se neprojeví a nadpisy zůstanou
    v motivovém Calibri Light, zatímco text okolo je Segoe UI. Viz :func:`_pismo`.

 2. ``run.bold = False`` zapíše *přímé* formátování a přebije tučnost ze stylu — nadpisy pak
    zůstanou základním řezem. Tučnost se proto jen zapíná, nikdy nevypíná. Viz :func:`psat`.

 3. Zarážky v zápatí nese styl ``Footer`` a přidání odstavcové zarážky je nepřebije: číslo
    stránky skočí doprostřed. Přepsat se musí rovnou styl. Viz :func:`nastav_zapati`.

A čtvrtá, která není o python-docx, nýbrž o ECMA-376: prvek vložený mimo předepsané pořadí
dětí ``w:pPr`` či ``w:tcPr`` umí skončit hláškou o poškozeném souboru — a ta se objeví až
u čtenáře. Proto :func:`_vloz_dle_poradi` a čtyři tabulky pořadí pod paletou.

Potřebuje ``python-docx``.
"""

import re

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ── paleta ───────────────────────────────────────────────────────────────────────────────
# Zelená je z loga aplikace; na papíře se drží tmavší odstín, protože ten obrazovkový je pod
# tiskovým rozlišením bledý.
AKCENT       = '4D8A19'
AKCENT_TMAVY = '3A6A12'
INKOUST      = '1F2419'
INKOUST_SVET = '5A6355'
LINKA        = 'D3DBCB'
PLOCHA       = 'EEF2E9'   # výplň záhlaví tabulek
PLOCHA_SVET  = 'F6F8F3'   # výplň vysvětlivek
RAMECEK      = 'FAFBF8'   # výplň rámečku pro obrázek

PISMO       = 'Segoe UI'
PISMO_LEHKE = 'Segoe UI Light'
PISMO_MONO  = 'Consolas'

# Šířka sazby při okrajích 2,2 cm na stránce A4 — tolik smí být široká tabulka i rámeček.
SAZBA_CM = 16.6

# Pořadí dětí w:pPr a w:tcPr podle ECMA-376. Word sice leccos odpustí, ale prvek vložený mimo
# pořadí umí skončit hláškou o poškozeném souboru — a ta se objeví až u čtenáře.
PPR_PORADI = (
    'pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr',
    'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
    'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
    'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap',
    'jc', 'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId',
    'cnfStyle', 'rPr', 'sectPr', 'pPrChange',
)
TCPR_PORADI = (
    'cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
    'textDirection', 'tcFitText', 'vAlign', 'hideMark', 'headers',
)
TBLPR_PORADI = (
    'tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
    'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd',
    'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription',
)
TRPR_PORADI = (
    'cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit',
    'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden',
)


# ── nízkoúrovňové pomůcky nad XML ────────────────────────────────────────────────────────

def _vloz_dle_poradi(rodic, dite, poradi):
    jmeno = dite.tag.split('}')[1]
    index = poradi.index(jmeno)
    for existujici in rodic:
        tag = existujici.tag.split('}')[1]
        if tag in poradi and poradi.index(tag) > index:
            existujici.addprevious(dite)
            return
    rodic.append(dite)


def _hrana(smer, styl, sz, barva, space=0):
    el = OxmlElement('w:' + smer)
    el.set(qn('w:val'), styl)
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), str(space))
    el.set(qn('w:color'), barva)
    return el


def spodni_linka(par, barva=AKCENT, sz=8, space=4):
    """Vodorovná linka pod odstavcem — tím se v dokumentu oddělují kapitoly."""
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    pBdr.append(_hrana('bottom', 'single', sz, barva, space))
    _vloz_dle_poradi(pPr, pBdr, PPR_PORADI)


def prostrkani(run, hodnota=40):
    """Rozestup písmen ve dvacetinách bodu; používá se jen pro nadtitulek na obálce."""
    rPr = run._r.get_or_add_rPr()
    el = OxmlElement('w:spacing')
    el.set(qn('w:val'), str(hodnota))
    rPr.append(el)


def vypln_bunky(bunka, barva):
    tcPr = bunka._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), barva)
    _vloz_dle_poradi(tcPr, shd, TCPR_PORADI)


def okraje_bunky(bunka, **hrany):
    """hrany: top/bottom/left/right = (styl, sz, barva) nebo None pro žádný okraj."""
    tcPr = bunka._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for smer in ('top', 'left', 'bottom', 'right'):
        if smer not in hrany:
            continue
        zadani = hrany[smer]
        if zadani is None:
            el = OxmlElement('w:' + smer)
            el.set(qn('w:val'), 'nil')
        else:
            styl, sz, barva = zadani
            el = _hrana(smer, styl, sz, barva)
        borders.append(el)
    _vloz_dle_poradi(tcPr, borders, TCPR_PORADI)


def okraje_tabulky(tab, vnitrni_vodorovny=(LINKA, 4)):
    tblPr = tab._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for smer in ('top', 'left', 'bottom', 'right', 'insideV'):
        el = OxmlElement('w:' + smer)
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    barva, sz = vnitrni_vodorovny
    borders.append(_hrana('insideH', 'single', sz, barva))
    _vloz_dle_poradi(tblPr, borders, TBLPR_PORADI)


def odsazeni_v_bunkach(tab, nahoře=70, dole=70, vlevo=100, vpravo=100):
    tblPr = tab._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for smer, hodnota in (('top', nahoře), ('left', vlevo), ('bottom', dole), ('right', vpravo)):
        el = OxmlElement('w:' + smer)
        el.set(qn('w:w'), str(hodnota))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    _vloz_dle_poradi(tblPr, mar, TBLPR_PORADI)


def opakovat_zahlavi(radek):
    trPr = radek._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    _vloz_dle_poradi(trPr, el, TRPR_PORADI)


def vyska_radku(radek, cm, pravidlo='atLeast'):
    trPr = radek._tr.get_or_add_trPr()
    el = OxmlElement('w:trHeight')
    el.set(qn('w:val'), str(int(cm * 567)))
    el.set(qn('w:hRule'), pravidlo)
    _vloz_dle_poradi(trPr, el, TRPR_PORADI)


def nedelit_radek(radek):
    trPr = radek._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    _vloz_dle_poradi(trPr, el, TRPR_PORADI)


def drz_s_dalsim(par):
    pPr = par._p.get_or_add_pPr()
    el = OxmlElement('w:keepNext')
    _vloz_dle_poradi(pPr, el, PPR_PORADI)


# ── text ─────────────────────────────────────────────────────────────────────────────────

ZNACKY = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')

# ==zvýrazněno== — žlutý podklad pro vyznačení změn k revizi (pokyn uživatele 9. 10. 2026).
# Uvnitř smí stát **tučně** i `neproporcionálně`; zvýraznění se nevnořuje.
ZVYRAZNENI = re.compile(r'(==.+?==)')


def psat(par, text, tucne=False, kurziva=False, velikost=None, barva=None, pismo=None):
    """Text s drobným značkováním: **tučně**, `neproporcionálně` a ==zvýrazněno==.

    Značkování je tu proto, aby se obsah kapitol dal psát jako běžné věty. Cokoli složitějšího
    (odkazy, poznámky) se skládá voláním po částech.
    """
    for usek in ZVYRAZNENI.split(text):
        if not usek:
            continue
        if usek.startswith('==') and usek.endswith('==') and len(usek) > 4:
            _psat_kusy(par, usek[2:-2], tucne, kurziva, velikost, barva, pismo, zvyraznit=True)
        else:
            _psat_kusy(par, usek, tucne, kurziva, velikost, barva, pismo, zvyraznit=False)
    return par


def _psat_kusy(par, text, tucne, kurziva, velikost, barva, pismo, zvyraznit):
    for kus in ZNACKY.split(text):
        if not kus:
            continue
        b, m, obsah = tucne, False, kus
        if kus.startswith('**') and kus.endswith('**'):
            b, obsah = True, kus[2:-2]
        elif kus.startswith('`') and kus.endswith('`'):
            m, obsah = True, kus[1:-1]
        run = par.add_run(obsah)
        # Jen zapnout, nikdy nevypínat: run.bold = False by zapsalo přímé formátování, které
        # přebije tučnost ze stylu — nadpisy by pak byly sázeny základním řezem.
        if b:
            run.bold = True
        if kurziva:
            run.italic = True
        if velikost:
            run.font.size = Pt(velikost)
        if barva:
            run.font.color.rgb = RGBColor.from_string(barva)
        if m:
            run.font.name = PISMO_MONO
            run.font.size = Pt((velikost or 10.5) - 1)
        elif pismo:
            run.font.name = pismo
        if zvyraznit:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def odstavec(doc, text='', styl=None, velikost=None, barva=None, kurziva=False,
             po=None, pred=None, zarovnani=None):
    par = doc.add_paragraph(style=styl)
    if text:
        psat(par, text, velikost=velikost, barva=barva, kurziva=kurziva)
    if po is not None:
        par.paragraph_format.space_after = Pt(po)
    if pred is not None:
        par.paragraph_format.space_before = Pt(pred)
    if zarovnani is not None:
        par.alignment = zarovnani
    return par


def nadpis1(doc, cislo, text):
    par = doc.add_paragraph(style='Heading 1')
    psat(par, (cislo + ' ' if cislo else '') + text)
    spodni_linka(par)
    return par


def nadpis2(doc, cislo, text):
    par = doc.add_paragraph(style='Heading 2')
    psat(par, (cislo + ' ' if cislo else '') + text)
    return par


def nadpis3(doc, text):
    par = doc.add_paragraph(style='Heading 3')
    psat(par, text)
    return par


def odrazky(doc, polozky):
    for text in polozky:
        par = doc.add_paragraph(style='List Bullet')
        psat(par, text)


def cislovane(doc, polozky):
    for text in polozky:
        par = doc.add_paragraph(style='List Number')
        psat(par, text)


# ── bloky ────────────────────────────────────────────────────────────────────────────────

def tabulka(doc, zahlavi, radky, sirky, klicovy_sloupec=True):
    """Tabulka bez mřížky: záhlaví s podtiskem, mezi řádky vlasová linka.

    Table Grid vypadá jako výpis z databáze; tady jde o dokument, který někdo čte.
    """
    tab = doc.add_table(rows=0, cols=len(zahlavi))
    tab.alignment = WD_TABLE_ALIGNMENT.LEFT
    tab.autofit = False
    okraje_tabulky(tab)
    odsazeni_v_bunkach(tab)

    r = tab.add_row()
    opakovat_zahlavi(r)
    for i, text in enumerate(zahlavi):
        bunka = r.cells[i]
        bunka.width = Cm(sirky[i])
        vypln_bunky(bunka, PLOCHA)
        okraje_bunky(bunka, bottom=('single', 8, AKCENT))
        par = bunka.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        psat(par, text, tucne=True, velikost=9.5, barva=AKCENT_TMAVY)

    for radek in radky:
        r = tab.add_row()
        nedelit_radek(r)
        for i, text in enumerate(radek):
            bunka = r.cells[i]
            bunka.width = Cm(sirky[i])
            bunka.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            par = bunka.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            psat(par, text, tucne=(klicovy_sloupec and i == 0), velikost=9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tab


def vysvetlivka(doc, popiska, text):
    """Odbočka k tématu: podtisk, svislá linka u levého okraje, popiska tučně."""
    tab = doc.add_table(rows=1, cols=1)
    tab.alignment = WD_TABLE_ALIGNMENT.LEFT
    tab.autofit = False
    okraje_tabulky(tab, vnitrni_vodorovny=(LINKA, 0))
    odsazeni_v_bunkach(tab, nahoře=120, dole=120, vlevo=160, vpravo=160)
    bunka = tab.rows[0].cells[0]
    bunka.width = Cm(SAZBA_CM)
    vypln_bunky(bunka, PLOCHA_SVET)
    okraje_bunky(bunka, left=('single', 18, AKCENT), top=None, bottom=None, right=None)
    par = bunka.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    psat(par, popiska + ' ', tucne=True, velikost=9.5, barva=AKCENT_TMAVY)
    psat(par, text, velikost=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tab


def citace(doc, radky):
    """Ukázka toho, co aplikace píše na obrazovku — sazbou odlišená od výkladu."""
    tab = doc.add_table(rows=1, cols=1)
    tab.alignment = WD_TABLE_ALIGNMENT.LEFT
    tab.autofit = False
    okraje_tabulky(tab, vnitrni_vodorovny=(LINKA, 0))
    odsazeni_v_bunkach(tab, nahoře=140, dole=140, vlevo=180, vpravo=180)
    bunka = tab.rows[0].cells[0]
    bunka.width = Cm(SAZBA_CM)
    okraje_bunky(bunka,
                 top=('single', 4, LINKA), bottom=('single', 4, LINKA),
                 left=('single', 4, LINKA), right=('single', 4, LINKA))
    prvni = True
    for popiska, text in radky:
        par = bunka.paragraphs[0] if prvni else bunka.add_paragraph()
        prvni = False
        par.paragraph_format.space_after = Pt(6)
        par.paragraph_format.space_before = Pt(0)
        if popiska:
            psat(par, popiska + ' ', tucne=True, velikost=9, barva=AKCENT_TMAVY)
        psat(par, text, velikost=9.5)
    bunka.paragraphs[-1].paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tab


_cislo_obrazku = [0]


def pocet_obrazku():
    """Kolik rámečků zatím vzniklo — do závěrečné hlášky generátoru."""
    return _cislo_obrazku[0]


def obrazek(doc, popis, navod, vyska=7.0):
    """Rámeček, do kterého se vloží snímek obrazovky, a popiska pod ním.

    Rámeček je tabulka o jedné buňce s přerušovaným okrajem. Obrázek se vkládá dovnitř
    (kliknout do rámečku, smazat šedý text, Vložení → Obrázky); rámeček pak drží obrázek na
    šířku sazby a popiska zůstane s ním. Kdo ho nechce, smaže celý řádek tabulky — popiska
    pod ním je obyčejný odstavec a zůstane.
    """
    _cislo_obrazku[0] += 1
    cislo = _cislo_obrazku[0]

    tab = doc.add_table(rows=1, cols=1)
    tab.alignment = WD_TABLE_ALIGNMENT.LEFT
    tab.autofit = False
    okraje_tabulky(tab, vnitrni_vodorovny=(LINKA, 0))
    odsazeni_v_bunkach(tab, nahoře=100, dole=100, vlevo=140, vpravo=140)
    radek = tab.rows[0]
    vyska_radku(radek, vyska)
    bunka = radek.cells[0]
    bunka.width = Cm(SAZBA_CM)
    bunka.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    vypln_bunky(bunka, RAMECEK)
    okraje_bunky(bunka,
                 top=('dashed', 8, LINKA), bottom=('dashed', 8, LINKA),
                 left=('dashed', 8, LINKA), right=('dashed', 8, LINKA))

    par = bunka.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)
    psat(par, 'sem patří obrázek ' + str(cislo), tucne=True, velikost=10, barva=INKOUST_SVET)
    par2 = bunka.add_paragraph()
    par2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par2.paragraph_format.space_before = Pt(2)
    par2.paragraph_format.space_after = Pt(0)
    psat(par2, navod, velikost=9, barva=INKOUST_SVET, kurziva=True)

    popiska = doc.add_paragraph(style='Caption')
    psat(popiska, 'Obrázek ' + str(cislo) + ' ', tucne=True, barva=AKCENT_TMAVY)
    psat(popiska, popis)
    return tab


# ── stránka, styly, zápatí ───────────────────────────────────────────────────────────────

def nastav_stranku(doc):
    for oddil in doc.sections:
        oddil.page_width = Cm(21)
        oddil.page_height = Cm(29.7)
        oddil.left_margin = Cm(2.2)
        oddil.right_margin = Cm(2.2)
        oddil.top_margin = Cm(2.2)
        oddil.bottom_margin = Cm(2.0)
        oddil.footer_distance = Cm(1.2)


def _pismo(styl, jmeno):
    """Písmo do stylu — a pryč s odkazy na motiv.

    Vestavěné nadpisové styly jmenují písmo přes motiv (`w:asciiTheme`). Ten má ve Wordu
    přednost před přímo zapsaným jménem, takže samotné nastavení `font.name` se neprojeví
    a nadpisy zůstanou v motivovém Calibri Light, zatímco text okolo je Segoe UI.
    """
    styl.font.name = jmeno
    rPr = styl.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for atribut in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rFonts.get(qn(atribut)) is not None:
            del rFonts.attrib[qn(atribut)]
    for atribut in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(atribut), jmeno)


def nastav_styly(doc):
    normal = doc.styles['Normal']
    _pismo(normal, PISMO)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INKOUST)
    pf = normal.paragraph_format
    pf.space_after = Pt(7)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15

    h1 = doc.styles['Heading 1']
    _pismo(h1, PISMO)
    h1.font.size = Pt(17)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(AKCENT_TMAVY)
    h1.paragraph_format.space_before = Pt(22)
    h1.paragraph_format.space_after = Pt(9)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles['Heading 2']
    _pismo(h2, PISMO)
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(INKOUST)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles['Heading 3']
    _pismo(h3, PISMO)
    h3.font.size = Pt(10.5)
    h3.font.bold = True
    h3.font.italic = False
    h3.font.color.rgb = RGBColor.from_string(AKCENT_TMAVY)
    h3.paragraph_format.space_before = Pt(11)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True

    popiska = doc.styles['Caption']
    _pismo(popiska, PISMO)
    popiska.font.size = Pt(8.5)
    popiska.font.italic = False
    popiska.font.bold = False
    popiska.font.color.rgb = RGBColor.from_string(INKOUST_SVET)
    popiska.paragraph_format.space_before = Pt(3)
    popiska.paragraph_format.space_after = Pt(12)

    for jmeno in ('List Bullet', 'List Number'):
        styl = doc.styles[jmeno]
        _pismo(styl, PISMO)
        styl.font.size = Pt(10.5)
        styl.font.color.rgb = RGBColor.from_string(INKOUST)
        styl.paragraph_format.space_after = Pt(4)
        styl.paragraph_format.line_spacing = 1.15


def nastav_cestinu(doc):
    """Bez tohohle kontroluje Word český text anglickým slovníkem a podtrhá skoro všechno."""
    rPr = doc.styles.element.find(qn('w:docDefaults'))
    if rPr is not None:
        default = rPr.find(qn('w:rPrDefault'))
        if default is not None:
            vnitrni = default.find(qn('w:rPr'))
            if vnitrni is None:
                vnitrni = OxmlElement('w:rPr')
                default.append(vnitrni)
            lang = vnitrni.find(qn('w:lang'))
            if lang is None:
                lang = OxmlElement('w:lang')
                vnitrni.append(lang)
            lang.set(qn('w:val'), 'cs-CZ')
    for jmeno in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption'):
        styl = doc.styles[jmeno]
        el = styl.element.get_or_add_rPr()
        lang = el.find(qn('w:lang'))
        if lang is None:
            lang = OxmlElement('w:lang')
            el.append(lang)
        lang.set(qn('w:val'), 'cs-CZ')


def pole(par, kod):
    zacatek = OxmlElement('w:fldChar')
    zacatek.set(qn('w:fldCharType'), 'begin')
    instrukce = OxmlElement('w:instrText')
    instrukce.set(qn('xml:space'), 'preserve')
    instrukce.text = kod
    konec = OxmlElement('w:fldChar')
    konec.set(qn('w:fldCharType'), 'end')
    run = par.add_run()
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(INKOUST_SVET)
    run._r.append(zacatek)
    run._r.append(instrukce)
    run._r.append(konec)


def nastav_zapati(doc, titulek):
    """Zápatí: jméno dokumentu vlevo, číslo stránky vpravo."""
    zapati = doc.sections[0].footer
    # Zarážky nese styl Footer (uprostřed a vpravo) a odstavcové přidání je nepřebije —
    # tabulátor by skočil na tu prostřední a číslo stránky by stálo uprostřed stránky.
    # Přepsat se proto musí rovnou styl.
    zarazky = doc.styles['Footer'].paragraph_format.tab_stops
    zarazky.clear_all()
    zarazky.add_tab_stop(Cm(SAZBA_CM), WD_TAB_ALIGNMENT.RIGHT)
    par = zapati.paragraphs[0]
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(0)
    psat(par, titulek, velikost=8.5, barva=INKOUST_SVET)
    par.add_run('\t')
    pole(par, ' PAGE ')


def zalom(doc):
    doc.add_page_break()


def zaklad(doc, titulek_zapati):
    """Vše, co musí zaznít nad prázdným dokumentem, než se začne psát obsah."""
    nastav_stranku(doc)
    nastav_styly(doc)
    nastav_cestinu(doc)
    nastav_zapati(doc, titulek_zapati)
