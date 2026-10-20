# -*- coding: utf-8 -*-
"""Udělá ze `schopnosti.html` dokument .docx, který jde otevřít a přepsat v MS Wordu.

Proč generátor a ne jednorázový převod: `schopnosti.html` se udržuje ručně a bude se měnit.
Bez tohohle skriptu by wordová podoba po první změně zastarala a nikdo by nevěděl, čím ji
obnovit.

Převádí se **struktura, ne vzhled**. Text se sází wordovými zabudovanými styly (Title,
Heading 1/2, List Bullet, Table Grid), takže dokument vypadá jako běžný wordový soubor a dá se
přeformátovat obvyklou cestou — změnou stylu, ne přepisováním odstavců. Barvy a písma
z HTML se nepřenášejí schválně: papírová a obrazovková podoba mají jinou práci a slepenec
z obojího by se v obou hodil špatně.

    python docs/gen-schopnosti-docx.py

⚠ Přepíše `docs/AnEP-schopnosti.docx`. Kdo do wordové podoby psal vlastní úpravy, přijde o ně —
zdrojem zůstává HTML.

Potřebuje `python-docx` a `lxml`.
"""

import os
import re
import sys

import lxml.html
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ZDE = os.path.dirname(os.path.abspath(__file__))
ZDROJ = os.path.join(ZDE, 'schopnosti.html')
CIL = os.path.join(ZDE, 'AnEP-schopnosti.docx')

MONO = 'Consolas'
ODKAZ = RGBColor(0x00, 0x31, 0x89)


# ── drobnosti ────────────────────────────────────────────────────────────────────────────

def mezery(text):
    """HTML mezery se slévají do jedné; odstavec zalomený ve zdroji je pořád jeden odstavec."""
    return re.sub(r'\s+', ' ', text)


def trida(el):
    return el.get('class') or ''


def pridej_run(par, text, bold=False, italic=False, mono=False, zvyraz=False):
    if not text:
        return
    # Mezera na začátku odstavce je jen zbytek odsazení ve zdroji, ne součást věty.
    if not par.runs and text.strip() == '':
        return
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    if mono:
        run.font.name = MONO
        run.font.size = Pt(10)
    # <mark> ve zdroji — vyznačené změny k revizi (pokyn uživatele 9. 10. 2026).
    if zvyraz:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def pridej_odkaz(par, url, text, mono=False):
    """Word nemá v python-docx své API pro hypertextový odkaz, musí se složit v XML."""
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    odkaz = OxmlElement('w:hyperlink')
    odkaz.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    barva = OxmlElement('w:color')
    barva.set(qn('w:val'), '003189')
    rPr.append(barva)
    podtrzeni = OxmlElement('w:u')
    podtrzeni.set(qn('w:val'), 'single')
    rPr.append(podtrzeni)
    if mono:
        pismo = OxmlElement('w:rFonts')
        pismo.set(qn('w:ascii'), MONO)
        pismo.set(qn('w:hAnsi'), MONO)
        rPr.append(pismo)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    odkaz.append(run)
    par._p.append(odkaz)


def text_z(el):
    return mezery(''.join(el.itertext())).strip()


# ── vnitřek odstavce ─────────────────────────────────────────────────────────────────────

def vloz_obsah(par, el, bold=False, italic=False, mono=False, zvyraz=False):
    if el.text:
        pridej_run(par, mezery(el.text), bold, italic, mono, zvyraz)
    for dite in el:
        if not isinstance(dite.tag, str):          # komentář
            if dite.tail:
                pridej_run(par, mezery(dite.tail), bold, italic, mono, zvyraz)
            continue
        tag = dite.tag.lower()
        b = bold or tag in ('strong', 'b')
        i = italic or tag in ('em', 'i')
        m = mono or tag in ('code', 'kbd', 'samp') or 'mono' in trida(dite)
        z = zvyraz or tag == 'mark'
        if tag == 'br':
            par.add_run().add_break()
        elif tag == 'a' and dite.get('href') and not dite.get('href').startswith('#'):
            pridej_odkaz(par, dite.get('href'), text_z(dite), m)
        else:
            vloz_obsah(par, dite, b, i, m, z)
        if dite.tail:
            pridej_run(par, mezery(dite.tail), bold, italic, mono, zvyraz)


def odstavec(doc, el, styl='Normal'):
    par = doc.add_paragraph(style=styl)
    vloz_obsah(par, el)
    return par


# ── bloky ────────────────────────────────────────────────────────────────────────────────

def tabulka(doc, el):
    radky = el.findall('.//tr')
    if not radky:
        return
    sirka = max(len(r.findall('./th')) + len(r.findall('./td')) for r in radky)
    tab = doc.add_table(rows=0, cols=sirka)
    tab.style = 'Table Grid'
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for poradi, r in enumerate(radky):
        bunky = [c for c in r if c.tag in ('th', 'td')]
        cilove = tab.add_row().cells
        zahlavi = all(c.tag == 'th' for c in bunky)
        for j, bunka in enumerate(bunky):
            if j >= sirka:
                break
            par = cilove[j].paragraphs[0]
            vloz_obsah(par, bunka)
            if zahlavi or 'key' in trida(bunka):
                for run in par.runs:
                    run.bold = True
        if zahlavi:
            # Aby se u tabulky přes dvě stránky opakovalo záhlaví — Word to jinak neudělá.
            trPr = cilove[0]._tc.getparent().get_or_add_trPr()
            opakuj = OxmlElement('w:tblHeader')
            opakuj.set(qn('w:val'), 'true')
            trPr.append(opakuj)
    doc.add_paragraph()


def poznamka(doc, el):
    """Vysvětlivka stranou: popiska tučně, text za ní, celé stylem Quote."""
    par = doc.add_paragraph(style='Quote')
    popiska = el.find('./span[@class="label"]')
    if popiska is not None:
        run = par.add_run(text_z(popiska) + ' — ')
        run.bold = True
    for p in el.findall('./p'):
        vloz_obsah(par, p)


def seznam(doc, el, cislovany=False):
    styl = 'List Number' if cislovany else 'List Bullet'
    for li in el.findall('./li'):
        par = doc.add_paragraph(style=styl)
        vloz_obsah(par, li)


def pruh(doc, el):
    """Řádka faktů z hlavičky — ve Wordu z ní je jeden odstavec s oddělovači."""
    kusy = [text_z(s) for s in el.findall('./span')]
    if not kusy:
        return
    par = doc.add_paragraph()
    par.add_run(' · '.join(kusy)).italic = True


def obsah_dokumentu(doc, nav):
    doc.add_paragraph('Obsah', style='Heading 1')
    for li in nav.findall('.//li'):
        doc.add_paragraph(text_z(li), style='List Number')
    doc.add_page_break()


def zpracuj(doc, el):
    tag = el.tag.lower() if isinstance(el.tag, str) else ''
    cls = trida(el)

    if tag in ('script', 'style'):
        return
    if tag == 'nav':
        obsah_dokumentu(doc, el)
        return
    if tag == 'table':
        tabulka(doc, el)
        return
    if tag == 'ul':
        seznam(doc, el)
        return
    if tag == 'ol':
        seznam(doc, el, cislovany=True)
        return
    if tag == 'h1':
        odstavec(doc, el, 'Title')
        return
    if tag == 'h2':
        odstavec(doc, el, 'Heading 1')
        return
    if tag == 'h3':
        odstavec(doc, el, 'Heading 2')
        return
    if tag == 'h4':
        odstavec(doc, el, 'Heading 3')
        return
    if tag == 'p':
        if 'eyebrow' in cls:
            par = odstavec(doc, el)
            for run in par.runs:
                run.italic = True
        elif 'subtitle' in cls:
            odstavec(doc, el, 'Subtitle')
        else:
            odstavec(doc, el)
        return
    if 'aside' in cls:
        poznamka(doc, el)
        return
    if 'factbar' in cls:
        pruh(doc, el)
        return
    if tag == 'footer':
        doc.add_paragraph()
        for p in el.findall('./p'):
            par = doc.add_paragraph()
            vloz_obsah(par, p)
            for run in par.runs:
                run.italic = True
                run.font.size = Pt(9)
        return

    for dite in el:
        zpracuj(doc, dite)


# ── stránka a jazyk ──────────────────────────────────────────────────────────────────────

def nastav_stranku(doc):
    for oddil in doc.sections:
        oddil.page_width = Cm(21)
        oddil.page_height = Cm(29.7)
        oddil.left_margin = Cm(2.5)
        oddil.right_margin = Cm(2.5)
        oddil.top_margin = Cm(2.5)
        oddil.bottom_margin = Cm(2.5)


def nastav_cestinu(doc):
    """Bez tohohle Word kontroluje český text anglickým slovníkem a podtrhá skoro všechno."""
    for jmeno in ('Normal',):
        styl = doc.styles[jmeno]
        rPr = styl.element.get_or_add_rPr()
        lang = rPr.find(qn('w:lang'))
        if lang is None:
            lang = OxmlElement('w:lang')
            rPr.append(lang)
        lang.set(qn('w:val'), 'cs-CZ')


def main():
    if not os.path.exists(ZDROJ):
        print('nenalezeno: ' + ZDROJ)
        return 1

    with open(ZDROJ, 'rb') as f:
        strom = lxml.html.fromstring(f.read())

    doc = Document()
    nastav_stranku(doc)
    nastav_cestinu(doc)

    telo = strom.find('.//body')
    for dite in telo:
        zpracuj(doc, dite)

    doc.core_properties.title = 'AnEP — vlastnosti a schopnosti'
    doc.core_properties.language = 'cs-CZ'
    doc.core_properties.comments = ('Vygenerováno z docs/schopnosti.html skriptem '
                                    'docs/gen-schopnosti-docx.py')

    doc.save(CIL)
    print('hotovo: ' + CIL)
    return 0


if __name__ == '__main__':
    sys.exit(main())
