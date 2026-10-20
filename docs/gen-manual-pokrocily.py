# -*- coding: utf-8 -*-
"""Vygeneruje `docs/AnEP-manual-pokrocily.docx` — uživatelský manuál pro pokročilé uživatele.

Tři texty o aplikaci a každý odpovídá na jinou otázku:

    prirucka.html                 rychlý start: vložit soubor a přečíst verdikt
    AnEP-schopnosti-s-obrazky.docx  co aplikace umí
    AnEP-manual-pokrocily.docx      jak to použít a o co se závěr opírá   ← tenhle

Do 16. 9. 2026 wordová podoba manuálu generátor neměla a upravovala se ručně. To přestalo
stačit v okamžiku, kdy se s aplikací rozešla: text psaný do dokumentu se nedá porovnat
s kódem, ze kterého mluví. Sazbu drží `docx_kit.py`, sdílený s `gen-schopnosti-s-obrazky.py`;
tady je jen obsah.

**Nemluví o situacích**: ty jsou vnitřní věc aplikace a uživatel je nikde nepotřebuje. Číslované
jsou pro čtenáře jen důvody.

    python docs/gen-manual-pokrocily.py

⚠ Přepíše `docs/AnEP-manual-pokrocily.docx` včetně vložených obrázků a ručních úprav. Kdo do
wordové podoby vkládal snímky, uloží si ji po vložení pod jiným jménem.

Potřebuje `python-docx`.
"""

import os
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor

from docx_kit import (AKCENT, AKCENT_TMAVY, INKOUST, INKOUST_SVET, PISMO_LEHKE,
                      cislovane, nadpis1, nadpis2, nadpis3, obrazek, odrazky, odstavec,
                      pocet_obrazku, prostrkani, psat, spodni_linka, tabulka, vysvetlivka,
                      zaklad, zalom)

ZDE = os.path.dirname(os.path.abspath(__file__))
CIL = os.path.join(ZDE, 'AnEP-manual-pokrocily.docx')

STAV_KE_DNI = '16. 9. 2026'
TITULEK = 'AnEP — uživatelský manuál pro pokročilé uživatele'


# ── obálka a obsah ───────────────────────────────────────────────────────────────────────

def obalka(doc):
    mezera = doc.add_paragraph()
    mezera.paragraph_format.space_after = Pt(30)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    run = par.add_run('UŽIVATELSKÝ MANUÁL PRO POKROČILÉ UŽIVATELE')
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(AKCENT)
    prostrkani(run, 60)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run('AnEP')
    run.font.name = PISMO_LEHKE
    run.font.size = Pt(48)
    run.font.color.rgb = RGBColor.from_string(INKOUST)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(10)
    run = par.add_run('Analyzátor elektronických podpisů')
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(INKOUST_SVET)
    spodni_linka(par, AKCENT, sz=12, space=8)

    odstavec(doc, 'Plný rozsah schopností aplikace — a to, oč se její závěry opírají.',
             velikost=12, po=14)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(24)
    psat(par, 'Windows · desktopová aplikace  ·  jádro ověřování DSS (Evropská komise)'
              '  ·  rozhraní česky  ·  licence EUPL-1.2',
         velikost=9, barva=INKOUST_SVET)

    obrazek(doc, 'Hlavní okno aplikace po ověření dokumentu.',
            'celé okno aplikace, záložka Zjednodušený přehled', vyska=9.5)

    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(10)
    psat(par, 'Stav ke dni ' + STAV_KE_DNI + '. Výstupy aplikace jsou určeny pro výuku '
              'a seznámení s problematikou a jsou poskytovány bez záruky; aplikace nevydává '
              'právní stanoviska.', velikost=8.5, barva=INKOUST_SVET)
    zalom(doc)


OSNOVA = [
    ('1', 'O tomto manuálu', []),
    ('2', 'Co AnEP je a co není', ['Tři možné výsledky', 'Kde jsou hranice']),
    ('3', 'Instalace, první spuštění a úložiště',
     ['Distribuce', 'První spuštění', 'Kde aplikace ukládá svá data']),
    ('4', 'Co aplikace přečte',
     ['Nosiče a formáty', 'Odloučené podpisy a evidenční záznamy', 'Heslem chráněná PDF',
      'Osiřelé podpisy', 'Soubor, který není podepsaným dokumentem']),
    ('5', 'Hlavní okno', ['Plocha, tlačítka a stavový řádek', 'Odpojení záložky',
                          'Hlavička o důvěryhodných seznamech']),
    ('6', 'Šest záložek', ['Zjednodušený přehled', 'Rozšířený přehled', 'Analýza',
                           'Digitální kontinuita', 'Detailní rozbor', 'Protokoly DSS']),
    ('7', 'Důvody a doporučení',
     ['Stavba důvodu', 'Proč se důvody potlačují', 'Náhradní znění', 'Číselník']),
    ('8', 'Vlastní prohlížeče',
     ['K čemu jsou', 'Kontejnery', 'Certifikáty', 'Seznamy CRL', 'Odpovědi OCSP',
      'Přílohy dokumentu', 'Inspektor podpisové struktury', 'Kudy se k nim dostat']),
    ('9', 'Důvěra a revokace',
     ['Odkud pocházejí kotvy důvěry', 'Revokační údaje a jejich původ', 'Tři skupiny údajů',
      'Vlastní zásoby', 'Obnova důvěryhodných seznamů']),
    ('10', 'Přísnost a ústupky',
     ['Čtyři ústupky', 'Dvě samostatná nastavení', 'Jednorázová validační politika',
      'Rozhodující ústupek']),
    ('11', 'Úkony nad dokumentem',
     ['Zachovat ověřitelnost', 'Souběh více podpisů', 'Forenzní balíček',
      'Jména výstupních souborů', 'Kdy se úkon nenabídne']),
    ('12', 'Digitální kontinuita',
     ['Tři čtenáři', 'Struktura záložky', 'Co určuje datum', 'Když lhůtu stanovit nelze']),
    ('13', 'Zdroje časových razítek', ['Sonda při zadání', 'Když vybraný zdroj nestačí']),
    ('14', 'Protokoly a export', ['Protokol o ověření', 'Dvě volby', 'Zprávy DSS']),
    ('15', 'Zvláštní obsahy',
     ['Datové zprávy ISDS', 'Evidenční záznamy', 'Odkaz z eLegalizace', 'Podpisy JAdES',
      'Vizualizace a původ podpisu', 'Zámek dokumentu', 'Prázdná podpisová pole',
      'Protipodpis', 'Podpisy v OOXML']),
    ('16', 'Nastavení — přehled', ['Základní nastavení', 'Pokročilé nastavení', 'O aplikaci']),
    ('17', 'Řešení potíží', []),
    ('18', 'Slovníček', []),
]


def obsah(doc):
    par = doc.add_paragraph(style='Heading 1')
    psat(par, 'Obsah')
    spodni_linka(par)
    for cislo, titul, podkapitoly in OSNOVA:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        psat(p, cislo + ' ', tucne=True, barva=AKCENT_TMAVY)
        psat(p, titul, tucne=True)
        if podkapitoly:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(7)
            p2.paragraph_format.left_indent = Cm(0.9)
            psat(p2, ' · '.join(podkapitoly), velikost=9, barva=INKOUST_SVET)
    zalom(doc)


# ── kapitoly ─────────────────────────────────────────────────────────────────────────────

def kapitola_1(doc):
    nadpis1(doc, '1', 'O tomto manuálu')
    odstavec(doc, 'Tento manuál popisuje AnEP v plném rozsahu: nejen jak dokument ověřit, ale '
                  'i co aplikace o výsledku tvrdí, o co se to tvrzení opírá a kde jsou jeho '
                  'meze. Je psán pro čtenáře, který se v problematice elektronických podpisů '
                  'orientuje — zná pojmy jako certifikát, CRL, OCSP, časové razítko, '
                  'důvěryhodný seznam — a chce z aplikace dostat všechno, co umí.')
    odstavec(doc, 'Kdo hledá jen „vložit soubor a přečíst verdikt“, najde to v samostatném '
                  'rychlém startu (`docs/prirucka.html`); tento manuál ho nenahrazuje, navazuje '
                  'na něj. Přehled schopností bez návodů je v dokumentu '
                  '„AnEP — přehled schopností“.')

    nadpis2(doc, '', 'Jak manuál číst')
    odrazky(doc, [
        'Kapitoly **2 až 6** popisují, co aplikace přečte a co ukáže.',
        'Kapitoly **7 až 10** vysvětlují, jak vzniká verdikt — důvody a doporučení, důvěra '
        'a revokace, přísnost a ústupky. To je jádro pro pokročilého uživatele.',
        'Kapitola **8** patří vlastním prohlížečům: nástrojům, které odpovídají na otázku '
        '„co v tom souboru je“, nikoli „platí to“.',
        'Kapitoly **11 až 13** popisují úkony nad dokumentem a otázku, do kdy zůstane '
        'ověřitelný.',
        'Kapitoly **14 až 18** pokrývají protokoly, zvláštní obsahy, nastavení, potíže a pojmy.',
    ])
    vysvetlivka(doc, 'Pozor',
                'Výstupy aplikace jsou určeny pro výuku a seznámení s problematikou a jsou '
                'poskytovány bez záruky. Aplikace nevydává právní stanoviska.')


def kapitola_2(doc):
    nadpis1(doc, '2', 'Co AnEP je a co není')
    odstavec(doc, 'AnEP ověřuje elektronické podpisy, pečeti a časová razítka podle nařízení '
                  'eIDAS (910/2014) a norem ETSI, především ETSI EN 319 102-1 a navazujících. '
                  'Vlastní ověřování obstarává otevřená knihovna DSS Evropské komise, tedy '
                  'referenční implementace uvedené normy. Aplikace nad ní staví to, co DSS '
                  'záměrně nedělá: mění strohé `INDETERMINATE` nebo `REVOKED_NO_POE` ve větu, '
                  'která říká, co se stalo, proč, a co se s tím dá dělat.')

    nadpis2(doc, '2.1', 'Tři možné výsledky')
    tabulka(doc,
            ['Výsledek', 'Co znamená'],
            [['platný',
              'Platnost byla prokázána za podmínek, které aplikace vypíše.'],
             ['stav platnosti neznámý',
              'Není to závěr o neplatnosti. Znamená to, že se platnost prokázat nepodařilo — '
              'a aplikace vždy uvede, co konkrétně chybělo a jestli to jde ještě opatřit.'],
             ['neplatný',
              'Neplatnost je prokázána: porušená integrita podepsaného obsahu, podpisová '
              'hodnota, kterou nelze ověřit veřejným klíčem z certifikátu, nebo revokace '
              'prokázaná k rozhodnému okamžiku.']],
            [4.6, 12.0])
    odstavec(doc, 'Prostřední případ je ten podstatný a nejčastější u starších dokumentů. '
                  'Většina obsahu aplikace se týká právě jeho: proč se platnost prokázat '
                  'nepodařilo a co s tím.')

    nadpis2(doc, '2.2', 'Kde jsou hranice')
    odstavec(doc, 'Co aplikace nedělá, a je to záměr:')
    odrazky(doc, [
        '**Nepodepisuje.** Umí přidat časové razítko a doplnit ověřovací materiál, ale podpis '
        'ani pečeť nevytváří.',
        '**Nevydává právní stanoviska.** Právní účinky uvádí jen tam, kde plynou přímo '
        'z výsledku ověření a z textu předpisu.',
        '**Nepředvídá cizí validátory s jistotou.** O jiných nástrojích mluví v možnosti '
        '(„mohly by“), protože nezná jejich verzi, nastavení ani obsah jejich zásob '
        'revokačních informací.',
        '**Nevytváří evidenční záznamy.** Ověřit je umí, vydat ne.',
        '**Neověřuje podpisy v kancelářských formátech OOXML.** Nahlédne do nich a řekne, co '
        'o sobě tvrdí; verdikt nevydá (kapitola 15.9).',
        '**Nezakrývá ústupky.** Kladný verdikt, který stojí na povoleném ústupku, to má u sebe '
        'napsáno.',
    ])


def kapitola_3(doc):
    nadpis1(doc, '3', 'Instalace, první spuštění a úložiště')

    nadpis2(doc, '3.1', 'Distribuce')
    odstavec(doc, 'Aplikace se dodává ve dvou podobách a obě obsahují vlastní běhové prostředí '
                  '— není třeba nic doinstalovávat:')
    odrazky(doc, [
        'instalátor `.msi` — zapíše aplikaci do Programů a vytvoří zástupce;',
        'přenosná verze — rozbalí se do libovolné složky a spouští se odtud.',
    ])

    nadpis2(doc, '3.2', 'První spuštění')
    odstavec(doc, 'Při prvním spuštění si aplikace stáhne evropské důvěryhodné seznamy (LOTL '
                  'a národní seznamy) a uloží je do místní mezipaměti. Trvá to desítky sekund '
                  'a děje se to jen jednou; při dalších startech se použije mezipaměť '
                  'a čerstvost se dokontroluje na pozadí, takže start nečeká na síť.')
    odstavec(doc, 'Hned po objevení okna proběhne na pozadí jedno zahazovací ověření nad '
                  'přibaleným vzorovým podpisem. Probouzí ověřovací cestu, takže první skutečné '
                  'ověření netrvá znatelně déle než druhé; nesahá přitom na síť ani na data '
                  'uživatele.')
    obrazek(doc, 'Hlavní okno aplikace hned po spuštění, prázdné.',
            'okno po spuštění, s vyznačenou plochou pro vložení souboru a řadou tlačítek',
            vyska=7.0)

    nadpis2(doc, '3.3', 'Kde aplikace ukládá svá data')
    odstavec(doc, 'Vše uživatelské leží v profilu, ve složce `%APPDATA%\\SignatureAnalyzer`:')
    tabulka(doc,
            ['Obsah', 'K čemu'],
            [['nastavení', 'volby z dialogu Nastavení, profily zdrojů časových razítek'],
             ['vlastní úložiště CRL', 'seznamy CRL, které si uživatel nashromáždil'],
             ['vlastní důvěryhodné certifikáty',
              'kotvy důvěry, které evropské seznamy neuvádějí'],
             ['mezipaměť důvěryhodných seznamů', 'stav EUTL, aby start nečekal na síť']],
            [5.4, 11.2])
    vysvetlivka(doc, 'Pozor',
                'Vlastní úložiště CRL a vlastní kotvy důvěry se po ověření nevracejí do '
                'výchozího stavu — na rozdíl od ústupků platí trvale, dokud je uživatel '
                'neodebere. Závěr, který na nich stojí, tedy jiný nástroj nemusí zopakovat; '
                'aplikace to u takového závěru vždy přizná.')


def kapitola_4(doc):
    nadpis1(doc, '4', 'Co aplikace přečte')
    odstavec(doc, 'Formát se pozná podle obsahu souboru, ne podle přípony.')
    tabulka(doc,
            ['Nosič', 'Co v něm aplikace ověří'],
            [['PDF',
              'PAdES všech úrovní (B-B až B-LTA), včetně podpisů PDF-NOT-ETSI z doby před '
              'eIDAS, podpisů PKCS #7 a dokumentových časových razítek'],
             ['ASiC-S / ASiC-E',
              '`.asice`, `.asics`, `.sce`, `.scs` i kontejner ASiC uložený s příponou `.zip`'],
             ['Archiv ZIP',
              'obyčejný archiv, který není kontejnerem ASiC, se otevře v prohlížeči obsahu; '
              'soubory v něm lze předat k ověření jednotlivě'],
             ['Odloučené a samostatné',
              'CAdES (`.p7s`, `.p7m`, `.cms`), XAdES (`.xml`), JAdES (`.json`, `.jose`, `.jws`)'],
             ['Datové zprávy ISDS',
              '`.zfo` — zprávy i doručenky, ověřené přes kvalifikované časové razítko nad '
              'datovou zprávou'],
             ['ODF',
              'ODT, ODS, ODP — podpisy XAdES v `META-INF/documentsignatures.xml`'],
             ['Kancelářské OOXML',
              'DOCX, XLSX, PPTX — podpis se plnohodnotně neověřuje; aplikace nahlédne dovnitř '
              'a ukáže, co o sobě podpis tvrdí, a odděleně to, co se ověřit dá'],
             ['Evidenční záznamy',
              'RFC 4998 (`.ers`, `.er`) a XMLERS, samostatné i vložené']],
            [4.2, 12.4])
    odstavec(doc, 'Bajty, které musely ustoupit, aby šel soubor vůbec přečíst — smetí před '
                  'hlavičkou PDF —, se spočítají a řeknou v Průběhu ověřování. Nic je '
                  'nepodepsalo, a tak to čtenář má vědět.')

    nadpis2(doc, '4.1', 'Odloučené podpisy a evidenční záznamy')
    odstavec(doc, 'K odloučenému podpisu se přikládá podepsaný obsah jako druhý soubor. Hlavní '
                  'okno má proto vedle plochy pro dokument ještě tlačítko pro externí obsah. '
                  'Bez podepsaného obsahu nelze odloučený podpis ověřit — otisk není s čím '
                  'porovnat, a aplikace to řekne jako důvod, ne jako chybu čtení.')
    odstavec(doc, 'Týmž tlačítkem se přikládá evidenční záznam k dokumentu. Není to druhý '
                  'podpis k ověření, nýbrž materiál dokládající, že dokument už existoval; '
                  'na pořadí obou souborů proto nezáleží.')

    nadpis2(doc, '4.2', 'Heslem chráněná PDF')
    odstavec(doc, 'PDF chráněné heslem se otevře po jeho zadání a ověří se stejně jako ostatní. '
                  'Heslo se nikam neukládá. Aplikace si o ně řekne až tehdy, když bez něj '
                  'dokument opravdu neotevře: u souboru, který nese jen vlastnické heslo '
                  '(omezuje tisk či úpravy, čtení ne), se neptá vůbec.')

    nadpis2(doc, '4.3', 'Osiřelé podpisy')
    odstavec(doc, 'Podpis, který knihovna DSS neumí zpracovat, se v přehledu neztratí — ohlásí '
                  'se jako nerozpoznaný druh prvku. To je důležité: dokument se dvěma podpisy, '
                  'z nichž jeden je nečitelný, nesmí vypadat jako dokument s jedním podpisem.')

    nadpis2(doc, '4.4', 'Soubor, který není podepsaným dokumentem')
    odstavec(doc, 'Holý certifikát, seznam CRL nebo odpověď OCSP se nečte ověřováním, nýbrž '
                  'očima. Vložíte-li takový soubor do aplikace, neproběhne ověření a rovnou se '
                  'otevře příslušný prohlížeč (kapitola 8). Poznají se podle obsahu, ne podle '
                  'přípony — `.cer` bývá i PEM a prvek vytažený z kontejneru nemá příponu '
                  'žádnou —, takže ověřitelný dokument tím o ověření nepřijde.')


def kapitola_5(doc):
    nadpis1(doc, '5', 'Hlavní okno')

    nadpis2(doc, '5.1', 'Plocha, tlačítka a stavový řádek')
    odstavec(doc, 'Okno má tři vodorovné části: plochu pro vstup s tlačítky, stavový řádek '
                  'a záložky s výsledkem.')
    tabulka(doc,
            ['Prvek', 'Funkce'],
            [['Plocha pro dokument',
              'Dokument nebo soubor s podepisovaným obsahem. Přetáhněte sem soubor, nebo '
              'klepněte a vyberte.'],
             ['↺ v rohu plochy', 'Obnovit naposledy zadanou dvojici souborů.'],
             ['✓', 'Spustit ověření.'],
             ['⚙', 'Nastavení.'],
             ['Tlačítko externího obsahu',
              'Odloučený podpis nebo evidenční záznam. Prázdné je tiché, s načteným souborem '
              'se rozsvítí; zrušení výběru nad obsazeným tlačítkem vloženou položku odebere.'],
             ['↗ ve stavovém řádku',
              'Otevřít dokument v přidružené aplikaci; u kontejnerů se rozdvojuje na „otevřít '
              'dokument“ a „otevřít jako kontejner“.'],
             ['Sponka ve stavovém řádku',
              'Zobrazit přílohy dokumentu; číslo v závorce udává, kolik jich je.']],
            [4.6, 12.0])
    obrazek(doc, 'Stavový řádek po ověření.',
            'jméno souboru, velikost, čas podpisu, počet nalezených prvků a vpravo ikony pro '
            'otevření dokumentu a pro přílohy', vyska=4.0)

    nadpis2(doc, '5.2', 'Odpojení záložky')
    odstavec(doc, 'Kteroukoli záložku lze odpojit do samostatného okna (pravé tlačítko na jejím '
                  'oušku). Hodí se, když je potřeba mít vedle sebe například Analýzu '
                  'a Detailní rozbor.')

    nadpis2(doc, '5.3', 'Hlavička o důvěryhodných seznamech')
    odstavec(doc, 'Nad kartami v Rozšířeném přehledu a v Analýze stojí řádek o stavu evropských '
                  'důvěryhodných seznamů — v obou týž, protože je to výrok o celém ověření, ne '
                  'o jedné záložce. Řekne, kdy aplikace své kopie seznamů naposledy ověřila '
                  'u zdroje, jestli obnova zrovna neběží, které národní seznamy se příliš dlouho '
                  'nepodařilo potvrdit u zdroje, a které země jsou v uloženém stavu sice vedeny, '
                  'ale nepřišel z nich jediný certifikát — typicky proto, že se jejich seznam '
                  'nepodařilo stáhnout ani ověřit.')
    odstavec(doc, 'Čas v hlavičce je okamžik porovnání se zdrojem, ne datum vydání seznamů. '
                  'Datum vydání ukáže klepnutí na ikonu ⓘ za řádkem: naposledy vydaný národní '
                  'seznam, jeho zemi a pořadové číslo. U jednotlivého prvku pak karta Analýzy '
                  'uvádí, podle kterého národního seznamu se jeho certifikát posuzoval — zemi, '
                  'pořadové číslo a datum vydání.')
    vysvetlivka(doc, 'Proč na tom záleží',
                'Je to údaj o tom, co aplikace mohla mít k dispozici, ne o dokumentu. Bez něj '
                'by chybějící kotva vypadala jako vlastnost podpisu — a kvalifikovaná pečeť '
                'by se bez varování hlásila jako pouhá zaručená.')


def kapitola_6(doc):
    nadpis1(doc, '6', 'Šest záložek')
    odstavec(doc, 'Záložky jsou řazené podle hloubky. Poslední dvě se ukazují jen v expertním '
                  'módu.')

    nadpis2(doc, '6.1', 'Zjednodušený přehled')
    odstavec(doc, 'Verdikt, druh podpisu i s výkladem pojmu, podepsaná osoba, vydavatel '
                  'certifikátu, čas podpisu a vizualizace podpisu tak, jak ji nese PDF. Z této '
                  'záložky se také generuje protokol.')
    obrazek(doc, 'Záložka Zjednodušený přehled s kladným verdiktem.',
            'barevný pruh s výsledkem, údaje o podepsané osobě a vizualizace podpisu', vyska=7.5)

    nadpis2(doc, '6.2', 'Rozšířený přehled')
    odstavec(doc, 'Totéž v devíti pojmenovaných rozbalovacích blocích:')
    odrazky(doc, [
        'podepsaná osoba — u pečeti pečetící osoba, u samostatného razítka jeho vydavatel,',
        'vizualizace podpisu,',
        'důvody podpisu (Reason, Location, kontakt),',
        'certifikát,',
        'časová razítka — u každého i vydavatel, konec platnosti jeho certifikátu a algoritmy, '
        'které razítko slibuje,',
        'evidenční záznamy,',
        'rozhodný okamžik ověření,',
        'kontrola revokace,',
        'použité kryptografické algoritmy.',
    ])
    odstavec(doc, 'U dokumentu, na který odkazuje příloha z eLegalizace, přibývá za kartami '
                  'prvků rámeček s odkazovanými dokumenty. Je vedle karet, ne v nich: odkaz '
                  'nese soubor, ne ten který podpis, a u dokumentu se třemi podpisy by se týž '
                  'blok opakoval třikrát.')
    vysvetlivka(doc, 'Kde jsou lhůty',
                'Dvě lhůty — do kdy lze platnost ověřit a do kdy je třeba zasáhnout — pod '
                'bloky Rozšířeného přehledu nestojí. Odpovídá na ně záložka Digitální '
                'kontinuita, a to nejen datem, ale i tím, čím je dané a jak se liší podle '
                'způsobu ověření. Dvě místa s týmž datem by znamenala dvě místa, která se '
                'musí držet stejného výpočtu.')
    obrazek(doc, 'Záložka Rozšířený přehled s rozbalenými bloky.',
            'zejména bloky „Kontrola revokace“ a „Časová razítka“', vyska=8.0)

    nadpis2(doc, '6.3', 'Analýza')
    odstavec(doc, 'Odpovídá na otázku „proč to dopadlo takhle“. Pro každý prvek jedna karta '
                  'a v ní bloky — které, to se řídí verdiktem a tím, o čem je u konkrétního '
                  'dokumentu co říct:')
    odrazky(doc, [
        '**Formát podpisu** — druh formátu, zapouzdření, úroveň a odpověď na otázku, jde-li '
        'o referenční formát podle prováděcího rozhodnutí Komise (EU) 2015/1506. Řádek '
        'o druhu formátu se ukáže jen u formátů z rodin AdES.',
        '**Zjištěné důvody a doporučení** — očíslované důvody: co se stalo a co pomůže '
        '(kapitola 7).',
        '**Průběh ověřování** — za jakých pravidel se ověřovalo, jak aplikace postupovala '
        'a čím to skončilo; uvnitř složený Podrobný výčet revokačních údajů, které měla po '
        'ruce, s odkazy „Ukaž“ a „Ulož“ u každého.',
        '**Poznámky k výsledku** — výhrady, které verdikt nezměnily, ale čtenář je má znát; '
        'uvnitř podblok s připomínkami, které si postěžovala sama knihovna DSS.',
        '**Možné právní účinky** — jen u platného podpisu a jen tam, kde plynou přímo '
        'z výsledku a z textu předpisu.',
        '**Certifikační podpis**, **Zámek dokumentu**, **Protipodpis (kontrasignatura)** — '
        'jsou-li čím naplnit.',
        '**Kde je šance na úspěšné ověření** a **Kde ověření nebude úspěšné** — odhad chování '
        'jiných nástrojů, vždy v možnosti, ne v jistotě.',
    ])
    odstavec(doc, 'U samostatného časového razítka přibývá blok „Časové razítko“ s jeho časem; '
                  'jméno podepsané osoby, čas podpisu a druh certifikátu jsou u ostatních prvků '
                  'v Rozšířeném přehledu, ne tady. Na kartě s kladným výsledkem by stály mezi '
                  'odpovědí a tím, co s ní dál.')
    obrazek(doc, 'Záložka Analýza u podpisu s neznámým stavem platnosti.',
            'seznam důvodů s doporučeními a pod ním blok Průběh ověřování', vyska=8.0)

    nadpis2(doc, '6.4', 'Digitální kontinuita')
    odstavec(doc, 'Jak dlouho dokument zůstane ověřitelný, kdo ho ještě ověří, a co a do kdy je '
                  'proto potřeba udělat. Tady jsou i tlačítka úkonů. Podrobně v kapitolách '
                  '11 a 12.')

    nadpis2(doc, '6.5', 'Detailní rozbor (expertní mód)')
    odstavec(doc, 'Vše, co o dokumentu vydá diagnostika DSS, po oddílech: výsledek ověření, '
                  'certifikát podepisující osoby, certifikační cesta, důvěryhodná služba '
                  'a její stav v seznamech, časová razítka i s jejich cestou, algoritmy, časy, '
                  'rozsah podpisu, podpisový slovník PDF, detekce změn po podpisu, záznamy VRI, '
                  'globální validační data ve slovníku DSS, struktura podpisů CMS a kryté '
                  'objekty. U certifikátů, seznamů CRL i odpovědí OCSP jsou odkazy „Ukaž“ '
                  'a „Ulož“, u nedůvěryhodného kořene i „Přidat do důvěryhodných“. '
                  'I z této záložky lze generovat protokol. Na začátku záložky stojí '
                  'tlačítko **Inspektor** podpisové struktury (kapitola 8.7).')
    odstavec(doc, 'Karty prvků se jmenují stejně jako v Rozšířeném přehledu — „Podpis“, '
                  'při více prvcích „Pečeť č.: 2“ — a nesou jméno podepsané osoby. Označení, '
                  'pod kterým prvek vede knihovna DSS (například `S-6A5A1DD8…` u podpisu, '
                  '`T-…` u razítka), je uvnitř karty jako řádek **Identifikátor DSS**; podle '
                  'něj se prvek dohledá v Protokolech DSS.')

    nadpis2(doc, '6.6', 'Protokoly DSS (expertní mód)')
    odstavec(doc, 'Surové zprávy knihovny DSS, jednoduchá i podrobná, ke stažení jako XML nebo '
                  'PDF. Užitečné při srovnávání s jinými validátory postavenými nad DSS.')
    vysvetlivka(doc, 'Pozor',
                'Expertní mód se zapíná v základním nastavení. Bez něj poslední dvě záložky '
                'nejsou vidět a některé úkony se nenabízejí.')


def kapitola_7(doc):
    nadpis1(doc, '7', 'Důvody a doporučení')
    odstavec(doc, 'Když ověření nedopadne kladně, nestačí to oznámit. **Důvod** je výklad pro '
                  'čtenáře: co se stalo, co to znamená a co se s tím dá dělat. Důvody jsou '
                  'očíslované a je jich **35**. Čísla se nemění a nepřečíslovávají — odkazuje '
                  'na ně kód, protokoly i tento manuál —, takže odkaz „důvod 21“ ve starším '
                  'protokolu platí i po letech.')
    odstavec(doc, 'Jeden důvod přitom shrnuje i několik různých zjištění pod jedno vysvětlení '
                  'a jedno doporučení: čtenáře nezajímá výčet nálezů, ale odpověď.')
    vysvetlivka(doc, 'Pozor na jmennou shodu',
                'Blok „Důvody podpisu“ v Rozšířeném přehledu je něco jiného: tam jde o důvod, '
                'který uvedla podepsaná osoba (pole Reason v PDF). Důvody v tomhle smyslu jsou '
                'jen v Analýze, v bloku „Zjištěné důvody a doporučení“.')

    nadpis2(doc, '7.1', 'Stavba důvodu')
    tabulka(doc,
            ['Část', 'Co obsahuje', 'Povinná'],
            [['Popis problému',
              'Co se stalo, běžnou řečí a bez zkratek. Stojí vždy první.', 'ano'],
             ['Technicky',
              'Totéž jazykem standardů — jména polí, mezní hodnoty, důvod, proč to standard '
              'zakazuje. Skryto pod odkazem „Podrobněji …“.', 'ne'],
             ['Co pomůže',
              'Co může udělat běžný uživatel: které nastavení přepnout, co doplnit, kdy to '
              'zkusit znovu.', 'ano'],
             ['Pro zkušenější',
              'Rada, která předpokládá pokročilé nastavení — typicky doplnit vlastní CRL '
              'seznam nebo vlastní kotvu důvěry. Je-li na kartě víc důvodů s takovou radou, '
              'stojí jednou, pod spodním z nich: je to rada k ověření, ne k jedné jeho '
              'položce.', 'ne']],
            [3.4, 11.0, 2.2])
    odstavec(doc, 'Tam, kde žádné řešení není, se žádné nepředstírá: u ztracené digitální '
                  'kontinuity se místo rady řekne, že vadu už napravit nelze.')

    nadpis2(doc, '7.2', 'Proč se důvody potlačují')
    odstavec(doc, 'Vypsat všechno naráz by čtenáře poslalo hledat příčinu tam, kde žádná není. '
                  'Silnější nález proto ten slabší umlčí: u podpisu s porušenou integritou nemá '
                  'smysl radit, ať se povolí tolerance k revokačním údajům — dokument je '
                  'změněný a žádné nastavení s tím nehne. Nemá-li aplikace čím výčet naplnit, '
                  'nenabídne ani jeho nadpis.')

    nadpis2(doc, '7.3', 'Náhradní znění')
    odstavec(doc, 'Některé důvody mají víc znění a aplikace vybírá podle toho, co o prvku ví — '
                  'jinak zní věta o razítku, jehož autorita mezitím expirovala, a jinak '
                  'o razítku, které neobstálo z jiné příčiny. Nepodaří-li se okolnosti určit, '
                  'použije se záložní znění. Údaje z dokumentu, jako jméno autority nebo datum '
                  'konce platnosti, se do vět doplňují.')

    nadpis2(doc, '7.4', 'Číselník')
    odstavec(doc, 'Úplný číselník důvodů se generuje přímo ze zdrojového kódu aplikace, takže '
                  'nemůže tvrdit něco jiného než ona, a dá se vytisknout: `docs/duvody-tisk.html`; '
                  've Wordu je jako `docs/AnEP-duvody.docx`.')
    obrazek(doc, 'Výřez ze záložky Analýza: důvody s doporučeními.',
            'blok „Zjištěné důvody a doporučení“ s rozbaleným odkazem „Podrobněji“', vyska=7.0)


def kapitola_8(doc):
    nadpis1(doc, '8', 'Vlastní prohlížeče')

    nadpis2(doc, '8.1', 'K čemu jsou')
    odstavec(doc, 'Ověření odpoví na otázku „platí to?“. Vedle ní stojí druhá, stejně častá: '
                  '**„co v tom souboru vlastně je?“** Na tu odpovídají vlastní prohlížeče. '
                  'Nejsou to náhledy: kterýkoli prvek z nich jde uložit, otevřít jinde, nebo '
                  'poslat zpět k plnému ověření.')
    odrazky(doc, [
        'Systémový prohlížeč certifikátů o odpovědi OCSP neví vůbec a u certifikátu ukazuje '
        'jinou sadu polí, než na kterou se čtenář v tomhle nástroji ptá.',
        'Kontejner není dokument: je to složka v jednom souboru a jinak než otevřením se nedá '
        'zjistit, co je uvnitř.',
        'Prohlížeče čtou soubor tak, jak je uložen — včetně toho, čeho si DSS nevšímá: '
        'nepodepsaných položek, pořadí prvků, metadat archivu. Pohled knihovny DSS na tentýž '
        'soubor ukazují záložky s výsledkem.',
        'Okna jsou nemodální; dá se jich otevřít vedle sebe kolik je třeba.',
    ])

    nadpis2(doc, '8.2', 'Prohlížeč kontejnerů')
    odstavec(doc, 'Otevírá se tlačítkem „Otevřít jako kontejner“ ve stavovém řádku a zvládne '
                  'víc druhů nosičů, než napovídá jméno: archiv ZIP (ASiC, ODF, OOXML '
                  'i libovolný jiný), objekt CMS/PKCS #7, XML, podpis JAdES, datovou zprávu '
                  'ISDS i PDF s přílohami.')
    odstavec(doc, 'Vlevo je strom položek se složkami; u známých jmen je u položky i její role '
                  '— podpis XAdES či CAdES, manifest, manifest ASiC s výčtem krytých souborů, '
                  'časové razítko, evidenční záznam, validační data. Vpravo se u vybrané '
                  'položky vypíšou její údaje (velikost uložená i původní, způsob komprese, '
                  'kontrolní součet, čas změny) a pod nimi náhled: text, obrázek, nebo '
                  'šestnáctkový výpis u toho, co ani jedno není. XML se pro čtení zformátuje.')
    tabulka(doc,
            ['Tlačítko', 'Co udělá'],
            [['Uložit prvek…', 'Vytáhne vybranou položku ven jako soubor.'],
             ['Uložit vše…', 'Rozbalí celý kontejner do zvoleného adresáře se zachováním '
                             'struktury.'],
             ['Uložit písemnosti…',
              'U datové zprávy uloží jen to, co odesílatel poslal, pod jeho názvy. Objeví se '
              'jen tam, kde je co ukládat.'],
             ['Otevřít v aplikaci',
              'Uloží prvek do dočasného souboru a otevře jej v programu přiřazeném jeho typu.'],
             ['Otevřít jako kontejner',
              'Otevře vnořený prvek v novém okně vedle rodičovského — archiv, objekt CMS, XML '
              'i PDF s vlastními přílohami.'],
             ['Ověřit v AnEP',
              'Předá prvek aplikaci k plnému ověření, jako byste jej sami vložili do vstupní '
              'plochy.'],
             ['Kopírovat SHA-256', 'Otisk vybraného prvku do schránky.']],
            [4.4, 12.2])
    odstavec(doc, 'Nad stromem je filtr podle názvu; s napsaným hledáním se rozbalí i to, co je '
                  'jinak složené. Validační data uložená jako samostatné položky mívají strojová '
                  'jména, ze kterých čtenář nepozná odpověď OCSP od certifikátu — takovým se ve '
                  'stromu dá čitelné jméno (`odpoved-ocsp-1.ors`, `certifikat-1.cer`), kdežto '
                  'skutečné jméno položky zůstává na kartě pod „Název v kontejneru“ a hledá se '
                  'v obou.')
    obrazek(doc, 'Prohlížeč kontejneru: strom položek, karta prvku a náhled.',
            'okno prohlížeče nad kontejnerem ASiC-E se sloupcem rolí', vyska=8.5)

    nadpis2(doc, '8.3', 'Prohlížeč certifikátů')
    odstavec(doc, 'Certifikát vypsaný ve slovech. Kromě subjektu, vydavatele, sériového čísla, '
                  'doby platnosti, algoritmu podpisu a veřejného klíče řekne i to, co bývá '
                  'schované v rozšířeních:')
    odrazky(doc, [
        'klíčová užití česky — digitální podpis, neodvolatelnost, šifrování klíče, '
        'podepisování certifikátů i CRL — a rozšířená užití zvlášť (časová razítka, '
        'podepisování odpovědí OCSP, ochrana e-mailu, podepisování dokumentů a další);',
        'základní omezení: koncový certifikát, nebo certifikační autorita, a do jaké hloubky;',
        'přístupový bod OCSP, adresa ke stažení certifikátu vydavatele a distribuční body CRL;',
        'prohlášení podle QcStatements — kvalifikovanost podle eIDAS, účel (podpis, pečeť, '
        'autentizace webových stránek), uložení klíče v prostředku QSCD, limit transakce, doba '
        'uchovávání údajů, právo kterých zemí se použije a role podle PSD2;',
        'zapsané osobní doklady — nese-li certifikát číslo občanského průkazu nebo cestovního '
        'pasu podle registračního schématu z ETSI EN 319 412-1 čl. 5.1.3, vypíše se druh '
        'dokladu a jeho číslo samostatným blokem.',
    ])
    vysvetlivka(doc, 'Proč zvlášť',
                'Číslo dokladu je i v atributu serialNumber jména držitele, mezi ostatními '
                'atributy a ve tvaru „IDCCZ-…“. Tam ho ale najde jen ten, kdo ví, co hledá — '
                'schéma není samovysvětlující a sériových čísel je v certifikátu víc.')

    nadpis2(doc, '8.4', 'Prohlížeč seznamu CRL')
    odstavec(doc, 'Reálný seznam CRL má tisíce záznamů a otázka se přitom týká jediného: '
                  '„je můj certifikát na seznamu, a od kdy“. Prohlížeč se proto **otevře rovnou '
                  'na záznamu toho certifikátu, o který jde**; není-li na seznamu, otevře se na '
                  'prvním záznamu a řekne to nad tabulkou.')
    odstavec(doc, 'Nad tabulkou stojí hlavička seznamu: vydavatel (s odkazem na jeho '
                  'certifikát), kdy byl seznam vydán a kdy má vyjít další, jeho číslo, je-li to '
                  'delta a k čemu, a kolik certifikátů uvádí. Tabulka nese sériové číslo, datum '
                  'odvolání a jeho důvod, dá se řadit podle sloupců a hledat v ní podle '
                  'sériového čísla — šestnáctkově i desítkově. Hledání je cesta pro seznam '
                  'otevřený samostatně, kde není od čeho se odrazit.')
    obrazek(doc, 'Prohlížeč seznamu CRL otevřený na hledaném záznamu.',
            'hlavička seznamu a tabulka odvolaných certifikátů', vyska=7.5)

    nadpis2(doc, '8.5', 'Prohlížeč odpovědi OCSP')
    odstavec(doc, 'Stav odpovědi jako takové (úspěšná, chybný požadavek, vnitřní chyba, zkuste '
                  'později, vyžadován podpis, neautorizováno), identifikace respondéra, čas '
                  'vydání, algoritmus podpisu a certifikáty, které odpověď nese. Pro každý '
                  'dotázaný certifikát pak sériové číslo, hashovací funkci v identifikátoru, '
                  'stav (dobrý, revokovaný, neznámý), u revokovaného datum a důvod pojmenovaný '
                  'slovem, a doby platnosti odpovědi.')

    nadpis2(doc, '8.6', 'Přílohy dokumentu')
    odstavec(doc, 'Přílohy PDF se sbírají ze tří různých míst — ze jmenného stromu vložených '
                  'souborů (to je ta „sponka“), z anotací umístěných na stránce a z polí '
                  'přidružených souborů, kterými nosí data dokumenty typu PDF/A-3 nebo '
                  'faktury — a u každé se pamatuje, kudy se do dokumentu dostala. Uživatel mezi '
                  'tím nerozlišuje, a proto je vidí pohromadě, ve stejném okně jako kontejner '
                  'a se všemi jeho možnostmi.')
    odstavec(doc, 'Příloha bývá sama podepsaným dokumentem, takže tlačítko „Ověřit v AnEP“ je '
                  'u ní to nejužitečnější. U heslem chráněného dokumentu se přílohy vytahují '
                  'už dešifrované, takže se heslo nezadává podruhé.')

    nadpis2(doc, '8.7', 'Inspektor podpisové struktury')
    odstavec(doc, 'Ověření říká, co z podpisu vyvodila knihovna DSS. Inspektor ukazuje, z čeho '
                  'vycházela: podpisovou strukturu tak, jak je v souboru uložena, včetně toho, '
                  'co DSS nepoužil nebo přehlédl — přebytečné položky v úložišti ověřovacích '
                  'údajů, záznam VRI bez podpisu, starší podobu úložiště v dřívějších revizích.')
    odstavec(doc, 'Otevírá se tlačítkem na začátku záložky Detailní rozbor, tedy jen '
                  'v expertním módu. Jméno tlačítka říká, co se rozebírá: **Inspektor PDF**, '
                  '**CAdES**, **ASiC**, **XAdES** nebo **JAdES**. Je-li k dokumentu předložen '
                  'externí podpis, rozebírá se podpis a dokument mu slouží k přepočtu otisku. '
                  'U formátu, který Inspektor nerozkládá, tlačítko není.')
    tabulka(doc,
            ['Formát', 'Co strom ukazuje'],
            [['PDF', 'podpisy a pečetě (slovník podpisu, rozsah /ByteRange, objekt CMS), '
                     'časová razítka dokumentu i razítka vnořená v podpisech, úložiště /DSS, '
                     'záznamy /VRI s tím, ke kterému podpisu patří, a revize dokumentu — u každé '
                     'co ji uzavírá a jak tehdy vypadalo /DSS; revizi bez jediného objektu '
                     '(hybridní soubor z exportu Wordu) pozná'],
             ['CAdES', 'objekt CMS: podepisující, podepsané a nepodepsané atributy, vložená '
                       'razítka, certifikáty, seznamy CRL a odpovědi OCSP'],
             ['ASiC', 'podpisy CAdES i XAdES, razítka kontejneru, manifesty ASiC s kontrolou '
                      'otisků i seznam souborů META-INF/manifest.xml, podepisované soubory'],
             ['XAdES', 'odkazy SignedInfo s přepočtem otisků, KeyInfo, podepsané '
                       'a nepodepsané vlastnosti, razítka, ověřovací údaje, kontrasignatury'],
             ['JAdES', 'chráněná hlavička (x5c, sigD, adoTst), podepsaný obsah a vlastnosti '
                       'etsiU v uloženém pořadí včetně protipodpisu; všechny tři serializace']],
            [2.6, 14.0])
    odstavec(doc, 'Strom začíná sbalený, rozbaluje se postupně. Barevná tečka u prvku značí '
                  'zjištění: zelená v pořádku, oranžová upozornění, červená chyba. Nad stromem '
                  'je hledání (Enter skočí na další shodu), u PDF pruh celého souboru s revizemi '
                  'a s rozsahem, který vybraný podpis kryje — klepnutím do pruhu se vybere revize.')
    odstavec(doc, 'Detail vybraného prvku má dvě oddělené části: **Zjištění o souboru** — co '
                  'Inspektor sám přepočítal (otisk podepsaných dat, otisk v razítku, otisk '
                  'podpisového certifikátu, soulad rozsahu s revizí) — a **Podle výsledku '
                  'ověření** — co s prvkem udělal DSS: výsledek podpisu či razítka, roli '
                  'certifikátu v řetězu, který revokační údaj rozhodl a který nebyl přijat nebo '
                  'použit. Pod nimi jsou údaje prvku; u řádků s odkazem „přejít →“ se skočí '
                  'na související prvek a oddíl „Týž prvek je v souboru také zde“ vede ke všem '
                  'výskytům téhož certifikátu či údaje (v CMS, v /DSS, ve VRI, v dřívějších '
                  'revizích). Certifikát a odpověď OCSP se vypíšou celé, seznam CRL se otevře '
                  've vlastním prohlížeči; každý prvek jde uložit.')
    vysvetlivka(doc, 'Pozor',
                'Inspektor o výsledku ověření nerozhoduje a jeho zjištění se do verdiktu '
                'nepromítají. Nesouhlasí-li otisk, řekne to i ověření samo — Inspektor jen '
                'ukáže, kde přesně.')
    obrazek(doc, 'Inspektor PDF — strom podpisové struktury, pruh revizí a detail prvku.',
            'okno Inspektoru nad podepsaným PDF s rozbaleným podpisem', vyska=8.5)

    nadpis2(doc, '8.8', 'Kudy se k nim dostat')
    tabulka(doc,
            ['Odkud', 'Co se otevře'],
            [['Stavový řádek → Otevřít jako kontejner',
              'prohlížeč kontejnerů nad právě ověřeným souborem'],
             ['Stavový řádek → sponka', 'přílohy dokumentu'],
             ['Analýza → Průběh ověřování → Podrobný výčet → Ukaž',
              'seznam CRL otevřený na dotčeném certifikátu, nebo odpověď OCSP'],
             ['Detailní rozbor → Ukaž',
              'certifikát, seznam CRL nebo odpověď OCSP z kterékoli cesty či slovníku'],
             ['Detailní rozbor → Inspektor', 'podpisová struktura dokumentu, nebo externího podpisu, byl-li předložen'],
             ['Prohlížeč kontejnerů', 'kterýkoli prvek uvnitř, i vnořený kontejner'],
             ['Vložení holého souboru do aplikace',
              'certifikát, seznam CRL nebo odpověď OCSP rovnou v prohlížeči']],
            [6.4, 10.2])
    odstavec(doc, 'Vedle „Ukaž“ stojí všude i „Ulož“: prvek se uloží přesně tak, jak v dokumentu '
                  'leží. Odkazy chybějí jen tam, kde aplikace bajty prvku nemá — slibovat '
                  'zobrazení a pak hlásit chybu by bylo horší.')


def kapitola_9(doc):
    nadpis1(doc, '9', 'Důvěra a revokace')

    nadpis2(doc, '9.1', 'Odkud pocházejí kotvy důvěry')
    odstavec(doc, 'Kotvy důvěry pocházejí z evropských důvěryhodných seznamů (LOTL) — kolem '
                  'třiceti národních seznamů a několika tisíc certifikátů. Uloží se do místní '
                  'mezipaměti. Vedle toho se čte i kořenové úložiště Windows.')
    odstavec(doc, 'Důvěryhodné seznamy nejsou jen seznam certifikátů: nesou i to, jakou službu '
                  'certifikát zastupuje, pod jakým poskytovatelem a kdy služba získala nebo '
                  'ztratila kvalifikovaný status. Právě z toho aplikace odvozuje, zda lze '
                  'časové razítko uznat i po expiraci jeho certifikátu — a kdy byla služba ze '
                  'seznamů vyřazena.')

    nadpis2(doc, '9.2', 'Revokační údaje a jejich původ')
    odstavec(doc, 'Revokační údaje se zjišťují živě: nejprve OCSP, pak CRL jako záloha, '
                  'a doplňují se o to, co dokument nese sám. U každého údaje aplikace sleduje '
                  'jeho původ:')
    tabulka(doc,
            ['Původ', 'Co z toho plyne'],
            [['z dokumentu',
              'Údaj nikam nezmizí. Závěr na něm postavený zopakuje kdokoli, komu dokument '
              'předložíte.'],
             ['z vlastního úložiště',
              'Údaj má po ruce jen tato aplikace. Jiný nástroj ho nemá odkud vzít — dokud se '
              'úkonem nevloží do dokumentu.'],
             ['ze sítě',
              'Údaj se právě podařilo stáhnout. Že půjde stáhnout i příště, zaručeno není: '
              'vydavatel jednou zveřejňování ukončí.']],
            [4.4, 12.2])
    vysvetlivka(doc, 'Pozor',
                'Původ rozhodujícího údaje je klíč k celé záložce Digitální kontinuita. Právě '
                'on rozhoduje, jestli aplikace umí říct, do kdy bude dokument ověřitelný — '
                'nebo jestli žádné takové datum neexistuje.')

    nadpis2(doc, '9.3', 'Tři skupiny údajů')
    odstavec(doc, 'V Podrobném výčtu leží revokační údaje ve třech skupinách, protože '
                  '„přípustný“ a „využitý“ není totéž:')
    odrazky(doc, [
        '**využita** — o tyto údaje se závěr opřel;',
        '**nevyužita** — standardy je připouštějí, ale závěr na nich nestojí. Typicky odpověď '
        'stažená dnes u podpisu, který se uzavřel podle údaje vloženého v dokumentu. Vada to '
        'není;',
        '**nedala se využít** — u každého i důvod, proč ne.',
    ])
    odstavec(doc, 'Je i čtvrtá možnost: na stav revokace se správně vůbec nikdo neptal. Stává '
                  'se to tam, kde je podpisový certifikát sám kotvou důvěry — pečeť '
                  'kvalifikované ověřovací služby bývá v důvěryhodném seznamu vedena jako '
                  'vlastní digitální identita té služby — a norma pro kotvu důvěry stav '
                  'revokace nezjišťuje. Aplikace to řekne výslovně, jinak by poctivé „nic se '
                  'nezjišťovalo, a správně“ vypadalo stejně jako porucha.')

    nadpis2(doc, '9.4', 'Vlastní zásoby')
    odrazky(doc, [
        '**Vlastní úložiště CRL** — seznamy, které si uživatel nashromáždil. Spravuje se '
        'v pokročilém nastavení (Vlastní CRL seznamy → Spravovat). Hodí se u starých podpisů, '
        'kde už vydavatel nic nezveřejňuje.',
        '**Vlastní kotvy důvěry** — certifikáty, které evropské seznamy neuvádějí, například '
        'interní firemní autorita. Aplikace umí jedním tlačítkem vložit vydávající autority '
        'tuzemských komerčních linek (I.CA, PostSignum, eIdentity, NCA, KB).',
    ])
    odstavec(doc, 'Vlastní úložiště CRL je **záložní zdroj**, ne přednostní: sáhne se po něm '
                  'teprve tehdy, když se použitelný údaj nepodařilo získat obvyklou cestou — '
                  'tedy když odpověď OCSP není vůbec, nebo je vydaná až po expiraci '
                  'certifikátu, a tedy o době jeho platnosti nic neříká. Ověřování dosud '
                  'platných certifikátů tím zůstává nedotčené.')
    odstavec(doc, 'Ze seznamů od téhož vydavatele si aplikace vybírá ten, který o certifikátu '
                  'skutečně může vypovídat. Seznam vydaný dřív, než certifikát vznikl, o něm '
                  'mlčet musí — jeho mlčení tedy není důkazem ničeho; a seznam vydaný před '
                  'rozhodným okamžikem nemůže vyloučit odvolání, ke kterému došlo potom. Oba '
                  'se proto vyřadí dřív, než se začne vybírat.')
    odstavec(doc, 'Aplikace vždy přizná, že závěr stojí na kotvě nebo na seznamu, který jiný '
                  'nástroj nemá odkud vzít — a v záložce Digitální kontinuita z toho vyvodí '
                  'důsledek pro to, kdo výsledek ověří. Při zachování ověřitelnosti se takový '
                  'materiál **zapíše do dokumentu**, jinak by uchovaný soubor postrádal právě '
                  'ten důkaz, na kterém jeho verdikt stál.')

    nadpis2(doc, '9.5', 'Obnova důvěryhodných seznamů')
    odstavec(doc, 'V pokročilém nastavení je tlačítko „Načíst seznam EUTL znovu“. Hodí se, když '
                  'je mezipaměť stará nebo prázdná, nebo když se změnil stav některé služby '
                  'a je potřeba to promítnout hned. Načtení trvá desítky sekund a jeho průběh '
                  'je v dialogu vidět; po dokončení se překreslí i hlavička výsledku, aby '
                  'varování o chybějících seznamech nezůstalo viset.')
    obrazek(doc, 'Pokročilé nastavení, sekce Důvěryhodné seznamy EU (EUTL).',
            'tlačítko pro opětovné načtení se stavovým textem o průběhu', vyska=6.0)


def kapitola_10(doc):
    nadpis1(doc, '10', 'Přísnost a ústupky')
    odstavec(doc, 'Ve výchozím stavu se ověřuje striktně podle standardů. Ústupky jsou '
                  'jednotlivě zapínatelné volby, které povolují to, co standardy zapovídají. '
                  'V základním nastavení je drží jediný přepínač „Ověřování podpisů a pečetí '
                  'dle standardů“ — jeho zrušením se povolí všechny naráz.')

    nadpis2(doc, '10.1', 'Čtyři ústupky')
    tabulka(doc,
            ['Volba', 'Co povolí'],
            [['Využij i revokační informace, vydané po expiraci certifikátu',
              'revokační údaj vydaný až po skončení platnosti certifikátu, o kterém vypovídá'],
             ['(ústupek na expirovaného vydavatele)',
              'revokační údaj od vydavatele, jemuž mezitím vypršel vlastní certifikát'],
             ['Ignorovat zastaralost algoritmu SHA1',
              'nepovažovat SHA-1 sám o sobě za důvod k odmítnutí'],
             ['Ověřovat i podpisy ve formátu PDF-NOT-ETSI',
              'ověřovat i nereferenční podpisy z doby před eIDAS']],
            [6.4, 10.2])

    nadpis2(doc, '10.2', 'Dvě samostatná nastavení')
    odrazky(doc, [
        '**Ověřit k zadanému datu a času** — ověření k okamžiku v minulosti. Výsledek pak '
        'neříká nic o platnosti k dnešnímu dni.',
        '**Ověřování časových razítek podle standardů** — vypnuto znamená přísnější '
        'posuzování: razítko se uzná jen do konce řádné doby platnosti jeho certifikátu, '
        'nikoli do vyřazení služby z důvěryhodných seznamů. Tím se aplikace přiblíží chování '
        'CzechPOINTu a podatelen. Je to přísnější čtení, ne ústupek.',
    ])
    obrazek(doc, 'Pokročilé nastavení: sekce Čas ověření a Uvolnění validační politiky DSS.',
            'všechna zaškrtávací pole ústupků a pole pro zadaný čas', vyska=7.0)

    nadpis2(doc, '10.3', 'Jednorázová validační politika')
    odstavec(doc, 'Každý ústupek i zadaný čas platí jen na jedno ověření. Po něm se nastavení '
                  'vrací do výchozího stavu. Důvod je prostý: kladný verdikt se nesmí tiše '
                  'opřít o volbu, kterou uživatel zapnul před týdnem a dávno na ni zapomněl.')
    vysvetlivka(doc, 'Pozor',
                'Vlastní kotvy důvěry ani vlastní úložiště CRL se nevracejí — ty platí trvale. '
                'Jednorázová je jen validační politika, tedy ústupky a zadaný čas. Expertní '
                'mód zůstává rovněž.')

    nadpis2(doc, '10.4', 'Rozhodující ústupek')
    odstavec(doc, 'Aplikace měří, jestli se ústupek na výsledku opravdu podílel: ověření '
                  'proběhne znovu bez něj a oba výsledky se porovnají. U kladného verdiktu se '
                  'proto vypíšou jen ty ústupky, bez kterých by závěr nevznikl — ne ty, které '
                  'byly jen zapnuté. Věta „ověřeno s ústupkem“ tedy vždy něco znamená. Týž '
                  'výčet jde do protokolu a platí i pro vlastní zásoby a pro zadaný čas.')


def kapitola_11(doc):
    nadpis1(doc, '11', 'Úkony nad dokumentem')
    odstavec(doc, 'Žádný z úkonů nesahá na podepsané bajty. Každý zapisuje nový soubor; původní '
                  'zůstává nedotčen. Tlačítka stojí v záložce Digitální kontinuita, pod textem, '
                  'který říká, co udělají.')

    nadpis2(doc, '11.1', 'Zachovat ověřitelnost')
    odstavec(doc, 'Podle toho, co dokument už nese a v jakém je stavu, vedou k témuž cíli tři '
                  'cesty. Aplikace vybere jednu a pojmenuje ji na tlačítku i ve jméně výstupního '
                  'souboru.')
    tabulka(doc,
            ['Varianta', 'Tlačítko', 'Co udělá', 'Kdo výsledek ověří'],
            [['Uchování', 'Uchovat',
              'Doplní chybějící revokační informace a přidá archivní časové razítko; vznikne '
              'podpis v úrovni LTA.',
              'všechny obvyklé validátory'],
             ['Prodloužení', 'Prodloužit',
              '==Doplní chybějící revokační informace a přidá další časové razítko tam, kde '
              'je řada razítek přerušená nebo kde poslednímu razítku už vypršel certifikát.==',
              'jen nástroje, které razítka posuzují podle standardů (AnEP, DSS a validátory '
              'nad ní)'],
             ['Přerazítkování', 'Přerazítkovat',
              'Přidá ==revokační informace, které o certifikátu něco vypovídají (ne odpověď '
              '„nerevokován“ vydanou po expiraci bez závazku vydavatele),== a další časové '
              'razítko.',
              'spíše jen nástroje s vlastní zásobou revokačních informací (CzechPOINT, '
              'podatelny); není garantováno']],
            [2.8, 2.4, 6.2, 5.2])
    vysvetlivka(doc, 'Pozor',
                'Tlačítko je vždy jen jedno — varianty se vylučují a uživatel mezi nimi '
                'nevybírá. Aplikace nabídne nejsilnější variantu, na kterou dosáhne. Úkon je '
                'přitom opakovatelný: dokument, který už jednou uchován byl, se dá uchovat '
                'znovu, až bude po čase potřeba.')

    nadpis2(doc, '11.2', 'Souběh více podpisů')
    odstavec(doc, 'Nové časové razítko kryje celý soubor, takže úkon se provádí jednou pro celý '
                  'dokument a nelze jej přidat jen k některému prvku. U dokumentu s několika '
                  'podpisy proto platí:')
    odrazky(doc, [
        'Úkon se nabídne, pokud dosáhne alespoň na jeden prvek.',
        'Prvku, na který nedosáhne, neuškodí — co je na něm dnes sporné, zůstane sporné '
        'i potom.',
        'Záložka Digitální kontinuita to říká výslovně: vypíše, kterým prvkům úkon prospěje '
        'a kterým ne, a u každého uvede proč.',
    ])

    nadpis2(doc, '11.3', 'Vytvořit forenzní balíček')
    odstavec(doc, 'Pro dokumenty, které zachovat nelze — a také pro soubory, jejichž formát '
                  'aplikace vůbec nepřečte. Zabalí nedotčený originál do kontejneru ASiC-E '
                  'spolu se vším ověřovacím materiálem, který se podařilo sehnat, a soupisem '
                  'toho, odkud co pochází — a opatří to časovým razítkem. Znalec pak má doklad, '
                  'že v okamžiku zafixování všechny podklady existovaly.')
    odstavec(doc, 'Do balíčku se dá přidat i další soubor: smlouva se svými přílohami, zpráva '
                  's tím, co k ní bylo připojeno. Zarazítkovat je zvlášť by každý zafixovalo '
                  'v jiném okamžiku a neřeklo by nic o tom, že patří k sobě.')
    odstavec(doc, 'Ověřitelnost samotného dokumentu to neprodlouží. Přesto je vhodné balíček '
                  'pořídit co nejdříve: později už nemusí být dostupné informace, které dnes '
                  'dostupné jsou.')
    obrazek(doc, 'Okno forenzního balíčku se seznamem souborů, které do něj půjdou.',
            'dialog Forenzní balíček', vyska=7.0)

    nadpis2(doc, '11.4', 'Jména výstupních souborů')
    odstavec(doc, 'Výstup nese značku varianty a datum, do kterého je ověřitelnost posunuta:')
    tabulka(doc,
            ['Varianta', 'Příklad jména výstupu'],
            [['Uchování', '`dokument_Uchovano_do_03_08_2034.pdf`'],
             ['Prodloužení', '`dokument_Prodlouzeno_do_03_08_2034.pdf`'],
             ['Přerazítkování', '`dokument_Prerazitkovano_do_03_08_2034.pdf`'],
             ['Forenzní balíček', '`FB_dokument_pdf_Do_03_08_2034.asice`']],
            [4.6, 12.0])
    odstavec(doc, 'Značka je bez diakritiky a bez mezer záměrně: jméno putuje e-mailem, do '
                  'spisových služeb a na cizí systémy. Značky se nevrší — opakovaný úkon tu '
                  'předchozí nahradí. Datum na konci není vlastností dokumentu ani podpisu, '
                  'nýbrž důsledkem volby zdroje časových razítek: sahá tak daleko, kam sahá '
                  'platnost certifikátu vybrané autority.')

    nadpis2(doc, '11.5', 'Kdy se úkon nenabídne')
    odstavec(doc, 'Úkon se nabídne jen tam, kde je možný a zároveň poctivý. Dokument, jehož '
                  'certifikát byl odvolán dříve, než se dá prokázat existence podpisu, uchovat '
                  'nelze — a aplikace to řekne, místo aby nabídla tlačítko, které nic nespraví. '
                  'Vetem přes všechny prvky je porušená integrita a přítomnost evidenčního '
                  'záznamu.')


def kapitola_12(doc):
    nadpis1(doc, '12', 'Digitální kontinuita')
    odstavec(doc, 'Otázka „do kdy je dokument ověřitelný“ nemá jednu odpověď, protože to není '
                  'vlastnost dokumentu, nýbrž vztah mezi dokumentem a tím, kdo ho ověřuje. '
                  'Záložka ji proto zodpovídá třema očima.')

    nadpis2(doc, '12.1', 'Tři čtenáři')
    tabulka(doc,
            ['Čtenář', 'Za jakých podmínek'],
            [['AnEP',
              'Tak, jak se právě ověřovalo — se zapnutými ústupky, vlastními zásobami '
              'a zadaným časem.'],
             ['Validátory podle standardů',
              'Bez jakýchkoli odchylek. Tento blok se ukazuje jen tam, kde se zdejší běh od '
              'standardů odchýlil.'],
             ['CzechPOINT, podatelny a další',
              'Nástroje s vlastní zásobou revokačních informací a shovívavostí k SHA-1, ale '
              'posuzující časová razítka podle řádné platnosti jejich certifikátů, bez ohledu '
              'na vyřazení služeb ze seznamů.']],
            [4.6, 12.0])
    vysvetlivka(doc, 'Pozor',
                'O třetím okruhu mluví aplikace vždy v možnosti („mohly by“) a končí poznámkou, '
                'že jde o odhad, ne o jistotu. Nezná jejich verzi, nastavení ani obsah jejich '
                'zásob.')

    nadpis2(doc, '12.2', 'Struktura záložky')
    odstavec(doc, 'Záložka má čtyři části:')
    cislovane(doc, [
        'Co je digitální kontinuita — výklad pojmu, nezávislý na dokumentu.',
        'Jaké možnosti udržování digitální kontinuity nabízí AnEP — rovněž výklad, se dvěma '
        'vnořenými bloky po variantách úkonu.',
        'Po prvcích — pro každý podpis, pečeť či samostatné razítko sada bloků, které '
        'odpovídají na otázku, kdo a do kdy jej ověří.',
        'Dokument jako celek — jen tam, kde je prvků víc. Obsahuje popis souběhu, oba bloky '
        'o úkonu a výčet toho, jak se úkon projeví na kterém prvku.',
    ])
    odstavec(doc, 'U každého prvku stojí čtyři odpovědi: jak dlouho dokáže platnost ověřit '
                  'AnEP, jak dlouho ji ověří validátory podle standardů, jak dlouho CzechPOINT '
                  'a podatelny, a jak dlouho je možné ověřitelnost zachovat. Nad nimi je '
                  'zavřený blok „Rozbor podpisu“ s fakty, ze kterých to všechno plyne.')
    odstavec(doc, 'U dokumentu s jediným podpisem čtvrtá část odpadá a oba bloky o úkonu stojí '
                  'přímo u něj. Všechny bloky jsou ve výchozím stavu sbalené — čtenář má '
                  'nejprve vidět nadpisy, podrobnosti až na vyžádání.')
    obrazek(doc, 'Záložka Digitální kontinuita nad dokumentem se dvěma podpisy.',
            'nahoře dva výkladové bloky, pod nimi prvky se svými bloky a dole část Dokument '
            'jako celek', vyska=8.5)

    nadpis2(doc, '12.3', 'Co určuje datum')
    odstavec(doc, 'Prvek visí v čase na posledním článku, který ho drží, a lhůtu určuje ten, '
                  'který povolí nejdřív:')
    odrazky(doc, [
        '**Důkaz existence** — s platným časovým razítkem smí být podpisový certifikát dávno '
        'po expiraci; bez něj končí všechno s ním.',
        '**Dostupnost revokačních informací** — údaj v dokumentu ani ve vlastním úložišti '
        'nezvětrá, ten stahovaný ze sítě jen dokud ho vydavatel zveřejňuje.',
        '**Uznatelnost vloženého údaje** — že údaj v souboru zůstane, ještě neznamená, že půjde '
        'uznat: po expiraci svého vydavatele je to údaj s prošlým certifikátem. Vyváže z toho '
        'jen razítko nad celým souborem.',
        '**Uznání razítka** — kvalifikované drží do vyřazení služby z důvěryhodných seznamů, '
        'což je datum, které nikdo předem nezná; jinak končí s certifikátem razítkovací '
        'autority.',
    ])

    nadpis2(doc, '12.4', 'Když lhůtu stanovit nelze')
    odstavec(doc, 'Zvláštní a poměrně častý případ: dnešní kladný závěr stojí na revokačním '
                  'údaji, který si aplikace právě stáhla ze sítě, a doba, po kterou se na jeho '
                  'dostupnost dalo spoléhat, už uplynula — typicky u podpisu v úrovni T, jehož '
                  'certifikát expiroval.')
    odstavec(doc, 'V takovém případě aplikace žádné datum neuvede. V nadpisu bloku „Jak dlouho '
                  'dokáže AnEP ověřit platnost?“ zůstane holá otázka a v nadpisu bloku o úkonu '
                  'stojí „Nejlépe ihned“. Uvnitř bloku je vysvětlení: údaj se teď stáhnout '
                  'podařilo, že bude dostupný i příště ale zaručeno není, a proto dobu stanovit '
                  'nelze — a kdyby dostupný zůstal, sahala by ověřitelnost do konkrétního data, '
                  'které je tam rovněž uvedeno.')
    vysvetlivka(doc, 'Pozor',
                'Datum se tím neztrácí, jen přestává být slibem. Čtenář ho má k dispozici '
                'i s podmínkou, za které platí. To je celý rozdíl proti stavu, kdy stálo '
                'v nadpisu jako tvrzení.')


def kapitola_13(doc):
    nadpis1(doc, '13', 'Zdroje časových razítek')
    odstavec(doc, 'Zdroje se spravují v pokročilém nastavení jako pojmenované profily: adresa, '
                  'přihlášení a algoritmus otisku. Aplikace umí předvyplnit veřejný zdroj '
                  'Sectigo Qualified.')

    nadpis2(doc, '13.1', 'Sonda při zadání')
    odstavec(doc, 'Při uložení profilu se zdroj otestuje. Zjistí se:')
    odrazky(doc, [
        'konec platnosti certifikátu autority — právě podle něj se pak počítá, kam až úkon '
        'posune ověřitelnost;',
        'výčet podporovaných otisků, aby se později nevybral algoritmus, který autorita '
        'odmítne. Nabídka nemusí být souvislá: autorita, která bere SHA-256 a SHA-512, nemusí '
        'brát SHA-384.',
    ])
    vysvetlivka(doc, 'Pozor',
                'Nedostupnost zdroje při sondě neznamená jeho odmítnutí — profil lze uložit '
                'i tak. Chybí pak ale údaj o konci platnosti certifikátu a aplikace nemůže '
                'dopředu říct, kam až úkon dosáhne; zjistí se to při pořízení nejbližšího '
                'razítka.')
    obrazek(doc, 'Dialog správy zdrojů časových razítek.',
            'seznam profilů a formulář pro zadání adresy, přihlášení a algoritmu otisku',
            vyska=7.0)

    nadpis2(doc, '13.2', 'Když vybraný zdroj nestačí')
    odstavec(doc, 'Certifikát vybrané autority může skončit dříve než dosavadní ověřitelnost '
                  'dokumentu. Aplikace na to upozorní a poradí buď vybrat zdroj s déle platným '
                  'certifikátem, nebo s úkonem počkat — nové razítko dosáhne tam, kam sahá '
                  'certifikát jeho autority, ať se přidá dnes nebo za rok.')


def kapitola_14(doc):
    nadpis1(doc, '14', 'Protokoly a export')

    nadpis2(doc, '14.1', 'Protokol o ověření')
    odstavec(doc, 'Ze Zjednodušeného přehledu, Rozšířeného přehledu i Detailního rozboru se '
                  'generuje protokol jako PDF/A-2b, tedy v archivním formátu, který přijímají '
                  'Formátová pravidla Národního archivu. Protokol nese záhlaví s údaji '
                  'o souboru a okamžikem ověření, vizualizaci podpisu v rámečku, výklad pojmů '
                  'a výčet toho, čím se ověřování řídilo. Protokol z obou přehledů se jmenuje '
                  'podle souboru, o kterém vypovídá.')
    odstavec(doc, 'Okamžik v záhlaví je okamžik **ověření**, ne vzniku protokolu: ten může být '
                  'o hodiny později, když uživatel mezitím listuje záložkami. Do protokolu jde '
                  'i to, co má čtenář na obrazovce sbalené — rozbalovací blok se exportuje celý.')

    nadpis2(doc, '14.2', 'Dvě volby')
    odstavec(doc, 'Před generováním se aplikace zeptá na dvě věci; obojí je ve výchozím stavu '
                  'vypnuté:')
    tabulka(doc,
            ['Volba', 'Důsledek'],
            [['Vložit ověřovaný soubor do protokolu jako přílohu',
              'Protokol ponese i soubor, o kterém vypovídá, takže se dá ověřit znovu z jediného '
              'souboru. Vysází se ale jako PDF/A-3, protože PDF/A-2 přílohu, která sama není '
              'PDF/A, nepřipouští (ISO 19005-2, klauzule 6.8).'],
             ['Opatřit protokol kvalifikovaným elektronickým časovým razítkem',
              'Razítko doloží, kdy protokol vznikl a že se od té doby nezměnil. Pořizuje se od '
              'zdroje vybraného v nastavení; není-li žádný, otázka se nenabídne.']],
            [6.0, 10.6])
    obrazek(doc, 'Dialog s volbami protokolu před jeho vygenerováním.',
            'dvě zaškrtávací pole a tlačítko Generovat', vyska=5.5)

    nadpis2(doc, '14.3', 'Zprávy DSS')
    odstavec(doc, 'V expertním módu lze v záložce Protokoly DSS stáhnout původní zprávy '
                  'knihovny DSS (jednoduchou i podrobnou) jako XML nebo PDF. To je nejpřímější '
                  'cesta, jak porovnat výsledek s jiným validátorem postaveným nad DSS.')


def kapitola_15(doc):
    nadpis1(doc, '15', 'Zvláštní obsahy')
    odstavec(doc, 'Věci, které se v českém prostředí potkávají a běžný validátor je nezná.')

    nadpis2(doc, '15.1', 'Datové zprávy ISDS (.zfo)')
    odstavec(doc, 'Ve zprávě ISDS je průkazný otisk uložený v kvalifikovaném časovém razítku '
                  'zprávy, ne v poli `dmHash`, které je jen jeho kopií. Aplikace pracuje se '
                  'syrovými bajty zprávy, otisk přepočítá a porovná s tím, co kryje razítko — '
                  'tím je zpráva, kterou čtenář vidí, svázána s okamžikem, kdy ji ISDS přijal.')
    odstavec(doc, 'Zpráva se přitom nečte jako obálka CMS s jedenácti kilobajty XML uvnitř, '
                  'nýbrž jako zpráva: kdo komu poslal, pod jakou značkou, kdy byla dodána a kdy '
                  'doručena, a co v ní bylo za písemnosti. Každá písemnost nese vlastní bajty, '
                  'takže se dá prohlédnout, uložit i ověřit samostatně; tlačítko „Uložit '
                  'písemnosti…“ je uloží všechny pod názvy, které jim dal odesílatel. Aplikace '
                  'rozumí číselníkům ISDS, takže zprávu popíše i věcně, nejen kryptograficky.')

    nadpis2(doc, '15.2', 'Evidenční záznamy (ERS, XMLERS)')
    odstavec(doc, 'Aplikace evidenční záznamy podle RFC 4998 i XMLERS ověří a řekne, jestli '
                  'a k jakému okamžiku dokládají existenci předloženého souboru — a co z nich '
                  'naopak neplyne. Vytvářet je neumí.')
    odstavec(doc, 'Kryje-li záznam víc souborů než ten předložený, otevře se okno, kde se '
                  'ostatní přiloží a zkontrolují proti otiskům záznamu. Plný verdikt se vydává '
                  'vždy k jednomu dokumentu, takže na kterýkoli z nich se odtud dá spustit '
                  'obvyklé ověření s týmž záznamem — míchat několik souborů do jednoho výsledku '
                  'by čtenáři nechalo hádat, který prvek patří ke kterému souboru.')

    nadpis2(doc, '15.3', 'Odkaz z eLegalizace')
    odstavec(doc, 'Příloha `legalisation-metadata.xml` odkazuje otisky na jiný dokument, který '
                  'v souboru není a nikdy nebude. Aplikace tu vazbu rozpozná, po ověření sama '
                  'otevře okno a nechá druhý dokument přiložit. Ověří přitom **všechny '
                  'deklarované otisky**, ne první, který umí spočítat: mluví o týchž bajtech, '
                  'takže neshoda mezi nimi je sama o sobě nález.')
    odstavec(doc, 'Kontrola je o odkazu, ne o podpisech: shodné otisky říkají, že dvojice '
                  'souborů patří k sobě, a vůbec nic o tom, zda je kterýkoli z nich platně '
                  'podepsán. Potvrzený řetěz jde i do protokolu. Nese-li odkázaný dokument '
                  'odkaz na další, pokračuje se stejně.')

    nadpis2(doc, '15.4', 'Podpisy JAdES')
    odstavec(doc, 'JSON s podpisem podle JWS se v prohlížeči kontejnerů rozebere na součástky: '
                  'vložený dokument — u obalujícího JAdES ho odtud jinak nikdo nevydá —, '
                  'chráněná hlavička po parametrech, certifikáty podepisujícího a každá '
                  'nepodepsaná vlastnost (časová razítka, certifikáty, revokační data) jako '
                  'samostatný prvek, v pořadí, ve kterém opravdu leží. To pořadí v JAdES něco '
                  'znamená: archivní razítko kryje to, co mu předchází.')

    nadpis2(doc, '15.5', 'Vizualizace a původ podpisu')
    odstavec(doc, 'Z PDF se vytáhne obrázek podpisu tak, jak ho vidí čtenář — včetně otočených '
                  'stránek —, a z metadat i program, kterým byl dokument podepsán. Aplikace '
                  'také rozlišuje podpisová pole, která zůstala prázdná, a pole, jejichž vzhled '
                  'se nevykreslí (skrytá anotace, nulová plocha).')

    nadpis2(doc, '15.6', 'Zámek dokumentu')
    odstavec(doc, 'Pokud podpis nese omezení dalších změn (`/Lock`, FieldMDP nebo certifikační '
                  'DocMDP), aplikace ukáže, co přesně je uzamčeno a jaké změny autor povolil. '
                  'Čtou se obě poloviny zápisu, protože každá nese jen část odpovědi: rozsah '
                  'plyne z podepsaného FieldMDP, úroveň povolených změn jen z `/Lock`. '
                  'Uzamčení po podpisu **není** certifikační podpis a běžný nástroj o něm '
                  'nemusí říct nic. Nebrání úkonům zachování ověřitelnosti.')

    nadpis2(doc, '15.7', 'Prázdná podpisová pole')
    odstavec(doc, 'Pole připravené k podpisu se vypíše i tehdy, když v něm podpis není — včetně '
                  'toho, jestli by byl podpis v něm vůbec vidět a co by se jím uzamklo. '
                  'Podepsání totiž vzhled pole nemění; předpovědět se tedy dá dopředu.')

    nadpis2(doc, '15.8', 'Protipodpis')
    odstavec(doc, 'Podpis, jehož předmětem není dokument, nýbrž jiný podpis pod ním, se na '
                  'kartě pozná a řekne se, čí podpis podepisuje. Je to jiný vztah než dva '
                  'podpisy vedle sebe a čtenáři, který to nerozliší, změní smysl celého '
                  'dokumentu.')

    nadpis2(doc, '15.9', 'Podpisy v OOXML')
    odstavec(doc, 'U DOCX, XLSX a PPTX se plné ověření nedělá: formát potřebuje transformaci '
                  '(RelationshipTransform), kterou DSS nezná, takže **není doloženo, že podpis '
                  'kryje obsah, který čtenář vidí**. Karta proto nemá barvu výsledku — má '
                  'vlastní šedomodré „neověřeno“, a je to jediná poctivá odpověď. Jméno '
                  'podepsané osoby v obvyklé kartě by čtenáři řeklo něco, co nikdo neprokázal.')
    odstavec(doc, 'Ukáže se, co podpis o sobě tvrdí — úroveň podle struktury (B / T / LT / LTA) '
                  'a jaké revokační informace nese —, a odděleně to, co se ověřit dá i tak, '
                  'protože na transformacích nezávisí: samotný token časového razítka '
                  'a vlastnosti certifikátu, tedy doba platnosti, kvalifikovanost podle '
                  'QcStatements a přítomnost vydavatele v evropských důvěryhodných seznamech. '
                  'To stačí k odhadu, jestli má ověření vůbec šanci. Forenzní balíček se '
                  'u takového souboru nabídne.')


def kapitola_16(doc):
    nadpis1(doc, '16', 'Nastavení — přehled')

    nadpis2(doc, '16.1', 'Základní nastavení')
    tabulka(doc,
            ['Volba', 'Význam'],
            [['Ověřování podpisů a pečetí dle standardů',
              'Vypnutí zapne tolerantní režim — čtveřici ústupků naráz.'],
             ['Ověřování časových razítek podle standardů',
              'Vypnutí znamená přísnější posuzování razítek podle řádné platnosti jejich '
              'certifikátů.'],
             ['Expertní mód',
              'Zpřístupní záložky Detailní rozbor a Protokoly DSS a některé úkony.']],
            [6.0, 10.6])
    odstavec(doc, 'Změny se projeví ihned, bez nutnosti restartu aplikace.')

    nadpis2(doc, '16.2', 'Pokročilé nastavení')
    odrazky(doc, [
        'Čas ověření — ověření k zadanému datu a času.',
        'Uvolnění validační politiky DSS — čtyři ústupky jednotlivě.',
        'Vlastní důvěryhodné certifikáty — správa kotev.',
        'Vlastní CRL seznamy — správa vlastního úložiště.',
        'Zdroj časových razítek (TSA) — profily se sondou.',
        'Důvěryhodné seznamy EU (EUTL) — opětovné načtení.',
    ])
    obrazek(doc, 'Dialog Nastavení v základní podobě.',
            'tři zaškrtávací pole a tlačítko Pokročilé nastavení', vyska=6.5)

    nadpis2(doc, '16.3', 'O aplikaci')
    odstavec(doc, 'Okno „O aplikaci“ uvádí, k čemu aplikace slouží a že její výstupy jsou bez '
                  'záruky, k tomu verzi, autora, licenci (EUPL 1.2) a odkaz na složku '
                  's licencemi použitých knihoven — kterých je přes sto.')


def kapitola_17(doc):
    nadpis1(doc, '17', 'Řešení potíží')
    tabulka(doc,
            ['Projev', 'Co s tím'],
            [['Ověření trvá dlouho při prvním spuštění',
              'Stahují se důvěryhodné seznamy. Děje se to jednou; dále se použije mezipaměť.'],
             ['Stav platnosti zůstal neznámý u staršího podpisu',
              'Podívejte se do Analýzy na důvody a doporučení. Nejčastěji chybějí revokační '
              'informace — zkuste ústupky, vlastní úložiště CRL, nebo přijměte, že už to nejde.'],
             ['Nedůvěryhodná kotva u českého komerčního certifikátu',
              'Vložte certifikáty komerčních autorit tlačítkem v pokročilém nastavení. Závěr '
              'pak ale stojí na kotvě, kterou jiný nástroj nemá — aplikace to u něj uvede.'],
             ['Chybí důvěryhodné seznamy některé země',
              'Hlavička nad kartami v Rozšířeném přehledu a v Analýze to řekne. V pokročilém '
              'nastavení použijte „Načíst seznam EUTL znovu“; kvalifikovaná pečeť by se jinak '
              'mohla hlásit jako pouhá zaručená.'],
             ['Úkon se nenabízí, ačkoli je podpis platný',
              'Záložka Digitální kontinuita vždy uvede důvod u konkrétního prvku. Nejčastěji '
              'jde o porušenou integritu, evidenční záznam, nebo o to, že není co doplnit.'],
             ['Vybraný zdroj razítek ověřitelnost neposune',
              'Certifikát autority končí dřív než dosavadní ověřitelnost dokumentu. Vyberte '
              'zdroj s déle platným certifikátem, nebo s úkonem počkejte.'],
             ['Aplikace nepřečte soubor',
              'Formát se pozná podle obsahu, ne podle přípony. U odloučeného podpisu je nutné '
              'přiložit i podepsaný obsah. Formát, který aplikace nezná, se dá aspoň zafixovat '
              'do forenzního balíčku.'],
             ['Vložený soubor se místo ověření otevřel v prohlížeči',
              'Nebyl to podepsaný dokument, nýbrž certifikát, seznam CRL nebo odpověď OCSP. '
              'Stavový řádek to říká.']],
            [5.4, 11.2])


POJMY = [
    ('Rozhodný okamžik (BST)',
     'Čas, ke kterému se posuzuje platnost podpisu. Odvozuje se z časového razítka; bez něj je '
     'to okamžik ověření.'),
    ('Důkaz existence (PoE)',
     'Doklad, že podpis existoval již v dřívějším čase — typicky časové razítko.'),
    ('Vyřazení služby',
     'Okamžik, kdy služba ztratí kvalifikovaný status v důvěryhodných seznamech. Předem se '
     'nedá zjistit. Razítka vydaná předtím zůstávají uznatelná, pozdější ne.'),
    ('Úroveň B / T / LT / LTA',
     'B — holý podpis; T — s časovým razítkem; LT — s vloženými revokačními informacemi; '
     'LTA — s archivním razítkem, které je fixuje.'),
    ('LOTL / EUTL',
     'Evropský seznam důvěryhodných seznamů a národní seznamy pod ním.'),
    ('Validační informace',
     'Vše, co je pro pozdější ověření potřeba doplnit — především revokační informace '
     'a certifikáty.'),
    ('Důvod',
     'Číslovaný výklad nálezu pro čtenáře: co se stalo, co to znamená a co s tím.'),
    ('Referenční formát',
     'Formát, který jako referenční stanoví prováděcí rozhodnutí Komise (EU) 2015/1506; orgány '
     'veřejné moci s ním musí umět pracovat.'),
    ('Protipodpis (kontrasignatura)',
     'Podpis, jehož předmětem není dokument, nýbrž jiný podpis pod ním.'),
    ('Evidenční záznam',
     'Doklad podle RFC 4998 nebo XMLERS o tom, že soubor existoval už v určitou dobu; '
     's podpisem nemá nic společného.'),
    ('Ústupek',
     'Povolená odchylka od standardů, kterou uživatel zapne vědomě a která platí jen na jedno '
     'ověření.'),
]


def kapitola_18(doc):
    nadpis1(doc, '18', 'Slovníček')
    tabulka(doc,
            ['Pojem', 'Význam'],
            [[pojem, vyklad] for pojem, vyklad in POJMY],
            [4.6, 12.0])
    odstavec(doc, 'AnEP — Analyzátor elektronických podpisů · ověřování podle nařízení eIDAS '
                  '(910/2014) a norem ETSI EN 319 102-1 a navazujících, nad knihovnou DSS '
                  'Evropské komise · EUPL 1.2',
             velikost=9, barva=INKOUST_SVET, pred=14)


# ── sestavení ────────────────────────────────────────────────────────────────────────────

def main():
    doc = Document()
    zaklad(doc, TITULEK)

    obalka(doc)
    obsah(doc)

    kapitola_1(doc)
    kapitola_2(doc)
    kapitola_3(doc)
    kapitola_4(doc)
    kapitola_5(doc)
    kapitola_6(doc)
    kapitola_7(doc)
    kapitola_8(doc)
    kapitola_9(doc)
    kapitola_10(doc)
    kapitola_11(doc)
    kapitola_12(doc)
    kapitola_13(doc)
    kapitola_14(doc)
    kapitola_15(doc)
    kapitola_16(doc)
    kapitola_17(doc)
    kapitola_18(doc)

    doc.core_properties.title = 'AnEP — uživatelský manuál pro pokročilé uživatele'
    doc.core_properties.subject = 'Analyzátor elektronických podpisů'
    doc.core_properties.language = 'cs-CZ'
    doc.core_properties.comments = ('Vygenerováno skriptem docs/gen-manual-pokrocily.py; '
                                    'obrázky se vkládají ručně do připravených rámečků.')

    doc.save(CIL)
    print('hotovo: ' + CIL + ' (' + str(pocet_obrazku()) + ' rámečků pro obrázky)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
