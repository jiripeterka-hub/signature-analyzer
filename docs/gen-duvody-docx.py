# -*- coding: utf-8 -*-
"""Vygeneruje `docs/AnEP-duvody.docx` — číselník důvodů ve Wordu.

Čte týž výpis jako `gen-duvody-tisk.py`: ten, který `ReasonDump` vytáhne přímo z kódu
(názvy, texty i vazby mezi důvody), takže ani wordová podoba nemůže tvrdit něco jiného než
aplikace. Postup:

    java -Dstdout.encoding=UTF-8 \
         -cp "app/target/test-classes;app/target/classes;$(cat cp.txt)" \
         cz.digitalnikontinuita.anep.ui.ReasonDump > duvody.txt
    python3 docs/gen-duvody-tisk.py duvody.txt      # HTML pro tisk
    python3 docs/gen-duvody-docx.py duvody.txt      # tenhle — Word

**Bez situací.** HTML list je technický číselník a u každého důvodu uvádí, z kterých situací
vzniká; wordová podoba je pro čtenáře a podle pravidla pro dokumenty v `docs/` o situacích
nemluví (rozhodnutí uživatele 14. 10. 2026). Zůstávají texty důvodu a vazby mezi důvody —
co který potlačí a čím je potlačen — a žebříček přednosti.

Sazba přes `docx_kit.py`, stejně jako manuál pro pokročilé a příručka.

⚠ Přepíše `docs/AnEP-duvody.docx`. Bez argumentu se výpis čte ze standardního vstupu.

Potřebuje `python-docx`.
"""

import os
import re
import sys
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor

from docx_kit import (AKCENT, AKCENT_TMAVY, INKOUST, INKOUST_SVET, PISMO_LEHKE, SAZBA_CM,
                      drz_s_dalsim, nadpis1, nadpis2, odstavec, prostrkani, psat, spodni_linka,
                      tabulka, vysvetlivka, zaklad, zalom)

ZDE = os.path.dirname(os.path.abspath(__file__))
CIL = os.path.join(ZDE, 'AnEP-duvody.docx')
TITULEK = 'AnEP — číselník důvodů'


# ── výpis z kódu ─────────────────────────────────────────────────────────────────────────

def nacti(src):
    """Důvody z výpisu ReasonDump: číslo → pole (NAZEV, PROBLEM, DETAIL, RESENI, RADA,
    POTLACI, POTLACEN, …). Řádky situací (@SIT) se čtou, ale nepoužívají."""
    duvody, cur = {}, None
    for raw in src:
        line = raw.rstrip('\n').rstrip('\r')
        m = re.match(r'### DUVOD (\d+)$', line)
        if m:
            cur = int(m.group(1))
            duvody[cur] = {}
            continue
        if cur and '\t' in line and not line.startswith('@'):
            k, v = line.split('\t', 1)
            duvody[cur][k] = v
    return duvody


def cisla(pole):
    return [int(x) for x in re.findall(r'R(\d+)', pole or '')]


def velke(text):
    return text[:1].upper() + text[1:] if text else text


# ── sazba ────────────────────────────────────────────────────────────────────────────────

def obalka(doc, pocet):
    mezera = doc.add_paragraph()
    mezera.paragraph_format.space_after = Pt(30)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    run = par.add_run('ČÍSELNÍK DŮVODŮ')
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(AKCENT)
    prostrkani(run, 60)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run('AnEP')
    run.font.name = PISMO_LEHKE
    run.font.size = Pt(48)
    run.font.color.rgb = RGBColor.from_string(INKOUST)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(10)
    run = par.add_run('Analyzátor elektronických podpisů')
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(INKOUST_SVET)
    spodni_linka(par, AKCENT, sz=12, space=8)

    odstavec(doc, 'Proč ověření nedopadlo kladně, co to znamená a co s tím — všech '
                  + str(pocet) + ' důvodů, jak je aplikace uvádí v záložce Analýza.',
             velikost=12, po=14)
    par = doc.add_paragraph()
    psat(par, 'Stav ke dni ' + date.today().strftime('%-d. %-m. %Y' if os.name != 'nt' else '%#d. %#m. %Y')
              + ' · vygenerováno z textů a pravidel přímo v kódu aplikace',
         velikost=9, barva=INKOUST_SVET)
    zalom(doc)


def uvod(doc):
    nadpis1(doc, '', 'Jak číselník číst')
    odstavec(doc, '**Co je důvod.** Výklad pro čtenáře: co ověření zjistilo, co to znamená a co '
                  's tím jde dělat. Jeden důvod může shrnout i několik různých zjištění pod jedno '
                  'vysvětlení a jedno doporučení. Čísla důvodů se nepřečíslovávají — zrušený '
                  'důvod zůstává neobsazený, aby ostatní dál znamenaly totéž.')
    odstavec(doc, '**Z čeho se skládá.** Popis problému (vždy), pod ním **Technicky** — totéž '
                  'jazykem standardů —, **Co pomůže** — co může udělat běžný uživatel — a **Pro '
                  'zkušenější** — rada předpokládající pokročilé nastavení. Volitelné části '
                  'u některých důvodů chybějí.')
    odstavec(doc, '**Proč na sobě důvody závisí.** Vypsat všechna zjištění naráz by čtenáře '
                  'poslalo hledat příčinu tam, kde žádná není. Silnější nález proto ten slabší '
                  'umlčí: u podpisu s porušenou integritou nemá smysl radit, ať se povolí '
                  'tolerance k revokačním údajům.')
    vysvetlivka(doc, 'Jak číst vazby',
                '**Potlačí** — vyskytnou-li se oba důvody, vypíše se jen tenhle. **Je potlačen** — '
                'opačný pohled na totéž. Některé důvody aplikace navíc vůbec neuvede, vyplývá-li '
                'z okolností jiný; to se z vazeb nevyčte a popisuje to výklad v záložce Analýza.')
    zalom(doc)


def duvod(doc, n, d, duvody):
    nazev = velke(d.get('NAZEV', ''))
    hlava = nadpis2(doc, '', 'Důvod ' + str(n) + ' — ' + nazev)
    drz_s_dalsim(hlava)
    if d.get('PROBLEM'):
        odstavec(doc, d['PROBLEM'])
    for klic, popiska in (('DETAIL', 'Technicky'), ('RESENI', 'Co pomůže'), ('RADA', 'Pro zkušenější')):
        if d.get(klic):
            par = doc.add_paragraph()
            psat(par, popiska + ': ', tucne=True, barva=AKCENT_TMAVY)
            psat(par, d[klic])

    def vycet(ns):
        return ', '.join(str(x) + ' (' + duvody.get(x, {}).get('NAZEV', '') + ')' for x in ns)

    vazby = []
    if cisla(d.get('POTLACI')):
        vazby.append('**Potlačí:** ' + vycet(cisla(d['POTLACI'])))
    if cisla(d.get('POTLACEN')):
        vazby.append('**Je potlačen:** ' + vycet(cisla(d['POTLACEN'])))
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(10)
    psat(par, ' · '.join(vazby) if vazby else 'Bez vazeb — nepotlačuje nic a nic ho nepotlačuje.',
         velikost=8.5, barva=INKOUST_SVET)


def zebricek(doc, duvody):
    nadpis1(doc, '', 'Žebříček přednosti')
    odstavec(doc, 'Kolik dalších důvodů který umlčí. Čím výš, tím spíš je právě on tím, co '
                  'čtenář uvidí jako jediný.')
    poradi = sorted(duvody, key=lambda r: (-len(cisla(duvody[r].get('POTLACI'))), r))
    radky = []
    for r in poradi:
        ns = cisla(duvody[r].get('POTLACI'))
        if ns:
            radky.append([str(r), velke(duvody[r].get('NAZEV', '')), str(len(ns)),
                          ', '.join(str(x) for x in ns)])
    tabulka(doc, ['Důvod', 'Název', 'Potlačí', 'Které'], radky,
            [1.6, SAZBA_CM - 1.6 - 1.6 - 4.2, 1.6, 4.2])
    bez = [r for r in sorted(duvody)
           if not cisla(duvody[r].get('POTLACI')) and not cisla(duvody[r].get('POTLACEN'))]
    odstavec(doc, 'Bez jakékoli vazby (nepotlačují nic a nic je nepotlačuje): '
                  + (', '.join(str(r) for r in bez) if bez else 'žádné') + '.', pred=6)


def main():
    src = open(sys.argv[1], encoding='utf-8') if len(sys.argv) > 1 else sys.stdin
    duvody = nacti(src)
    if not duvody:
        print('ve výpisu nejsou žádné důvody')
        return 1

    doc = Document()
    zaklad(doc, TITULEK)
    obalka(doc, len(duvody))
    uvod(doc)
    nadpis1(doc, '', 'Důvody')
    for n in sorted(duvody):
        duvod(doc, n, duvody[n], duvody)
    zalom(doc)
    zebricek(doc, duvody)

    doc.core_properties.title = TITULEK
    doc.core_properties.subject = 'Analyzátor elektronických podpisů — důvody výsledku ověření'
    doc.core_properties.language = 'cs-CZ'
    doc.core_properties.comments = ('Vygenerováno skriptem docs/gen-duvody-docx.py z výpisu '
                                    'ReasonDump (texty a vazby přímo z kódu).')
    doc.save(CIL)
    print('hotovo: ' + CIL + ' (' + str(len(duvody)) + ' důvodů)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
