# -*- coding: utf-8 -*-
"""Vygeneruje `docs/AnEP-politika-overovani.docx` z `docs/politika-overovani.md`.

Politika ověřování má jediný zdroj — markdown vedle tohoto skriptu. Wordová podoba se z něj
jen vysází, aby se ty dvě nemohly rozejít: kdo mění politiku, mění markdown a pustí tenhle
skript. Sazbu drží `docx_kit.py`, sdílený s ostatními generátory v `docs/`.

Rozumí jen té podmnožině markdownu, kterou politika používá:

    # Titul                      obálka (první nadpis prvního řádu)
    ## 3 Nadpis kapitoly         nadpis1 — číslo se oddělí od textu
    ### 3.1 Nadpis oddílu        nadpis2
    - odrážka                    odrazky
    1. položka                   cislovane
    | a | b |  + |---|---|       tabulka; první řádek je záhlaví
    **tučně**, `kód`             uvnitř textu (zpracuje docx_kit.psat)

    python docs/gen-politika-overovani.py

⚠ Přepíše `docs/AnEP-politika-overovani.docx`.

Potřebuje `python-docx`.
"""

import os
import re

from docx import Document
from docx.shared import Pt, RGBColor

from docx_kit import (AKCENT, INKOUST_SVET, SAZBA_CM, cislovane, nadpis1, nadpis2, odrazky,
                      odstavec, prostrkani, psat, spodni_linka, tabulka, zaklad, zalom)

ZDE = os.path.dirname(os.path.abspath(__file__))
ZDROJ = os.path.join(ZDE, 'politika-overovani.md')
CIL = os.path.join(ZDE, 'AnEP-politika-overovani.docx')
TITULEK = 'AnEP — politika ověřování elektronických podpisů'


def cislo_a_text(nadpis):
    """„3.1 Kotvy“ → („3.1“, „Kotvy“); bez čísla → („“, celý text)."""
    m = re.match(r'^(\d+(?:\.\d+)*)\s+(.*)$', nadpis)
    return (m.group(1), m.group(2)) if m else ('', nadpis)


def bunky(radek):
    return [b.strip() for b in radek.strip().strip('|').split('|')]


def sirky_sloupcu(zahlavi, radky):
    """Rozdělí šířku sazby podle délky obsahu, s rozumným minimem pro úzké sloupce."""
    n = len(zahlavi)
    vahy = []
    for i in range(n):
        delky = [len(zahlavi[i])] + [len(r[i]) for r in radky if i < len(r)]
        vahy.append(min(max(delky), 70) + 8)
    soucet = sum(vahy)
    sirky = [max(2.4, SAZBA_CM * v / soucet) for v in vahy]
    nadbytek = sum(sirky) - SAZBA_CM
    if nadbytek > 0:
        nejsirsi = sirky.index(max(sirky))
        sirky[nejsirsi] -= nadbytek
    return sirky


def obalka(doc, titul, podtitul, udaje, stitek='POLITIKA OVĚŘOVÁNÍ PODPISU'):
    mezera = doc.add_paragraph()
    mezera.paragraph_format.space_after = Pt(30)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    run = par.add_run(stitek)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(AKCENT)
    prostrkani(run, 60)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run('AnEP')
    run.font.size = Pt(40)
    run.font.color.rgb = RGBColor.from_string(AKCENT)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(10)
    run = par.add_run(titul)
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(INKOUST_SVET)
    spodni_linka(par, AKCENT, sz=12, space=8)

    if podtitul:
        odstavec(doc, podtitul.strip('*'), velikost=12, po=14)
    if udaje:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(24)
        psat(par, udaje, velikost=9, barva=INKOUST_SVET)
    zalom(doc)


def main(zdroj=ZDROJ, cil=CIL, titulek=TITULEK, stitek='POLITIKA OVĚŘOVÁNÍ PODPISU'):
    """Vysází politiku; parametry použije i gen-politika-uchovavani.py."""
    with open(zdroj, encoding='utf-8') as f:
        radky = f.read().splitlines()

    doc = Document()
    zaklad(doc, titulek)

    # Obálka: první nadpis, za ním tučný podtitul a řádek s verzí — až po první kapitolu.
    i = 0
    titul = podtitul = udaje = ''
    while i < len(radky) and not radky[i].startswith('## '):
        r = radky[i].strip()
        if r.startswith('# '):
            titul = r[2:]
        elif r.startswith('**') and not podtitul:
            podtitul = r
        elif r and not udaje:
            udaje = r
        i += 1
    obalka(doc, titul, podtitul, udaje, stitek)

    odstavec_buf, odrazky_buf, cislovane_buf = [], [], []

    def vyprazdni():
        if odstavec_buf:
            odstavec(doc, ' '.join(odstavec_buf))
            odstavec_buf.clear()
        if odrazky_buf:
            odrazky(doc, list(odrazky_buf))
            odrazky_buf.clear()
        if cislovane_buf:
            cislovane(doc, list(cislovane_buf))
            cislovane_buf.clear()

    while i < len(radky):
        r = radky[i].rstrip()
        s = r.strip()
        if s.startswith('### '):
            vyprazdni()
            nadpis2(doc, *cislo_a_text(s[4:]))
        elif s.startswith('## '):
            vyprazdni()
            nadpis1(doc, *cislo_a_text(s[3:]))
        elif s.startswith('|'):
            vyprazdni()
            zahlavi = bunky(s)
            i += 1                                      # oddělovač |---|---|
            data = []
            while i + 1 < len(radky) and radky[i + 1].strip().startswith('|'):
                i += 1
                data.append(bunky(radky[i]))
            tabulka(doc, zahlavi, data, sirky_sloupcu(zahlavi, data))
        elif s.startswith('- '):
            if odstavec_buf or cislovane_buf:
                vyprazdni()
            odrazky_buf.append(s[2:])
        elif re.match(r'^\d+\.\s', s):
            if odstavec_buf or odrazky_buf:
                vyprazdni()
            cislovane_buf.append(re.sub(r'^\d+\.\s+', '', s))
        elif not s:
            vyprazdni()
        else:
            if odrazky_buf or cislovane_buf:
                vyprazdni()
            odstavec_buf.append(s)
        i += 1
    vyprazdni()

    doc.save(cil)
    print('zapsáno:', cil)


if __name__ == '__main__':
    main()
