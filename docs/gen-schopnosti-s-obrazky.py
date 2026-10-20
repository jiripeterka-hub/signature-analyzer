# -*- coding: utf-8 -*-
"""Vygeneruje `docs/AnEP-schopnosti-s-obrazky.docx` — ilustrovaný přehled schopností aplikace.

Proč ještě jeden dokument o schopnostech: `AnEP-schopnosti.docx` vzniká převodem ze
`schopnosti.html` a je to obrazovková stránka přelitá do Wordu — souvislý text bez obrázků.
Tenhle je psaný rovnou jako wordový dokument k dalším úpravám: má obálku, číslované kapitoly,
rámečky pro snímky obrazovky, které do něj vloží autor, a dodatky. Zdrojem je tenhle skript,
ne HTML — obojí se udržuje zvlášť, protože každé odpovídá na jinou otázku.

Sazbu drží `docx_kit.py`, sdílený s `gen-manual-pokrocily.py`; tady je jen obsah.

**Nemluví o situacích**: ty jsou vnitřní věc aplikace a uživatel je nikde nepotřebuje. Důvody
ano, včetně jejich stavby — nejprve popis problému, pak jeho řešení.

    python docs/gen-schopnosti-s-obrazky.py

⚠ Přepíše `docs/AnEP-schopnosti-s-obrazky.docx` včetně vložených obrázků. Kdo do wordové
podoby vkládal snímky nebo psal vlastní úpravy, přijde o ně — po vložení obrázků si dokument
uložte pod jiným jménem, nebo skript už nepouštějte.

Potřebuje `python-docx`.
"""

import os
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor

from docx_kit import (AKCENT, AKCENT_TMAVY, INKOUST, INKOUST_SVET, PISMO_LEHKE,
                      citace, nadpis1, nadpis2, nadpis3, obrazek, odrazky, odstavec,
                      pocet_obrazku, prostrkani, psat, spodni_linka, tabulka, vysvetlivka,
                      zaklad, zalom)

ZDE = os.path.dirname(os.path.abspath(__file__))
CIL = os.path.join(ZDE, 'AnEP-schopnosti-s-obrazky.docx')

STAV_KE_DNI = '9. 10. 2026'
TITULEK = 'AnEP — přehled schopností'


# ── obálka a obsah ───────────────────────────────────────────────────────────────────────

def obalka(doc):
    mezera = doc.add_paragraph()
    mezera.paragraph_format.space_after = Pt(30)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    run = par.add_run('PŘEHLED SCHOPNOSTÍ')
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(AKCENT)
    prostrkani(run, 60)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run('AnEP')
    run.font.name = PISMO_LEHKE
    run.font.size = Pt(48)
    run.font.color.rgb = RGBColor.from_string(INKOUST)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(10)
    run = par.add_run('Analyzátor elektronických podpisů')
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(INKOUST_SVET)
    spodni_linka(par, AKCENT, sz=12, space=8)

    odstavec(doc, 'Co aplikace umí přečíst, co o tom řekne a co s tím dokáže udělat.',
             velikost=12, po=14)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(24)
    psat(par, 'Windows · desktopová aplikace  ·  jádro ověřování DSS (Evropská komise)'
              '  ·  rozhraní česky  ·  licence EUPL-1.2',
         velikost=9, barva=INKOUST_SVET)

    obrazek(doc, 'Hlavní okno aplikace s kladným výsledkem ověření.',
            'celé okno aplikace, záložka Zjednodušený přehled', vyska=9.5)

    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(10)
    psat(par, 'Stav ke dni ' + STAV_KE_DNI + '. Dokument je určen k dalším úpravám: obrázky se '
              'vkládají do připravených rámečků, text je sázen wordovými styly, takže se dá '
              'přeformátovat změnou stylu.', velikost=8.5, barva=INKOUST_SVET)
    zalom(doc)


OSNOVA = [
    ('1', 'K čemu je AnEP', ['Tři možné výsledky', 'Pro koho je určen']),
    ('2', 'Co aplikace přečte', ['Nosiče a formáty', 'Zvláštní případy při čtení']),
    ('3', 'Práce s aplikací', ['Hlavní okno', 'Co se děje při ověření']),
    ('4', 'Šest záložek', ['Zjednodušený přehled', 'Rozšířený přehled', 'Analýza',
                           'Digitální kontinuita', 'Detailní rozbor', 'Protokoly DSS']),
    ('5', 'Důvody: popis problému a jeho řešení',
     ['Z čeho se důvod skládá', 'Ukázka', 'Proč se důvody navzájem potlačují', 'Číselník']),
    ('6', 'Vlastní prohlížeče',
     ['Proč vlastní', 'Prohlížeč kontejnerů', 'Prohlížeč certifikátů', 'Prohlížeč seznamu CRL',
      'Prohlížeč odpovědi OCSP', 'Přílohy dokumentu', 'Inspektor podpisové struktury',
      'Kudy se k prohlížečům dostat']),
    ('7', 'Důvěra a revokační údaje', ['Odkud pochází důvěra', 'Revokační údaje', 'Vlastní zásoby']),
    ('8', 'Přísnost a ústupky', ['Čtyři ústupky', 'Dvě samostatná nastavení',
                                 'Jednorázová politika', 'Rozhodující ústupek']),
    ('9', 'Co aplikace udělá s dokumentem',
     ['Zachovat ověřitelnost', 'Souběh více podpisů', 'Forenzní balíček', 'Jména výstupů',
      'Kdy se úkon nenabídne']),
    ('10', 'Digitální kontinuita', ['Tři čtení jedné otázky', 'Co určuje datum',
                                    'Když lhůtu stanovit nelze', 'Co udělat']),
    ('11', 'Protokoly a export', ['Tři protokoly', 'Dvě volby', 'Zprávy DSS']),
    ('12', 'Zvláštní obsahy', ['Datové zprávy ISDS', 'Evidenční záznamy', 'Odkaz z eLegalizace',
                               'Podpisy JAdES', 'Vizualizace a původ podpisu', 'Zámek dokumentu',
                               'Prázdná podpisová pole', 'Protipodpis', 'Podpisy v OOXML']),
    ('13', 'Nastavení a úložiště', ['Základní nastavení', 'Pokročilé nastavení',
                                    'Zdroje časových razítek', 'Kde aplikace ukládá svá data']),
    ('14', 'Provoz, měření a licence', ['Distribuce a start', 'Sondy a skeny', 'Licence']),
    ('15', 'Kde jsou hranice', []),
    ('A', 'Dodatek: přehled důvodů', []),
    ('B', 'Dodatek: pojmy, které dokument používá', []),
]


def obsah(doc):
    par = doc.add_paragraph(style='Heading 1')
    psat(par, 'Obsah')
    spodni_linka(par)
    for cislo, titul, podkapitoly in OSNOVA:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        psat(p, cislo + ' ', tucne=True, barva=AKCENT_TMAVY)
        psat(p, titul, tucne=True)
        if podkapitoly:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(7)
            p2.paragraph_format.left_indent = Cm(0.8)
            psat(p2, ' · '.join(podkapitoly), velikost=9, barva=INKOUST_SVET)
    zalom(doc)


# ── kapitoly ─────────────────────────────────────────────────────────────────────────────

def kapitola_1(doc):
    nadpis1(doc, '1', 'K čemu je AnEP')
    odstavec(doc, 'AnEP ověřuje elektronické podpisy, pečeti a časová razítka podle nařízení '
                  '**eIDAS** a norem **ETSI** — a hlavně vysvětlí výsledek česky někomu, kdo '
                  'není odborník na PKI.')
    odstavec(doc, 'Vlastní ověřování obstarává otevřená knihovna **DSS** Evropské komise, tedy '
                  'referenční implementace ETSI EN 319 102-1. Aplikace nad ní staví to, co DSS '
                  'záměrně nedělá: mění strohé `INDETERMINATE/REVOKED_NO_POE` ve větu, která '
                  'říká, co se stalo, proč, a co se s tím dá dělat.')

    nadpis2(doc, '1.1', 'Tři možné výsledky')
    odstavec(doc, 'Výsledek má tři podoby a každá nese svou barvu, kterou pak drží celá '
                  'aplikace i protokol:')
    tabulka(doc,
            ['Výsledek', 'Co znamená'],
            [['platný (zelená)',
              'Platnost se podařilo prokázat za pravidel, za kterých se ověřovalo. Pokud se '
              'o to zasloužil povolený ústupek, je to u výsledku napsáno.'],
             ['stav platnosti neznámý (žlutá)',
              'Není to závěr o neplatnosti. Platnost se nepodařilo prokázat — a aplikace vždy '
              'řekne, co konkrétně chybělo a jestli to jde ještě opatřit.'],
             ['neplatný (červená)',
              'Podpis prokazatelně neobstál: porušená integrita, nesprávná dvojice souboru '
              'a podpisu, revokace prokázaná k rozhodnému okamžiku.']],
            [4.6, 12.0])
    odstavec(doc, 'Prostřední případ je ten podstatný. Většina nesnází s elektronickými podpisy '
                  'nekončí slovem „neplatný“, ale tím, že chybí důkaz — a rozdíl mezi „neplatné“ '
                  'a „neprokázané“ je přesně to, co běžné nástroje zamlčují.')
    obrazek(doc, 'Tři barvy výsledku, jak je aplikace ukazuje ve Zjednodušeném přehledu.',
            'tři snímky karty výsledku pod sebou — zelená, žlutá, červená', vyska=8.0)

    nadpis2(doc, '1.2', 'Pro koho je určen')
    odstavec(doc, 'Pro toho, kdo dostal podepsaný dokument a potřebuje vědět, na čem je: '
                  'pro úředníka, archiváře, právníka i pro zvědavého uživatele. Aplikace přitom '
                  'slouží také k seznámení s problematikou elektronických podpisů — proto u '
                  'každého pojmu, který použije, řekne, co znamená. Své výstupy vydává '
                  '**bez záruky**; říká to okno „O aplikaci“ i každý protokol.')


def kapitola_2(doc):
    nadpis1(doc, '2', 'Co aplikace přečte')
    odstavec(doc, 'Formát se pozná podle **obsahu souboru, ne podle přípony**. Přejmenovaný '
                  'soubor tedy aplikaci nezmate a kontejner uvnitř kontejneru se otevře také.')

    nadpis2(doc, '2.1', 'Nosiče a formáty')
    tabulka(doc,
            ['Nosič', 'Co v něm aplikace ověří'],
            [['PDF',
              'PAdES všech úrovní (B-B až B-LTA), včetně podpisů PDF-NOT-ETSI z doby před '
              'eIDAS, podpisů PKCS #7 a dokumentových časových razítek'],
             ['ASiC-S / ASiC-E',
              '`.asice`, `.asics`, `.sce`, `.scs` i kontejner ASiC uložený s příponou `.zip`'],
             ['Archiv ZIP',
              'obyčejný archiv, který není kontejnerem ASiC, se otevře v prohlížeči obsahu; '
              'soubory v něm lze předat k ověření jednotlivě'],
             ['Odloučené a samostatné',
              'CAdES (`.p7s`, `.p7m`, `.cms`), XAdES (`.xml`), JAdES (`.json`, `.jose`, `.jws`)'],
             ['Datové zprávy',
              '`.zfo` — zprávy ISDS i doručenky, ověřené přes kvalifikované časové razítko '
              'nad zprávou'],
             ['ODF',
              'ODT, ODS, ODP — podpisy XAdES v `META-INF/documentsignatures.xml`'],
             ['Kancelářské OOXML',
              'DOCX, XLSX, PPTX — podpis se **neověřuje**, ale aplikace nahlédne dovnitř '
              '(viz kapitola 12)'],
             ['Evidenční záznamy',
              'RFC 4998 (`.ers`, `.er`) a XMLERS, samostatné i vložené v dokumentu'],
             ['Holý certifikát, CRL, OCSP',
              'nejsou to podepsané dokumenty, a tak se neověřují — otevřou se rovnou '
              've vlastním prohlížeči (kapitola 6)']],
            [4.2, 12.4])

    nadpis2(doc, '2.2', 'Zvláštní případy při čtení')
    odrazky(doc, [
        '**Odloučený podpis** — podepsaný obsah se přiloží jako druhý soubor; aplikace má pro '
        'něj vlastní tlačítko vedle plochy pro dokument.',
        '**Heslem chráněné PDF** — aplikace si o heslo řekne, ale až tehdy, když bez něj '
        'dokument opravdu neotevře. Vlastnické heslo, které čtení nebrání, se tedy neptá.',
        '**Osiřelý podpis** — podpis, který se nepodařilo zpracovat, se v přehledu neztratí: '
        'ohlásí se jako nerozpoznaný druh prvku, ne jako by tam nebyl.',
        '**Prázdné podpisové pole** — pole připravené k podpisu, ve kterém zatím podpis není, '
        'se vypíše také, včetně toho, jestli by bylo vidět a co by po podpisu uzamklo.',
        '**Smetí před hlavičkou souboru** — bajty, které musely ustoupit, aby šel dokument '
        'vůbec přečíst, se spočítají a řeknou; nic je nepodepsalo.',
    ])


def kapitola_3(doc):
    nadpis1(doc, '3', 'Práce s aplikací')

    nadpis2(doc, '3.1', 'Hlavní okno')
    odstavec(doc, 'Okno má nahoře plochu pro soubor a vedle ní tři tlačítka, pod nimi stavový '
                  'řádek a pod ním záložky s výsledkem.')
    odrazky(doc, [
        '**Plocha pro dokument** — soubor se do ní přetáhne, nebo se klepnutím vybere.',
        '**↺ Obnovit naposledy zadaný obsah** — v rohu plochy; vrátí poslední dvojici souborů, '
        'aby se ověření dalo zopakovat po změně nastavení bez opětovného hledání v adresářích.',
        '**Ověřit** a **Nastavení** — dvě tlačítka, nic víc.',
        '**Externí obsah** — tlačítko pro odloučený podpis nebo pro evidenční záznam '
        'k předloženému dokumentu. Dokud je prázdné, je tiché; s načteným souborem se rozsvítí.',
        '**Stavový řádek** jmenuje ověřený soubor, jeho velikost a počet nalezených prvků '
        'a nabízí, co k němu patří: otevřít jej v přidružené aplikaci, otevřít jako kontejner '
        'a podívat se dovnitř, ukázat přílohy dokumentu i s jejich počtem.',
        '**Záložku lze odpojit** do vlastního okna (pravé tlačítko na jejím oušku) a mít tak '
        'vedle sebe dvě části výsledku najednou.',
    ])
    obrazek(doc, 'Prázdné okno před ověřením: plocha pro přetažení souboru a tlačítka.',
            'okno po spuštění, bez načteného dokumentu', vyska=7.0)

    nadpis2(doc, '3.2', 'Co se děje při ověření')
    odstavec(doc, 'Ověření běží na pozadí, okno zůstává ovladatelné a průběh je vidět. Start '
                  'aplikace nečeká na síť: důvěryhodné seznamy se drží v místní mezipaměti '
                  'a jejich čerstvost se dokontroluje až za běhu. Ověřovací cesta se navíc při '
                  'startu předehřeje nad vzorovým dokumentem — bez sítě a bez uživatelských dat '
                  '—, takže první ověření netrvá znatelně déle než druhé.')
    odstavec(doc, 'Chybí-li důvěryhodné seznamy některé země nebo jsou-li uložené seznamy staré, '
                  'řekne to aplikace sama — v hlavičce nad kartami v Rozšířeném přehledu '
                  'a v Analýze. Je to údaj o tom, **co mohla mít k dispozici**, a bez něj by '
                  'chybějící kotva vypadala jako vlastnost dokumentu.')


def kapitola_4(doc):
    nadpis1(doc, '4', 'Šest záložek')
    odstavec(doc, 'Řazeno podle hloubky. První čtyři vidí každý, poslední dvě se ukazují jen '
                  'v **expertním módu**.')

    nadpis2(doc, '4.1', 'Zjednodušený přehled')
    odstavec(doc, 'Verdikt, druh prvku i s výkladem pojmu (zaručený, uznávaný, kvalifikovaný), '
                  'podepsaná osoba, vydavatel certifikátu, čas podpisu a vizualizace podpisu '
                  'tak, jak ji nese PDF. Kdo chce jen odpověď, dál už jít nemusí. '
                  'Změnil-li se po podpisu obsah dokumentu, upozorní na to hnědožlutým '
                  'řádkem s čísly dotčených stran a odkazem na Rozšířený přehled — jinak by '
                  'zelené „platný“ svádělo k domněnce, že dokument je ten podepsaný.')
    obrazek(doc, 'Zjednodušený přehled — karta jednoho podpisu.',
            'obsah záložky Zjednodušený přehled', vyska=7.5)

    nadpis2(doc, '4.2', 'Rozšířený přehled')
    odstavec(doc, 'Totéž v devíti pojmenovaných rozbalovacích blocích:')
    odrazky(doc, [
        'podepsaná osoba (u pečeti pečetící osoba, u razítka vydavatel razítka),',
        'vizualizace podpisu,',
        'důvody podpisu — tedy to, co podepsaná osoba uvedla jako důvod, místo a kontakt,',
        'certifikát,',
        'časová razítka — u každého i to, kdo je vydal, do kdy platí jeho certifikát '
        'a jaké algoritmy razítko slibuje,',
        'evidenční záznamy,',
        'rozhodný okamžik ověření,',
        'kontrola revokace,',
        'použité kryptografické algoritmy — podpisu i každého časového razítka, které ho '
        'kryje, s délkou klíče a u razítek zvlášť s hašovací funkcí jejich podpisu a otisku '
        'orazítkovaných dat.',
    ])
    odstavec(doc, 'U dokumentu, na který odkazuje příloha z eLegalizace, přibývá za kartami '
                  'prvků ještě rámeček s odkazovanými dokumenty. Je vedle karet, ne v nich: '
                  'odkaz nese soubor, ne ten který podpis.')
    odstavec(doc, 'Nad kartami prvků stojí blok o prázdných podpisových polích a — změnil-li '
                  'se po podpisu obsah stránek — rozbalený blok s hnědožlutým záhlavím '
                  '**Po podepsání se v dokumentu změnil obsah**. Uvede dotčené strany, '
                  'upozorní, že jiné validátory mohou takovou změnu hodnotit jinak (například '
                  'neplatností posledního podpisu nebo neznámým stavem jeho platnosti), '
                  'a nabídne náhled dokumentu v podobě kryté podpisem i v dnešní podobě '
                  '(po změnách).')
    obrazek(doc, 'Rozšířený přehled — rozbalovací bloky.',
            'obsah záložky Rozšířený přehled, ideálně s jedním rozbaleným blokem', vyska=8.0)

    nadpis2(doc, '4.3', 'Analýza')
    odstavec(doc, 'Proč to dopadlo takhle. Pro každý podpis, pečeť či samostatné razítko jedna '
                  'karta; podle toho, jak ověření dopadlo a o čem je u konkrétního dokumentu co '
                  'říct, se v ní objeví '
                  'tyto bloky:')
    tabulka(doc,
            ['Blok', 'Co v něm stojí'],
            [['Formát podpisu',
              'Druh formátu, zapouzdření, úroveň a odpověď na otázku, jestli jde o referenční '
              'formát podle prováděcího rozhodnutí Komise (EU) 2015/1506.'],
             ['Zjištěné důvody a doporučení',
              'Očíslované důvody — co se stalo a co s tím. Celá kapitola 5.'],
             ['Průběh ověřování',
              'Za jakých pravidel se ověřovalo, jak aplikace postupovala a čím to skončilo; '
              'uvnitř složený Podrobný výčet revokačních údajů, které měla po ruce, s odkazy '
              '„Ukaž“ a „Ulož“ u každého.'],
             ['Poznámky k výsledku',
              'Co na výsledku stojí za zmínku, i když verdikt nezměnilo. Uvnitř podblok '
              's připomínkami, které si postěžovala sama knihovna DSS.'],
             ['Možné právní účinky',
              'Jen tam, kde plynou přímo z výsledku a z textu předpisu — u platného '
              'kvalifikovaného podpisu odkaz na čl. 25 eIDAS a na § 6 zákona č. 12/2020 Sb. '
              'Zapsaný osobní doklad se hledá ve všech položkách serialNumber předmětu '
              'certifikátu, nejen v první.'],
             ['Certifikační podpis',
              'Že podpis certifikuje dokument a jaké změny po sobě povolil.'],
             ['Zámek dokumentu',
              'Co podpis uzamkl a v jakém rozsahu, je-li dokument zamčen.'],
             ['Protipodpis (kontrasignatura)',
              'U podpisu, který podepisuje jiný podpis, se řekne čí a co z toho plyne.'],
             ['Kde je šance na úspěšné ověření',
              'Kterým dalším nástrojům se dokument ověřit podaří a kterým ne — v možnosti, '
              'ne v jistotě.']],
            [5.0, 11.6])
    odstavec(doc, 'V expertním módu stojí nad kartami ještě bloky o dokumentu jako celku: '
                  '**Shoda s normou PDF/A**, **Co v souboru přibylo po podpisu** (přibyla-li '
                  'jen validační data a obsah se nezměnil) a **Požadavky na vkladovou '
                  'listinu**. Běžnému uživateli se neukazují, aby Analýzu nezahlcovaly.')
    odstavec(doc, 'Blok **Požadavky na vkladovou listinu** u každého PDF s vnitřními podpisy '
                  'posoudí technické požadavky ČÚZK — splňuje, nesplňuje, nebo nelze posoudit. '
                  'Vychází z postupu podatelny katastru, která ověřuje nejprve podle standardů '
                  'a jinak může využít službu SecuSign Validation: podpis platný s ústupky '
                  'kromě SHA-1 požadavek splňuje, SHA-1 vede k „nesplňuje“, podpis neznámé '
                  'platnosti k „nelze posoudit“ s odkazem „Ověřit s ústupky podatelny '
                  'katastru“, a u nereferenčního formátu blok upozorní, že podatelna jeho '
                  'přijímání chystá omezit.')
    obrazek(doc, 'Analýza — blok Zjištěné důvody a doporučení.',
            'karta podpisu s rozbaleným blokem důvodů', vyska=8.0)

    nadpis2(doc, '4.4', 'Digitální kontinuita')
    odstavec(doc, 'Jak dlouho dokument zůstane ověřitelný, kdo ho ještě ověří, a co a do kdy '
                  'je proto potřeba udělat. Tady také stojí tlačítka úkonů. Podrobně '
                  'v kapitolách 9 a 10.')
    obrazek(doc, 'Digitální kontinuita — lhůty a úkon, který je posune.',
            'obsah záložky Digitální kontinuita', vyska=8.0)

    nadpis2(doc, '4.5', 'Detailní rozbor')
    odstavec(doc, 'Vše, co o dokumentu vydá diagnostika DSS, po oddílech: výsledek ověření, '
                  'certifikát podepisující osoby, certifikační cesta, důvěryhodná služba '
                  'a její stav v seznamech, časová razítka i s jejich cestou, algoritmy, časy, '
                  'rozsah podpisu, podpisový slovník PDF, detekce změn po podpisu, záznamy VRI '
                  'a globální validační data, struktura podpisů CMS a kryté objekty. '
                  'U certifikátů, seznamů CRL i odpovědí OCSP jsou odkazy „Ukaž“ a „Ulož“ '
                  'a u nedůvěryhodného kořene i „Přidat do důvěryhodných“. '
                  'Jen v expertním módu. ==Na začátku záložky stojí tlačítko **Inspektor** podpisové struktury — podle formátu Inspektor PDF, CAdES, ASiC, XAdES nebo JAdES (kapitola 6.7). Karty prvků se jmenují stejně jako v Rozšířeném přehledu a nesou jméno podepsané osoby; označení, pod kterým prvek vede knihovna DSS, je v kartě jako řádek „Identifikátor DSS“, podle kterého se prvek dohledá v Protokolech DSS.==')

    nadpis2(doc, '4.6', 'Protokoly DSS')
    odstavec(doc, 'Surové zprávy knihovny DSS, jednoduchá i podrobná, ke stažení jako XML nebo '
                  'PDF. Jen v expertním módu.')


def kapitola_5(doc):
    nadpis1(doc, '5', 'Důvody: popis problému a jeho řešení')
    odstavec(doc, 'Když ověření nedopadne kladně, nestačí to oznámit. **Důvod** je výklad pro '
                  'čtenáře: co se stalo, co to znamená a co se s tím dá dělat. Důvody jsou '
                  'očíslované a je jich ==**35**==; jeden důvod přitom shrnuje i několik různých '
                  'zjištění pod jedno vysvětlení a jedno doporučení, protože čtenáře nezajímá '
                  'výčet nálezů, ale odpověď.')
    vysvetlivka(doc, 'Pozor na jmennou shodu',
                'Blok „Důvody podpisu“ v Rozšířeném přehledu je něco jiného: tam jde o důvod, '
                'který uvedla podepsaná osoba (pole Reason v PDF). Důvody v tomhle smyslu jsou '
                'jen v Analýze, v bloku „Zjištěné důvody a doporučení“.')

    nadpis2(doc, '5.1', 'Z čeho se důvod skládá')
    odstavec(doc, 'Každý důvod má stejnou stavbu a **stejné pořadí částí: nejprve popis '
                  'problému, pak jeho řešení.** Čtenář se tedy nejdřív dozví, na čem je, '
                  'a teprve potom, co s tím.')
    tabulka(doc,
            ['Část', 'Co obsahuje', 'Povinná'],
            [['Popis problému',
              'Co se stalo, běžnou řečí a bez zkratek. Stojí vždy první.', 'ano'],
             ['Technicky',
              'Totéž jazykem standardů — jména polí, mezní hodnoty, důvod, proč to standard '
              'zakazuje. Ukáže se pod odkazem „Podrobněji …“, pro toho, kdo chce nález ověřit '
              'sám.', 'ne'],
             ['Co pomůže',
              'Co může udělat běžný uživatel: které nastavení přepnout, co doplnit, kdy to '
              'zkusit znovu. Stojí vždy po popisu problému.', 'ano'],
             ['Pro zkušenější',
              'Rada, která předpokládá pokročilé nastavení — typicky doplnit vlastní CRL '
              'seznam nebo vlastní kotvu důvěry. Je-li na kartě víc důvodů s touto radou, '
              'stojí jednou, pod spodním z nich.', 'ne']],
            [3.4, 11.0, 2.2])
    odstavec(doc, 'Tam, kde žádné řešení není, se žádné nepředstírá: u ztracené digitální '
                  'kontinuity se místo rady řekne, že vadu už napravit nelze.')
    odstavec(doc, 'Některé důvody mají **náhradní znění** a aplikace vybírá podle toho, co '
                  'o prvku ví — jinak zní věta o razítku, jehož autorita mezitím expirovala, '
                  'a jinak o razítku, které neobstálo z jiné příčiny. Údaje z dokumentu '
                  '(jméno autority, datum konce platnosti) se do vět doplňují.')

    nadpis2(doc, '5.2', 'Ukázka')
    odstavec(doc, 'Takhle vypadá důvod 2 — revokační informace vydaná po expiraci certifikátu:')
    citace(doc, [
        ('Popis problému',
         'Aplikaci se podařilo získat (on-line) revokační informace, prokazující stav revokace '
         'certifikátu k rozhodnému okamžiku. Tyto ale byly vydány až po expiraci (skončení '
         'řádné doby platnosti) certifikátu, ke kterému se vztahují. Technické standardy však '
         'zakazují využití takovýchto revokačních informací.'),
        ('Technicky',
         'Jde o CRL seznam či odpověď OCSP serveru s hodnotou thisUpdate, která následuje až po '
         'notAfter toho certifikátu, ke kterému se vztahuje.'),
        ('Co pomůže',
         'Zvolit tolerantní režim ověřování: v Nastavení zrušit zaškrtnutí volby „Ověřování '
         'podpisů a pečetí dle standardů“, případně v rámci pokročilého nastavení zapnout jen '
         'volbu „Využij i revokační informace, vydané po expiraci certifikátu“, a opakovat '
         'ověření.'),
        ('Pro zkušenější',
         'Pokud je možné získat použitelný CRL seznam z vhodné doby, lze jej (v Pokročilém '
         'nastavení) vložit do úložiště vlastních CRL seznamů.'),
    ])

    nadpis2(doc, '5.3', 'Proč se důvody navzájem potlačují')
    odstavec(doc, 'Vypsat všechno naráz by čtenáře poslalo hledat příčinu tam, kde žádná není. '
                  'Silnější nález proto ten slabší umlčí: u podpisu s **porušenou integritou** '
                  'nemá smysl radit, ať se povolí tolerance k revokačním údajům — dokument je '
                  'změněný a žádné nastavení s tím nehne. Aplikace tedy vypisuje jen ty důvody, '
                  'které za daného stavu ještě něco znamenají. A nemá-li čím naplnit výčet, '
                  'nenabídne ani jeho nadpis.')

    nadpis2(doc, '5.4', 'Číselník')
    odstavec(doc, 'Úplný číselník důvodů se generuje **přímo ze zdrojového kódu aplikace** — '
                  'názvy, texty i vazby mezi důvody pocházejí z tříd, které o nich rozhodují, '
                  'takže list nemůže tvrdit něco jiného než aplikace. Dá se vytisknout '
                  'a rozdávat. Seznam všech ==35== důvodů je v dodatku A.')


def kapitola_6(doc):
    nadpis1(doc, '6', 'Vlastní prohlížeče')
    odstavec(doc, 'Ověření odpoví na otázku „platí to?“. Vedle ní ale stojí druhá, stejně častá: '
                  '**„co v tom souboru vlastně je?“** — co jsem to podepsal, co ten kontejner '
                  'nese, co přesně říká ta odpověď OCSP, o kterou se závěr opřel. Na tu odpovídají '
                  'vlastní prohlížeče a jsou to plnohodnotné nástroje, ne náhledy.')

    nadpis2(doc, '6.1', 'Proč vlastní')
    odrazky(doc, [
        'Systémový prohlížeč certifikátů **o odpovědi OCSP neví vůbec** a u certifikátu ukazuje '
        'jinou sadu polí, než na kterou se čtenář v tomhle nástroji ptá.',
        'Kontejner **není dokument**: je to složka v jednom souboru a jinak než otevřením se '
        'nedá zjistit, co je uvnitř.',
        'Prohlížeče čtou soubor **tak, jak je uložen** — včetně toho, čeho si DSS nevšímá: '
        'nepodepsaných položek, pořadí prvků, metadat archivu. Pohled knihovny DSS na tentýž '
        'soubor ukazují záložky s výsledkem; tohle je pohled na soubor samotný.',
        'Okna jsou **nemodální**: čtenář si otevře certifikát a dál se pohybuje po záložkách, '
        'nebo si vedle sebe položí dva prvky.',
    ])

    nadpis2(doc, '6.2', 'Prohlížeč kontejnerů')
    odstavec(doc, 'Otevře se tlačítkem „Otevřít jako kontejner“ ve stavovém řádku a zvládne '
                  'víc druhů nosičů, než napovídá jméno:')
    tabulka(doc,
            ['Nosič', 'Co se ukáže'],
            [['ZIP (ASiC, ODF, OOXML, prostý archiv)',
              'strom položek se složkami, u známých jmen i jejich role — podpis XAdES či CAdES, '
              'manifest, časové razítko, evidenční záznam, validační data'],
             ['CMS / PKCS #7',
              'podepsaný obsah, podepisující, jeho certifikáty a revokační data jako '
              'samostatné prvky'],
             ['XML',
              'strom prvků tak, jak jdou v souboru za sebou'],
             ['JAdES (JSON/JWS)',
              'vložený dokument, chráněná hlavička po parametrech, certifikáty podepisujícího '
              'a každá nepodepsaná vlastnost zvlášť'],
             ['Datová zpráva ISDS',
              'obálka zprávy a písemnosti, každá s vlastními bajty'],
             ['PDF s přílohami',
              'přílohy dokumentu; samotné PDF se otevírá v přidružené aplikaci']],
            [5.4, 11.2])
    odstavec(doc, 'U vybraného prvku se vypíšou jeho údaje — velikost uloženou i původní, '
                  'způsob komprese, kontrolní součet, čas změny — a vedle nich náhled: text, '
                  'obrázek, nebo šestnáctkový výpis u toho, co ani jedno není. XML se pro '
                  'čtení zformátuje. S prvkem pak jde udělat tohle:')
    odrazky(doc, [
        '**Uložit prvek…** — vytáhnout jej ven jako soubor;',
        '**Uložit vše…** — rozbalit celý kontejner do adresáře se zachováním struktury;',
        '**Uložit písemnosti…** — u datové zprávy jen to, co odesílatel poslal, pod jeho jmény;',
        '**Otevřít v aplikaci** — předat prvek programu přiřazenému jeho typu;',
        '**Otevřít jako kontejner** — vnořený kontejner ve vlastním okně vedle rodičovského;',
        '**Ověřit v AnEP** — poslat prvek zpět k plnému ověření, jako by jej uživatel vložil sám;',
        '**Kopírovat SHA-256** — otisk prvku do schránky.',
    ])
    odstavec(doc, 'Nad stromem je **filtr podle názvu**; s napsaným hledáním se rozbalí i to, '
                  'co je jinak složené. Validační data uložená jako samostatné položky mívají '
                  'strojová jména, ze kterých čtenář nepozná odpověď OCSP od certifikátu — '
                  'takovým se ve stromu dá čitelné jméno (`odpoved-ocsp-1.ors`), kdežto '
                  'skutečné jméno položky zůstává na kartě pod „Název v kontejneru“ a hledá se '
                  'v obou.')
    obrazek(doc, 'Prohlížeč kontejnerů — strom položek, karta prvku a náhled.',
            'okno prohlížeče kontejnerů nad kontejnerem ASiC-E', vyska=8.5)

    nadpis2(doc, '6.3', 'Prohlížeč certifikátů')
    odstavec(doc, 'Certifikát vypsaný ve slovech, ne v šestnáctkové soustavě. Kromě subjektu, '
                  'vydavatele, sériového čísla, doby platnosti a algoritmů řekne i to, co bývá '
                  'schované v rozšířeních:')
    odrazky(doc, [
        '**klíčová užití** česky (digitální podpis, neodvolatelnost, podepisování certifikátů '
        'i CRL, časová razítka, podepisování odpovědí OCSP a další) — a rozšířená zvlášť;',
        '**základní omezení** — jde o koncový certifikát, nebo o certifikační autoritu, a do '
        'jaké hloubky;',
        '**kde se hledá vydavatel a odkud se stahuje revokační údaj** — přístupový bod OCSP '
        'i distribuční body CRL;',
        '**prohlášení podle QcStatements** — je-li certifikát kvalifikovaný podle eIDAS, k čemu '
        'slouží (podpis, pečeť, autentizace webu), leží-li klíč v prostředku QSCD, jaký je '
        'limit transakce, jak dlouho se uchovávají údaje, podle práva kterých zemí byl vydán '
        'a jaké role podle PSD2 nese;',
        '**zapsané osobní doklady** — nese-li certifikát číslo občanského průkazu nebo pasu '
        'podle ETSI EN 319 412-1, vypíše se druh dokladu a jeho číslo samostatně, aby je '
        'čtenář nemusel hledat mezi atributy jména — i když předmět nese víc položek '
        'serialNumber, třeba číslo dokladu vedle klientského čísla vydavatele.',
    ])
    odstavec(doc, 'Tlačítkem „Uložit…“ jde certifikát vytáhnout do souboru přesně tak, jak '
                  'v dokumentu leží.')

    nadpis2(doc, '6.4', 'Prohlížeč seznamu CRL')
    odstavec(doc, 'Reálný seznam CRL má tisíce záznamů a otázka se přitom týká jediného. '
                  'Prohlížeč proto nad tabulkou vypíše hlavičku seznamu — vydavatele (s odkazem '
                  'na jeho certifikát), kdy byl vydán a kdy má vyjít další, číslo seznamu, '
                  'jestli je to delta, a kolik certifikátů uvádí — a **otevře se rovnou na '
                  'záznamu toho certifikátu, o který jde**. Není-li na seznamu, řekne to nad '
                  'tabulkou.')
    odstavec(doc, 'Tabulka nese sériové číslo, datum odvolání a jeho důvod, dá se **řadit podle '
                  'sloupců** a nad ní je **hledání sériového čísla, které bere šestnáctkový '
                  'i desítkový zápis**. To je cesta pro seznam otevřený samostatně, kde není '
                  'od čeho se odrazit.')
    obrazek(doc, 'Prohlížeč seznamu CRL otevřený na hledaném záznamu.',
            'okno prohlížeče CRL s hlavičkou a tabulkou', vyska=7.5)

    nadpis2(doc, '6.5', 'Prohlížeč odpovědi OCSP')
    odstavec(doc, 'Odpověď OCSP vypsaná ve slovech: stav odpovědi jako takové, kdo odpověděl, '
                  'kdy ji vydal, jaké certifikáty v sobě nese — a pro každý dotázaný certifikát '
                  'jeho sériové číslo, použitou hashovací funkci, stav (dobrý, revokovaný, '
                  'neznámý), datum a **důvod odvolání pojmenovaný slovem**, a doby platnosti '
                  'odpovědi. Uložit ji lze rovněž.')

    nadpis2(doc, '6.6', 'Přílohy dokumentu')
    odstavec(doc, 'Přílohy PDF se sbírají ze **tří různých míst** — ze jmenného stromu vložených '
                  'souborů (to je ta „sponka“), z anotací umístěných na stránce a z polí '
                  'přidružených souborů, kterými nosí data dokumenty typu PDF/A-3 — a u každé '
                  'se pamatuje, kudy se do dokumentu dostala. Uživatel mezi tím nerozlišuje, '
                  'a proto je vidí pohromadě. Ukazují se v témže prohlížeči jako kontejner, '
                  'se všemi jeho možnostmi; příloha bývá sama podepsaným dokumentem, takže '
                  '„Ověřit v AnEP“ je u ní ta nejužitečnější.')
    odstavec(doc, 'U heslem chráněného dokumentu se přílohy vytahují už dešifrované — na heslo '
                  'se tedy podruhé neptá.')

    nadpis2(doc, '6.7', 'Inspektor podpisové struktury')
    odstavec(doc, '==Jen v expertním módu, tlačítkem na začátku Detailního rozboru; podle formátu se jmenuje Inspektor PDF, CAdES, ASiC, XAdES nebo JAdES, a je-li k dokumentu předložen externí podpis, rozebírá ten. Ukazuje podpisovou strukturu tak, jak je v souboru uložena, ve stromu, který začíná sbalený: podpisy a pečetě, všechna časová razítka včetně vnořených v podpisech, u PDF slovník /DSS, záznamy /VRI a revize dokumentu — /DSS ve výsledné podobě s údajem, ve které revizi položka přibyla, i v podobě každé revize zvlášť; revizi bez jediného objektu pozná. Rozkládá podepisující, atributy a vlastnosti podpisu, razítka, certifikáty, seznamy CRL a odpovědi OCSP a přepočítává otisky podepsaných dat, razítek i podpisového certifikátu. Pruh nad stromem ukazuje revize PDF a rozsah, který vybraný podpis kryje. U prvků spárovaných s výsledkem ověření uvede odděleně od zjištění o souboru, co s nimi udělal DSS — výsledek podpisu či razítka, roli certifikátu v řetězu, který revokační údaj rozhodl a který nebyl přijat nebo použit. Odkazy vedou mezi souvisejícími prvky a ke všem výskytům téhož prvku; každý prvek jde uložit. O výsledku ověření Inspektor nerozhoduje.==')
    obrazek(doc, 'Inspektor PDF — strom podpisové struktury, pruh revizí a detail prvku.',
            'okno Inspektoru nad podepsaným PDF s rozbaleným podpisem', vyska=8.5)

    nadpis2(doc, '6.8', 'Kudy se k prohlížečům dostat')
    tabulka(doc,
            ['Odkud', 'Co se otevře'],
            [['Stavový řádek → Otevřít jako kontejner',
              'prohlížeč kontejnerů nad právě ověřeným souborem'],
             ['Stavový řádek → sponka',
              'přílohy dokumentu; v závorce je jejich počet'],
             ['Analýza → Podrobný výčet → Ukaž',
              'seznam CRL otevřený na dotčeném certifikátu, nebo odpověď OCSP'],
             ['Detailní rozbor → Ukaž',
              'certifikát, seznam CRL nebo odpověď OCSP z kterékoli cesty či slovníku'],
             ['==Detailní rozbor → Inspektor==', '==podpisová struktura dokumentu, nebo externího podpisu, byl-li předložen=='],
             ['Prohlížeč kontejnerů',
              'kterýkoli prvek uvnitř — i vnořený kontejner ve vlastním okně'],
             ['Vložení holého souboru do aplikace',
              'certifikát, seznam CRL nebo odpověď OCSP se rovnou otevře v prohlížeči; '
              'poznají se podle obsahu, ne podle přípony']],
            [6.0, 10.6])


def kapitola_7(doc):
    nadpis1(doc, '7', 'Důvěra a revokační údaje')

    nadpis2(doc, '7.1', 'Odkud pochází důvěra')
    odstavec(doc, 'Kotvy důvěry pocházejí z **evropských důvěryhodných seznamů (LOTL/EUTL)** — '
                  'kolem třiceti národních seznamů a několika tisíc certifikátů. Uloží se do '
                  'místní mezipaměti, takže start aplikace nečeká na síť; čerstvost se '
                  'dokontroluje na pozadí a seznam jde kdykoli načíst znovu ručně. Vedle toho '
                  'se čte i kořenové úložiště Windows.')
    odstavec(doc, 'Aplikace si přitom pamatuje, **co ze seznamu plyne**: nejen že je certifikát '
                  'důvěryhodný, ale i jakou službu poskytuje, pod jakým poskytovatelem a do kdy '
                  'byla ta služba v seznamu vedena. Bez toho by se nedalo posoudit časové '
                  'razítko od služby, která mezitím ze seznamu zmizela.')

    nadpis2(doc, '7.2', 'Revokační údaje')
    odstavec(doc, 'Stav certifikátu se zjišťuje živě — **napřed OCSP, pak CRL** jako záloha — '
                  'a doplňuje se o to, co dokument nese sám. U každého údaje aplikace sleduje, '
                  '**odkud pochází**: z dokumentu, z vlastního úložiště, nebo ze sítě. Právě to '
                  'rozhoduje, jestli půjde závěr zopakovat i jinde a později, a proto se to '
                  'v Průběhu ověřování píše u každého údaje zvlášť.')
    odstavec(doc, 'Údaje se rozdělují do tří skupin, protože „přípustný“ a „využitý“ není totéž: '
                  'o které se závěr opřel, které standardy připouštějí, ale závěr na nich '
                  'nestojí, a které použít nešlo — u každého z nich i proč. Tam, kde se na stav '
                  'revokace správně vůbec nikdo neptal (certifikát je sám kotvou důvěry), řekne '
                  'se i to; jinak by poctivé „nic se nezjišťovalo, a správně“ vypadalo jako '
                  'porucha.')

    nadpis2(doc, '7.3', 'Vlastní zásoby')
    odrazky(doc, [
        '**Vlastní úložiště CRL** — seznamy, které si uživatel nashromáždil. Hodí se u starých '
        'podpisů, kde už vydavatel nic nezveřejňuje. Je to **záložní zdroj**: sáhne se po něm '
        'tehdy, když se použitelný údaj nepodařilo získat obvyklou cestou. Ze seznamů od téhož '
        'vydavatele si aplikace vybere ten, který o certifikátu skutečně může vypovídat — '
        'seznam vydaný dřív, než certifikát vznikl, nebo dřív než rozhodný okamžik, mlčí '
        'o něčem jiném, než co se právě zjišťuje.',
        '**Vlastní kotvy důvěry** — certifikáty, které evropské seznamy neuvádějí, například '
        'interní firemní autorita. Jedním tlačítkem se vloží vydávající autority tuzemských '
        'komerčních linek (I.CA, PostSignum, eIdentity, NCA, KB); doplnit se dá libovolná '
        'další. Aplikace vždy přizná, že závěr stojí na kotvě, kterou jiný nástroj nemá odkud '
        'vzít.',
    ])
    odstavec(doc, 'Co ověření z vlastních zásob opravdu použilo, se při zachování ověřitelnosti '
                  '**zapíše do dokumentu** — jinak by uchovaný soubor postrádal právě ten '
                  'důkaz, na kterém jeho verdikt stál, a nikdo jiný by si ho už neopatřil.')
    obrazek(doc, 'Pokročilé nastavení — správa vlastních důvěryhodných certifikátů a CRL seznamů.',
            'okno Pokročilé nastavení', vyska=8.0)


def kapitola_8(doc):
    nadpis1(doc, '8', 'Přísnost a ústupky')
    odstavec(doc, 'Ve výchozím stavu se ověřuje **striktně podle standardů**; v základním '
                  'nastavení to drží jediná volba „Ověřování podpisů a pečetí dle standardů“. '
                  'Kdo ji zruší, povolí všechny ústupky naráz; v pokročilém nastavení se dají '
                  'povolit jednotlivě.')

    nadpis2(doc, '8.1', 'Čtyři ústupky')
    tabulka(doc,
            ['Ústupek', 'Co povolí'],
            [['Po expiraci',
              'revokační informaci vydanou až po skončení platnosti certifikátu, o kterém '
              'vypovídá'],
             ['Expirovaný vydavatel',
              'revokační informaci od vydavatele, jemuž mezitím vypršel vlastní certifikát'],
             ['SHA-1',
              'ignorovat zastaralost algoritmu SHA-1'],
             ['PDF-NOT-ETSI',
              'ověřovat i podpisy v nereferenčním formátu z doby před eIDAS']],
            [4.6, 12.0])

    nadpis2(doc, '8.2', 'Dvě samostatná nastavení')
    odrazky(doc, [
        '**Ověřit k zadanému datu a času** — celé ověření se posune do minulosti, což je '
        'jediný způsob, jak se dozvědět, jak by dokument dopadl tehdy.',
        '**Ověřování časových razítek podle standardů** — razítka se posuzují podle řádné '
        'platnosti certifikátů razítkovacích autorit namísto podle data, kdy byla služba '
        'vyřazena z důvěryhodných seznamů. Je to přísnější čtení, ne ústupek.',
    ])

    nadpis2(doc, '8.3', 'Jednorázová politika')
    vysvetlivka(doc, 'Každý ústupek platí jen na jedno ověření',
                'Po něm se nastavení vrací do výchozího stavu. Kladný verdikt se tedy nemůže '
                'tiše opřít o volbu, kterou uživatel zapnul před měsícem a dávno na ni '
                'zapomněl. Vlastních zásob se to netýká — ty platí, dokud je uživatel '
                'neodebere.')
    odstavec(doc, 'Pro jedno ověření zapne tři z ústupků — všechny kromě SHA-1 — i odkaz '
                  '„Ověřit s ústupky podatelny katastru“ v posouzení vkladové listiny; tak se '
                  'zjistí, zda by podpis obstál u podatelny katastru, která SHA-1 '
                  'nepřijímá.')

    nadpis2(doc, '8.4', 'Rozhodující ústupek')
    odstavec(doc, 'Aplikace měří, jestli se ústupek na výsledku opravdu **podílel**: ověření '
                  'proběhne znovu bez něj a porovná se. U kladného verdiktu se proto vypíšou '
                  'jen ty ústupky, bez kterých by závěr nevznikl — ne ty, které byly jen '
                  'zapnuté. Totéž platí pro vlastní zásoby a pro zadaný čas ověření, a do '
                  'protokolu jde týž výčet.')
    obrazek(doc, 'Pokročilé nastavení — jednotlivé ústupky a čas ověření.',
            'okno Pokročilé nastavení, horní část s volbami', vyska=7.5)


def kapitola_9(doc):
    nadpis1(doc, '9', 'Co aplikace udělá s dokumentem')
    odstavec(doc, 'Žádný z úkonů nesahá na podepsané bajty. Každý zapisuje **nový soubor** '
                  'a původní nechává být. Tlačítka jsou v záložce Digitální kontinuita, hned '
                  'pod textem, který říká, co udělají a co z toho čtenář bude mít.')

    nadpis2(doc, '9.1', 'Zachovat ověřitelnost')
    odstavec(doc, 'Doplní chybějící revokační informace a zafixuje je dalším časovým '
                  'razítkem, aby dokument přežil expiraci podpisového certifikátu i certifikátů '
                  'dosavadních razítek. Podle stavu dokumentu se nabízí jedna ze tří variant, '
                  'které se na tlačítku jmenují:')
    tabulka(doc,
            ['Varianta', 'Kdy se nabídne'],
            [['Uchovat',
              'ověření podle standardů obstálo bez ústupků, řada časových razítek je '
              'nepřerušená a krytá platným certifikátem; doplní, co chybí, a přidá archivní '
              'razítko (B-LTA) — výsledek by měl ověřit kterýkoli validátor'],
             ['Přerazítkovat',
              'jen u PDF, kde platnost obstojí jen při mírnějším posuzování (revokační údaje, '
              'které přísné čtení nepřipouští, ústupek, kotva z vlastního úložiště, razítko '
              'neobstálé kvůli vyřazení služby); přidá razítko nad dnešním stavem pro validátory, '
              'které razítka posuzují podle platnosti jejich certifikátů'],
             ['Prodloužit',
              'ověření obstálo jen díky tomu, že služba posledního razítka si dosud drží '
              'kvalifikovaný status, ačkoli certifikát toho razítka už vypršel nebo je řada '
              'přerušená; nové razítko doloží existenci dokumentu, dokud to ještě platí']],
            [4.0, 12.6])
    odstavec(doc, 'Tlačítko je vždy jen jedno — varianty se vylučují a uživatel mezi nimi '
                  'nevybírá; aplikace nabídne nejsilnější, na kterou dosáhne. Úkon je přitom '
                  'opakovatelný: dokument, který už jednou uchován byl, se dá uchovat znovu, '
                  'až bude po čase potřeba.')

    nadpis2(doc, '9.2', 'Souběh více podpisů')
    odstavec(doc, 'Nové razítko kryje **celý soubor**, takže se nedá přidat pro jeden podpis '
                  'a pro druhý ne. U dokumentu s více podpisy je proto úkon jeden a jeho účinek '
                  'se vypíše po prvcích: kterému podpisu pomůže, kterému ne a proč. Prvku, na '
                  'který úkon nedosáhne, přitom neuškodí — co je na něm dnes sporné, zůstane '
                  'sporné i potom.')

    nadpis2(doc, '9.3', 'Forenzní balíček')
    odstavec(doc, 'Pro dokumenty, které zachovat nelze — a také pro soubory, jejichž formát '
                  'aplikace vůbec nepřečte. Zabalí **nedotčený** originál do kontejneru ASiC-E '
                  'spolu se vším ověřovacím materiálem, který se podařilo sehnat, a se soupisem '
                  'toho, odkud co pochází — a opatří to časovým razítkem. Znalec pak má jistotu, '
                  'že v okamžiku zafixování všechny podklady existovaly.')
    odstavec(doc, 'Do balíčku se dá přidat i další soubor: smlouva se svými přílohami, zpráva '
                  's tím, co k ní bylo připojeno. Zarazítkovat je zvlášť by každý zafixovalo '
                  'v jiném okamžiku a neřeklo by to nic o tom, že patří k sobě.')
    obrazek(doc, 'Okno forenzního balíčku se seznamem souborů, které do něj půjdou.',
            'dialog Forenzní balíček', vyska=7.5)

    nadpis2(doc, '9.4', 'Jména výstupů')
    odstavec(doc, 'Výsledný soubor si nese v názvu značku úkonu a datum, do kdy vydrží — '
                  'například `smlouva_Uchovano_do_06_03_2029.pdf`. Značka je bez diakritiky '
                  'a bez mezer záměrně: jméno putuje e-mailem i do spisových služeb. Značky se '
                  'nevrství: opakovaný úkon tu předchozí nahradí, aby z názvu nevznikl řetěz.')
    odstavec(doc, 'Datum na konci není vlastností dokumentu ani podpisu, nýbrž důsledkem volby '
                  'zdroje časových razítek — sahá tak daleko, kam sahá platnost certifikátu '
                  'vybrané autority.')

    nadpis2(doc, '9.5', 'Kdy se úkon nenabídne')
    odstavec(doc, 'Úkon se nabídne jen tam, kde je možný **a zároveň poctivý**. Dokument, jehož '
                  'certifikát byl odvolán dřív, než se dá prokázat existence podpisu, uchovat '
                  'nelze — a aplikace to řekne, místo aby nabídla tlačítko, které nic nespraví. '
                  'Řekne také, když vybraný zdroj razítek ověřitelnost prakticky neposune, '
                  'protože jeho vlastní certifikát končí sotva později než dokument sám.')


def kapitola_10(doc):
    nadpis1(doc, '10', 'Digitální kontinuita')
    odstavec(doc, 'Otázka „do kdy je dokument ověřitelný“ nemá jednu odpověď, protože to není '
                  'vlastnost dokumentu, nýbrž vztah mezi dokumentem a tím, kdo ho ověřuje.')

    nadpis2(doc, '10.1', 'Tři čtení jedné otázky')
    tabulka(doc,
            ['Čtení', 'Za jakých podmínek'],
            [['AnEP',
              'tak, jak se právě ověřovalo — se zapnutými ústupky, vlastními zásobami '
              'a zadaným časem'],
             ['Podle standardů',
              'validátor bez jakýchkoli odchylek; kde to dnes už nejde, řekne se, dokdy to šlo'],
             ['CzechPOINT a podatelny',
              'nástroje s vlastní zásobou revokačních informací a shovívavostí k SHA-1, ale bez '
              'ohledu na vyřazení služby z důvěryhodných seznamů']],
            [4.6, 12.0])
    odstavec(doc, 'O třetím okruhu mluví aplikace vždy v možnosti („mohly by“): nezná jejich '
                  'verzi, nastavení ani obsah jejich zásob. Kde jde odhadnout, že se dokument '
                  'jinde ověřit **podaří**, řekne to rovněž — a řekne, čím to bude.')

    nadpis2(doc, '10.2', 'Co určuje datum')
    odstavec(doc, 'Prvek visí v čase na posledním článku, který ho drží, a lhůtu určuje ten, '
                  'který povolí nejdřív:')
    odrazky(doc, [
        '**Důkaz existence** — s platným časovým razítkem smí být podpisový certifikát dávno '
        'po expiraci; bez něj končí všechno s ním.',
        '**Dostupnost revokačních informací** — údaj v dokumentu ani ve vlastním úložišti '
        'nezvětrá, ten stahovaný ze sítě jen dokud ho vydavatel zveřejňuje.',
        '**Uznatelnost vloženého údaje** — že údaj v souboru zůstane, ještě neznamená, že půjde '
        'uznat: po expiraci svého vydavatele je to údaj s prošlým certifikátem. Vyváže z toho '
        'jen razítko nad celým souborem.',
        '**Uznání razítka** — kvalifikované drží do vyřazení služby z důvěryhodných seznamů, '
        'což je datum, které nikdo předem nezná; jinak končí s certifikátem razítkovací '
        'autority.',
    ])

    nadpis2(doc, '10.3', 'Když lhůtu stanovit nelze')
    odstavec(doc, 'Zvláštní a poměrně častý případ: dnešní kladný závěr stojí na revokačním '
                  'údaji staženém právě teď ze sítě a doba, po kterou se na jeho dostupnost dalo '
                  'spoléhat, už uplynula. Aplikace pak **žádné datum neuvede** — v nadpisu '
                  'zůstane holá otázka a u úkonu stojí „Nejlépe ihned“. Uvnitř bloku je '
                  'vysvětlení a s ním i datum, kam by ověřitelnost sahala, kdyby údaj dostupný '
                  'zůstal. Datum se tím neztrácí, jen přestává být slibem.')
    odstavec(doc, 'Druhý důvod pro „Nejlépe ihned“: nejnovější časové razítko prvku je '
                  'podepsané algoritmem, který podle ETSI TS 119 312 přestane být spolehlivý '
                  'dřív, než lhůta uplyne — dnes typicky klíč RSA o 2 048 bitech s datem '
                  '1. 1. 2027. Aplikace AnEP na to datum reaguje jen varováním, jiné validátory '
                  'ale takové razítko mohou odmítnout; nové razítko se silnějším algoritmem je '
                  'proto třeba přidat dřív. Podepisuje-li slabým klíčem i vybraný zdroj razítek, '
                  'upozorní na to blok „Jak dlouho je možné zachovat ověřitelnost?“.')

    nadpis2(doc, '10.4', 'Co udělat')
    odstavec(doc, 'Záložka má čtyři části: výklad pojmu, přehled možností, které aplikace '
                  'nabízí, pak bloky po jednotlivých prvcích a nakonec — má-li dokument prvků '
                  'víc — pohled na soubor jako celek. U každého prvku stojí čtyři odpovědi: '
                  'dokdy platnost ověří AnEP, dokdy ji ověří validátor podle standardů, dokdy '
                  'CzechPOINT a podatelny, a **co a do kdy je třeba udělat**. Kde lhůta '
                  'neexistuje, řekne se i to, a kde je pozdě, řekne se to rovnou.')


def kapitola_11(doc):
    nadpis1(doc, '11', 'Protokoly a export')

    nadpis2(doc, '11.1', 'Tři protokoly')
    odstavec(doc, 'Protokol se generuje ze Zjednodušeného přehledu, z Rozšířeného přehledu '
                  'i z Detailního rozboru — každý o té hloubce, ve které vznikl — a vždy jako '
                  '**PDF/A-2b**, tedy v archivním formátu, který přijímají Formátová pravidla '
                  'Národního archivu. Nese záhlaví s údaji o souboru a s okamžikem ověření, '
                  'vizualizaci podpisu v rámečku, výklad pojmů uvozený „tj.:“ a výčet toho, '
                  'čím se ověřování řídilo. Protokol z obou přehledů se jmenuje podle souboru, '
                  'o kterém vypovídá.')
    odstavec(doc, 'Do protokolu jde i to, co má čtenář na obrazovce sbalené — rozbalovací blok '
                  'se exportuje celý.')
    obrazek(doc, 'Protokol o ověření vysázený jako PDF/A.',
            'první strana vygenerovaného protokolu', vyska=9.0)

    nadpis2(doc, '11.2', 'Dvě volby')
    odstavec(doc, 'Před generováním se aplikace zeptá na dvě věci; obojí je ve výchozím stavu '
                  'vypnuté:')
    tabulka(doc,
            ['Volba', 'Důsledek'],
            [['Vložit ověřovaný soubor',
              'protokol ponese i soubor, o kterém vypovídá, takže se dá ověřit znovu z jediného '
              'souboru — vysází se ale jako PDF/A-3, protože PDF/A-2 přílohu, která sama není '
              'PDF/A, nepřipouští'],
             ['Opatřit protokol časovým razítkem',
              'razítko doloží, kdy protokol vznikl a že se od té doby nezměnil; pořizuje se od '
              'zdroje vybraného v nastavení, a nabídne se jen tehdy, když nějaký je']],
            [5.2, 11.4])

    nadpis2(doc, '11.3', 'Zprávy DSS')
    odstavec(doc, 'Vedle toho jdou v expertním módu stáhnout původní zprávy knihovny DSS — '
                  'jednoduchá i podrobná — jako XML nebo PDF. Hodí se, když je třeba ukázat, '
                  'co přesně řekla referenční implementace, bez jakéhokoli výkladu.')


def kapitola_12(doc):
    nadpis1(doc, '12', 'Zvláštní obsahy')
    odstavec(doc, 'Věci, které se v českém prostředí potkávají a běžný validátor je nezná.')

    nadpis3(doc, 'Datové zprávy ISDS (.zfo)')
    odstavec(doc, 'Průkazný otisk je v kvalifikovaném časovém razítku zprávy, ne v poli '
                  '`dmHash`, které je jen jeho kopií — aplikace ho proto přepočítá ze syrových '
                  'bajtů a porovná s tím, co kryje razítko. Zpráva se přitom nečte jako obálka '
                  'CMS s jedenácti kilobajty XML uvnitř, nýbrž jako zpráva: kdo komu poslal, '
                  'pod jakou značkou, kdy byla dodána a kdy doručena, a co v ní bylo za '
                  'písemnosti. Aplikace rozumí číselníkům ISDS, takže o zprávě řekne i to, '
                  'čím byla.')

    nadpis3(doc, 'Evidenční záznamy (ERS, XMLERS)')
    odstavec(doc, 'Ověří se a řekne se, jestli a k jakému okamžiku dokládají existenci '
                  'předloženého souboru — a co z nich naopak neplyne. Kryje-li záznam víc '
                  'souborů než ten předložený, otevře se okno, kde se ostatní přiloží '
                  'a zkontrolují proti otiskům záznamu; kterýkoli z nich jde odtud poslat '
                  'k plnému ověření s týmž záznamem.')

    nadpis3(doc, 'Odkaz z eLegalizace')
    odstavec(doc, 'Příloha `legalisation-metadata.xml` odkazuje otisky na jiný dokument, který '
                  'v souboru není a nikdy nebude. Aplikace tu vazbu rozpozná, po ověření sama '
                  'otevře okno a nechá druhý dokument přiložit — a **ověří všechny deklarované '
                  'otisky**, ne první, který umí spočítat: mluví o týchž bajtech, takže '
                  'neshoda mezi nimi je sama o sobě nález. Potvrzený řetěz jde i do protokolu. '
                  'Nese-li odkázaný dokument odkaz na další, pokračuje se stejně.')

    nadpis3(doc, 'Podpisy JAdES')
    odstavec(doc, 'JSON s podpisem podle JWS se rozebere na součástky: vložený dokument '
                  '(u obalujícího JAdES ho odtud jinak nikdo nevydá), chráněná hlavička po '
                  'parametrech, certifikáty podepisujícího a každá nepodepsaná vlastnost — '
                  'časová razítka, certifikáty, revokační data — jako samostatný prvek '
                  'v pořadí, ve kterém opravdu leží. To pořadí v JAdES něco znamená: archivní '
                  'razítko kryje to, co mu předchází.')

    nadpis3(doc, 'Vizualizace a původ podpisu')
    odstavec(doc, 'Z PDF se vytáhne obrázek podpisu tak, jak ho vidí čtenář — včetně otočených '
                  'stránek —, a z metadat i program, kterým byl dokument podepsán. Obojí je '
                  'vidět v obou přehledech i v protokolu.')

    nadpis3(doc, 'Zámek dokumentu')
    odstavec(doc, 'PDF může být uzamčeno samotným podpisem — a **není to certifikační podpis**: '
                  'takový dokument nemusí nést nic, čeho by si běžný nástroj všiml. Aplikace '
                  'čte obě poloviny zápisu (rozsah i úroveň povolených změn) a řekne, co je '
                  'uzamčeno. Zámek nic neříká o platnosti a nebrání zachování ověřitelnosti.')

    nadpis3(doc, 'Prázdná podpisová pole')
    odstavec(doc, 'Pole připravené k podpisu se vypíše i tehdy, když v něm podpis není — '
                  'včetně toho, jestli by byl podpis v něm vůbec vidět a co by se jím uzamklo.')

    nadpis3(doc, 'Protipodpis')
    odstavec(doc, 'Podpis, který nepodepisuje dokument, nýbrž jiný podpis, se na kartě pozná '
                  'a řekne se, čí podpis to podepisuje. Je to jiný vztah než dva podpisy vedle '
                  'sebe a čtenáři, který to nerozliší, změní smysl celého dokumentu.')

    nadpis3(doc, 'Podpisy v OOXML')
    odstavec(doc, 'DOCX, XLSX a PPTX aplikace **neověřuje** — formát potřebuje transformaci, '
                  'kterou DSS nezná, takže není doloženo, že podpis kryje obsah, který čtenář '
                  'vidí. Karta proto nemá barvu výsledku; má vlastní šedomodré „neověřeno“, '
                  'a je to jediná poctivá odpověď. Ukáže se, co podpis o sobě tvrdí (úroveň '
                  'podle struktury B / T / LT / LTA, vložené revokační informace) a odděleně '
                  'to, co se ověřit **dá i tak** a na transformacích nezávisí: samotný token '
                  'časového razítka a vlastnosti certifikátu — doba platnosti, kvalifikovanost '
                  'podle QcStatements a to, zda je jeho vydavatel v evropských důvěryhodných '
                  'seznamech. Nabídne se u nich forenzní balíček.')
    obrazek(doc, 'Nahlédnutí do podpisu v kancelářském dokumentu — karta bez barvy verdiktu.',
            'okno Podpis v kancelářském dokumentu', vyska=7.5)


def kapitola_13(doc):
    nadpis1(doc, '13', 'Nastavení a úložiště')

    nadpis2(doc, '13.1', 'Základní nastavení')
    odrazky(doc, [
        '**Ověřování podpisů a pečetí dle standardů** — jediný přepínač přísnosti pro běžného '
        'uživatele.',
        '**Ověřování časových razítek podle standardů** — přísnější posuzování razítek.',
        '**Expertní mód** — zpřístupní Detailní rozbor a Protokoly DSS a bloky '
        'o dokumentu v Analýze.',
        '**O aplikaci** — k čemu aplikace slouží, že její výstupy jsou bez záruky, dále verze, '
        'autor, licence a složka s licencemi použitých knihoven.',
    ])
    odstavec(doc, 'Změny se projeví ihned, bez restartu.')
    obrazek(doc, 'Základní nastavení.', 'okno Nastavení, základní podoba', vyska=7.0)

    nadpis2(doc, '13.2', 'Pokročilé nastavení')
    odstavec(doc, 'Jednotlivé ústupky, ověření k zadanému datu a času, správa vlastních '
                  'důvěryhodných certifikátů a vlastních CRL seznamů, zdroje časových razítek '
                  'a ruční obnova důvěryhodných seznamů EU.')

    nadpis2(doc, '13.3', 'Zdroje časových razítek')
    odstavec(doc, 'Zdroje se spravují jako pojmenované profily s adresou, přihlášením '
                  'a algoritmem otisku; veřejný zdroj Sectigo Qualified umí aplikace '
                  'předvyplnit. **Při zadání se zdroj otestuje**: zjistí se konec platnosti '
                  'jeho certifikátu — právě podle něj se pak počítá, kam až úkon posune '
                  'ověřitelnost — a výčet otisků, které autorita skutečně podporuje, aby se '
                  'pak nevybíral algoritmus, který odmítne. Nedostupnost zdroje se přitom '
                  'nezamění za odmítnutí: profil jde uložit i tak, jen se dopředu neví, kam '
                  'dosáhne. Zjistí se i **algoritmus a délka klíče**, kterým autorita razítka '
                  'podepisuje; přestane-li klíč být podle ETSI TS 119 312 spolehlivý dřív, než '
                  'skončí certifikát autority — dnes typicky RSA o 2 048 bitech s datem '
                  '1. 1. 2027 —, správa zdrojů na to upozorní a doporučí zdroj s delším klíčem.')
    obrazek(doc, 'Správa zdrojů časových razítek s výsledkem sondy.',
            'okno správy TSA profilů', vyska=7.0)

    nadpis2(doc, '13.4', 'Kde aplikace ukládá svá data')
    odstavec(doc, 'Vše uživatelské leží v `%APPDATA%\\SignatureAnalyzer`: nastavení a profily '
                  'zdrojů razítek, vlastní úložiště CRL, vlastní důvěryhodné certifikáty '
                  'a mezipaměť důvěryhodných seznamů. Do systému aplikace jinak nezasahuje.')


def kapitola_14(doc):
    nadpis1(doc, '14', 'Provoz, měření a licence')

    nadpis2(doc, '14.1', 'Distribuce a start')
    odstavec(doc, 'Aplikace se dodává jako instalátor `.msi` i jako přenosná verze, obojí '
                  's vlastním běhovým prostředím — uživatel nemusí nic doinstalovávat. Start je '
                  'zrychlený předehřátím ověřovací cesty a mezipamětí důvěryhodných seznamů.')

    nadpis2(doc, '14.2', 'Sondy a skeny')
    odstavec(doc, 'Vedle běžných jednotkových testů má projekt desítky **sond a skenů**, '
                  'které pouštějí skutečnou ověřovací cestu nad sbírkou podepsaných dokumentů '
                  'a hlásí, co vyšlo. Existují proto, že většina otázek v téhle oblasti se '
                  'přečtením kódu zodpovědět nedá: jestli změna nezhoršila jiné dokumenty, '
                  'jestli lhůta vychází i tam, kde ji nikdo nečeká, co karta dokumentu opravdu '
                  'řekne. Nejedno pravidlo v aplikaci vypadá, jak vypadá, jen proto, že měření '
                  'odporovalo tomu, co se zdálo zřejmé.')

    nadpis2(doc, '14.3', 'Licence')
    odstavec(doc, 'AnEP se šíří pod **Evropskou unijní veřejnou licencí (EUPL) verze 1.2**, '
                  'která má závazné české znění. Licence použitých knihoven a oznámení jejich '
                  'autorů se dodávají s každou distribucí; seznam se generuje z těch souborů, '
                  'které se opravdu balí, takže se od skutečnosti nemůže odchýlit.')


def kapitola_15(doc):
    nadpis1(doc, '15', 'Kde jsou hranice')
    odstavec(doc, 'Co aplikace **nedělá**, a je to záměr:')
    odrazky(doc, [
        '**Nepodepisuje.** Umí přidat časové razítko a doplnit ověřovací materiál, ale podpis '
        'ani pečeť nevytváří.',
        '**Nevydává právní stanoviska.** Právní účinky uvádí jen tam, kde plynou přímo '
        'z výsledku ověření a z textu předpisu.',
        '**Nepředvídá cizí validátory.** O jiných nástrojích mluví v možnosti, ne v jistotě — '
        'nezná jejich verzi, nastavení ani obsah jejich zásob.',
        '**Nevytváří evidenční záznamy.** Ověřit je umí, vydat ne.',
        '**Neověřuje podpisy v kancelářských formátech OOXML.** Nahlédne do nich a řekne, co '
        'o sobě tvrdí, ale verdikt nevydá.',
        '**Nezakrývá ústupky.** Kladný verdikt, který stojí na povoleném ústupku, to má u sebe '
        'napsáno.',
        '**Nedává záruku.** Výstupy slouží k seznámení s problematikou a k vlastnímu posouzení, '
        'ne jako závazný posudek.',
    ])


DUVODY = [
    (1, 'zastaralý formát (nereferenční)'),
    (2, 'revokační informace vydané po expiraci certifikátu'),
    (3, 'neznámý certifikát'),
    (4, 'expirovaný podpisový certifikát, bez časového razítka'),
    (5, 'expirovaný certifikát posledního časového razítka, při ověřování dle certifikátů'),
    (6, 'vložená revokační data nejde použít — jejich vydavatel mezitím expiroval'),
    (7, 'porušená integrita podpisu'),
    (8, 'nesprávná dvojice soubor + podpis'),
    (9, 'revokovaný certifikát'),
    (10, 'neznámý stav revokace'),
    (11, 'revokační informace nelze využít'),
    (12, 'revokační informace nedostupné'),
    (13, 'podpis před začátkem platnosti certifikátu'),
    (14, 'revokovaný certifikát certifikační autority'),
    (15, 'zastaralá hashovací funkce SHA-1'),
    (16, 'kryptograficky nedůvěryhodná hashovací funkce (MD5, RIPEMD-160)'),
    (17, 'neznámý SubFilter'),
    (18, 'chybí atribut signing-time v podpisu'),
    (19, 'nedůvěryhodné časové razítko — autorita není v důvěryhodných seznamech'),
    (20, 'podpis po expiraci certifikátu'),
    (21, 'některé časové razítko neslouží jako důkaz existence'),
    (22, 'on-line revokační informace prokazující revokaci; nelze ji dle standardů použít'),
    (23, 'porušená integrita časového razítka — vadu nelze napravit'),
    (24, 'vydavatel revokačních dat mimo dobu své platnosti k rozhodnému okamžiku'),
    (25, 'stažená revokační informace o certifikátu nic neříká'),
    (26, 'revokovaný certifikát; chybí důkaz, kdy podpis vznikl'),
    (27, 'razítko ve slovníku VRI nezakládá důkaz existence podpisu'),
    (28, 'podpisová hodnota se neověřila, obsah je přitom nezměněný'),
    (29, 'vložená revokační data od vydavatele, který neplatil už při jejich vydání'),
    (30, 'chybí jen důkaz, že vložená revokační informace v dokumentu byla včas'),
    (31, 'samostatné časové razítko neobstálo'),
    (32, 'evidenčnímu záznamu něco schází'),
    (33, 'chybí podepsaný soubor (externí podpis)'),
    (34, '==deklarovaný čas podpisu leží mimo platnost certifikátu=='),
    (35, '==hashovací funkce, kterou standardy nepřipouštějí (SHA3-224)=='),
]


def dodatek_a(doc):
    nadpis1(doc, 'A', 'Dodatek: přehled důvodů')
    odstavec(doc, 'Krátká jména všech ==35== důvodů. Každý z nich má v aplikaci popis problému '
                  'a jeho řešení v podobě podle kapitoly 5; úplné znění je v generovaném '
                  'číselníku.')
    tabulka(doc,
            ['č.', 'Důvod'],
            [[str(cislo), jmeno] for cislo, jmeno in DUVODY],
            [1.2, 15.4],
            klicovy_sloupec=True)


POJMY = [
    ('Rozhodný okamžik',
     'Čas, ke kterému se platnost posuzuje. Určuje ho časové razítko, jinak je to přítomnost.'),
    ('Důkaz existence',
     'Doklad o tom, že podpis existoval už v určitou dobu — typicky časové razítko. Bez něj '
     'nelze o expirovaném certifikátu tvrdit nic dobrého.'),
    ('Revokační informace',
     'Údaj o tom, zda byl certifikát odvolán: odpověď OCSP nebo seznam CRL.'),
    ('Kotva důvěry',
     'Certifikát, kterému se věří bez dalšího dokazování. Pochází z evropských důvěryhodných '
     'seznamů, z úložiště Windows nebo z vlastního úložiště uživatele.'),
    ('Úroveň podpisu (B, T, LT, LTA)',
     'Jak je podpis vybaven pro budoucnost: B je holý podpis, T má časové razítko, LT nese '
     'i revokační údaje, LTA navíc archivní razítko.'),
    ('Referenční formát',
     'Formát, který jako referenční stanoví prováděcí rozhodnutí Komise (EU) 2015/1506; orgány '
     'veřejné moci s ním musí umět pracovat.'),
    ('Digitální kontinuita',
     'Schopnost dokumentu zůstat ověřitelný i po letech — a to, co je pro ni třeba udělat, '
     'dokud je čas.'),
    ('Zachování ověřitelnosti',
     'Doplnění revokačních údajů a archivního časového razítka, kterým se dokument dostane na '
     'úroveň LTA. Do dokumentu se přidává, nic se v něm nepřepisuje.'),
    ('Vyřazení služby',
     'Okamžik, od kterého už služba není vedena v důvěryhodných seznamech. Razítka vydaná '
     'předtím zůstávají uznatelná, pozdější ne.'),
    ('Ústupek',
     'Povolená odchylka od standardů, kterou uživatel zapne vědomě a která platí jen na jedno '
     'ověření.'),
    ('Protipodpis',
     'Podpis, jehož předmětem není dokument, nýbrž jiný podpis pod ním.'),
    ('Evidenční záznam',
     'Doklad podle RFC 4998 nebo XMLERS o tom, že soubor existoval už v určitou dobu; '
     's podpisem nemá nic společného.'),
]


def dodatek_b(doc):
    nadpis1(doc, 'B', 'Dodatek: pojmy, které dokument používá')
    tabulka(doc,
            ['Pojem', 'Co znamená'],
            [[pojem, vyklad] for pojem, vyklad in POJMY],
            [4.4, 12.2])
    odstavec(doc, 'Aplikace sama vykládá pojmy na místě, kde je použije — v protokolu uvozením '
                  '„tj.:“, na obrazovce v textu karty. Tenhle dodatek je jen shrnutí pro '
                  'čtenáře dokumentu.', velikost=9.5, barva=INKOUST_SVET)


# ── sestavení ────────────────────────────────────────────────────────────────────────────

def main():
    doc = Document()
    zaklad(doc, TITULEK)

    obalka(doc)
    obsah(doc)

    kapitola_1(doc)
    kapitola_2(doc)
    kapitola_3(doc)
    kapitola_4(doc)
    kapitola_5(doc)
    kapitola_6(doc)
    kapitola_7(doc)
    kapitola_8(doc)
    kapitola_9(doc)
    kapitola_10(doc)
    kapitola_11(doc)
    kapitola_12(doc)
    kapitola_13(doc)
    kapitola_14(doc)
    kapitola_15(doc)
    dodatek_a(doc)
    dodatek_b(doc)

    doc.core_properties.title = 'AnEP — přehled schopností'
    doc.core_properties.subject = 'Analyzátor elektronických podpisů'
    doc.core_properties.language = 'cs-CZ'
    doc.core_properties.comments = ('Vygenerováno skriptem docs/gen-schopnosti-s-obrazky.py; '
                                   'obrázky se vkládají ručně do připravených rámečků.')

    doc.save(CIL)
    print('hotovo: ' + CIL + ' (' + str(pocet_obrazku()) + ' rámečků pro obrázky)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
